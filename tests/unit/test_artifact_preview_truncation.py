"""Audit #10 — a capped artifact preview must not look complete.

xlsx previews are capped at 10 sheets / 200 rows / 30 columns and pdf
previews at 30 pages, but ``truncated`` used to reflect only the
100000-char total cap. A 12-sheet workbook whose shown text is short
therefore reported ``truncated: False``. These black-box route tests
pin the corrected flag for both caps, plus the uncapped case.
"""

from __future__ import annotations

import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from forgeflow.api.dependencies import get_current_user
from forgeflow.api.hub_deps import resolve_tenant
from forgeflow.api.routers import runs
from forgeflow.config import get_settings
from forgeflow.documents.store import DocArtifactStore
from forgeflow.rbac.models import UserContext
from forgeflow.runtime.orchestrator import RunRecord, get_run_store, reset_run_store

TENANT = "tenant-preview"
OWNER = "owner-preview"


@pytest.fixture(autouse=True)
def _clean_state(tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "resource_store_root", str(tmp_path))
    reset_run_store()
    yield
    reset_run_store()


def _client() -> TestClient:
    app = FastAPI()
    app.dependency_overrides[resolve_tenant] = lambda: TENANT
    app.dependency_overrides[get_current_user] = lambda: UserContext(
        user_id=OWNER, role="admin"
    )
    app.include_router(runs.router, prefix="/prefix")
    return TestClient(app)


def _seed(artifact: dict) -> None:
    get_run_store().save(
        RunRecord(
            run_id="run-preview",
            thread_id="thread-preview",
            tenant_id=TENANT,
            agent_id=None,
            intent="preview truncation",
            status="running",
            outcome="pending",
            steps=[],
            errors=[],
            created_at="2026-10-29T00:00:00+00:00",
            actor_user_id=OWNER,
            session_id="session-preview",
            artifacts=[artifact],
        )
    )


def _xlsx_bytes(*, sheets: int, rows: int) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Sheet1"
    for index in range(1, sheets):
        book.create_sheet(f"Sheet{index + 1}")
    for sheet in book.worksheets:
        for row in range(1, rows + 1):
            sheet.cell(row=row, column=1, value=f"r{row}")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(*, pages: int) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_xlsx_over_sheet_cap_reports_truncated():
    ref = DocArtifactStore().put(_xlsx_bytes(sheets=12, rows=2), ".xlsx")
    _seed({"id": "book.xlsx", "format": "xlsx", "content_ref": ref})
    response = _client().get("/prefix/run-preview/artifact-preview?artifact_id=book.xlsx")
    assert response.status_code == 200, response.text
    assert response.json()["truncated"] is True


def test_xlsx_over_row_cap_reports_truncated():
    ref = DocArtifactStore().put(_xlsx_bytes(sheets=1, rows=201), ".xlsx")
    _seed({"id": "book.xlsx", "format": "xlsx", "content_ref": ref})
    response = _client().get("/prefix/run-preview/artifact-preview?artifact_id=book.xlsx")
    assert response.status_code == 200, response.text
    assert response.json()["truncated"] is True


def test_pdf_over_page_cap_reports_truncated():
    ref = DocArtifactStore().put(_pdf_bytes(pages=31), ".pdf")
    _seed({"id": "doc.pdf", "format": "pdf", "content_ref": ref})
    response = _client().get("/prefix/run-preview/artifact-preview?artifact_id=doc.pdf")
    assert response.status_code == 200, response.text
    assert response.json()["truncated"] is True


def test_small_preview_still_reports_not_truncated():
    ref = DocArtifactStore().put(_xlsx_bytes(sheets=1, rows=2), ".xlsx")
    _seed({"id": "book.xlsx", "format": "xlsx", "content_ref": ref})
    response = _client().get("/prefix/run-preview/artifact-preview?artifact_id=book.xlsx")
    assert response.status_code == 200, response.text
    assert response.json()["truncated"] is False
