"""Typed upload inputs must reach real handlers in the ReAct model loop."""

from __future__ import annotations

import io

import pytest
from docx import Document
from langchain_core.messages import AIMessage

from forgeflow.config import get_settings
from forgeflow.documents.store import DocArtifactStore
from forgeflow.repositories.memory.resource_repo import clear_resource_store
from forgeflow.resources.service import ResourceService, reset_resource_index
from forgeflow.runtime import orchestrator as orch
from forgeflow.runtime.orchestrator import RequestContext, TaskCreate
from forgeflow.runtime.react_executor import ReactExecutor
from forgeflow.runtime.tool_registry import load_default_bindings, reset_registry


@pytest.fixture(autouse=True)
def isolated(force_memory_backend, tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "resource_store_root", str(tmp_path / "resources"))
    clear_resource_store()
    reset_resource_index()
    orch.reset_run_store()
    reset_registry()
    load_default_bindings()
    yield
    clear_resource_store()
    reset_resource_index()
    orch.reset_run_store()
    reset_registry()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "filename,tool,path_key,name_key",
    [
        ("report.docx", "document.inspect", "document_paths", "document_names"),
        ("slides.pptx", "document.inspect", "document_paths", "document_names"),
        ("table.xlsx", "sheet.inspect", "sheet_paths", "sheet_names"),
        ("memo.txt", "textfile.inspect", "text_paths", "text_names"),
        ("report.pdf", "pdf.inspect", "pdf_paths", "pdf_names"),
    ],
)
async def test_typed_upload_paths_are_injected_and_model_paths_cannot_override(
    filename, tool, path_key, name_key
):
    resource = await ResourceService().register_file("tenant", name=filename, data=b"fixture")
    task = TaskCreate(context={"resources": [resource.id]})
    args, injected = ReactExecutor._effective_args(
        tool,
        {
            path_key: ["/forged/secret"],
            name_key: ["forged"],
            "paths": ["/forged/secret"],
        },
        task,
        RequestContext(tenant_id="tenant", role="manager"),
        [],
    )
    assert args[name_key] == [filename]
    assert "/forged/secret" not in args[path_key]
    assert path_key in injected
    assert name_key in injected


def test_missing_typed_paths_do_not_accept_model_fabrications():
    args, _ = ReactExecutor._effective_args(
        "document.edit",
        {
            "document_paths": ["/private/data.docx"],
        },
        TaskCreate(),
        RequestContext(),
        [],
    )
    assert "document_paths" not in args


@pytest.mark.asyncio
@pytest.mark.parametrize("name", ["说明.txt", "报告.docx", "数据.xlsx", "说明.md"])
async def test_document_upload_cannot_be_sent_to_python_lint(name):
    resource = await ResourceService().register_file("tenant", name=name, data=b"fixture")
    task = TaskCreate(intent="检查该文件是否正确，并进行优化", context={"resources": [resource.id]})
    tools = ReactExecutor._narrow_tools_for_declared_data(
        task,
        RequestContext(tenant_id="tenant", role="manager"),
        ["code.lint", "code.run", "textfile.inspect", "document.edit", "artifact.save"],
    )
    assert "code.lint" not in tools
    assert "code.run" not in tools
    assert "artifact.save" in tools


@pytest.mark.asyncio
async def test_python_attachment_still_offers_code_validation():
    resource = await ResourceService().register_file("tenant", name="main.py", data=b"print(1)")
    task = TaskCreate(context={"resources": [resource.id]})
    tools = ReactExecutor._narrow_tools_for_declared_data(
        task,
        RequestContext(tenant_id="tenant", role="manager"),
        ["code.lint", "textfile.inspect"],
    )
    assert "code.lint" in tools


class DocumentModel:
    model = "scripted-doc-model"
    _llm_type = "test"

    def __init__(self):
        self.tools = iter(["document.inspect", "document.edit", "artifact.save"])

    def bind_tools(self, tools):
        return self

    def bind(self, **kwargs):
        return self

    async def ainvoke(self, messages, **kwargs):
        tool = next(self.tools, None)
        if tool is None:
            return AIMessage(content="文档已修改，文件可下载。")
        # The model does not know any storage path, filename, or explicit edits.
        return AIMessage(content="", tool_calls=[{"name": tool, "args": {}, "id": tool}])


@pytest.mark.asyncio
@pytest.mark.parametrize("answer", ["已读取上一轮记录。", ""])
async def test_empty_model_answer_is_retried_and_not_silently_successful(monkeypatch, answer):
    class EmptyFirstModel(DocumentModel):
        calls = 0

        async def ainvoke(self, messages, **kwargs):
            self.calls += 1
            return AIMessage(content=answer if self.calls > 1 else "")

    model = EmptyFirstModel()
    monkeypatch.setattr(get_settings(), "agent_runtime_mode", "react")
    monkeypatch.setattr(
        orch,
        "_build_planner_models",
        lambda: (model, model, [{"slot": "strong", "model": model.model}]),
    )
    handle = await orch.run_task(
        TaskCreate(intent="回顾刚才的记录"),
        RequestContext(tenant_id="tenant", user_id="user", role="manager"),
    )
    record = orch.get_run_store().get(handle.run_id)
    assert model.calls >= 2
    if answer:
        assert record.llm["final_answer"]["text"] == answer
        assert record.status == "completed"
    else:
        assert record.status == "failed"
        assert any("未返回最终答复" in error for error in record.errors)


@pytest.mark.asyncio
async def test_react_loop_edits_the_uploaded_word_with_empty_model_arguments(monkeypatch):
    doc = Document()
    doc.add_paragraph("原名称需要更新。")
    buffer = io.BytesIO()
    doc.save(buffer)
    resource = await ResourceService().register_file(
        "tenant", name="uploaded.docx", data=buffer.getvalue()
    )
    explicit = [{"op": "replace_text", "match": "原名称", "replace": "新名称"}]
    task = TaskCreate(
        intent="修改文档并返回文件", context={"resources": [resource.id], "edits": explicit}
    )
    ctx = RequestContext(tenant_id="tenant", user_id="user", role="manager")
    args, _ = ReactExecutor._effective_args(
        "document.edit",
        {
            "edits": [{"op": "replace_text", "match": "原名称", "replace": "模型篡改"}],
        },
        task,
        ctx,
        [],
    )
    assert args["edits"] == explicit
    monkeypatch.setattr(get_settings(), "agent_runtime_mode", "react")
    model = DocumentModel()
    monkeypatch.setattr(
        orch,
        "_build_planner_models",
        lambda: (model, model, [{"slot": "strong", "model": model.model}]),
    )
    handle = await orch.run_task(task, ctx)
    record = orch.get_run_store().get(handle.run_id)
    artifact = next(a for a in record.artifacts if a.get("format") == "docx")
    output = DocArtifactStore().get(artifact["content_ref"])
    assert Document(io.BytesIO(output)).paragraphs[0].text == "新名称需要更新。"
    assert all(step["status"] == "ok" for step in record.steps)
