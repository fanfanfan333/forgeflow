"""INC43 T04 — DOCX document-editing layered pins (offline; no real model, no PG).

What this file nails down
-------------------------
1. **Layering** — ``apply_edits`` (the Tool layer) really rewrites the bytes,
   while ``resolve_intent`` (the LLM layer) never touches them.
2. **Recomputable diff** — ``compute_diff``'s four figures are exact for known
   inputs, including the "only the number changed" case.
3. **Honest verification** — ``verify_docx`` reports ``False`` only for a
   *measured* violation and ``None`` for an unmeasured dimension.
4. **Handler honesty** — ``document_edit`` never fabricates success: a missing
   intent + no model is ``ok=False``, a non-``.docx`` path fails honestly, and an
   unsatisfiable constraint fails honestly after the bounded retry.
5. **Artifact projection** — a successful ``document.edit`` payload projects to
   the ``document_docx`` artifact (with ``content_ref`` + the diff), and a
   base64-in-payload design can never project (the payload ceiling eats it).
6. **Routing** — ``_is_document_task`` is code-first / mutually exclusive with
   analysis, and ``_candidates_for`` injects the three tools before the report.
"""

from __future__ import annotations

import base64
import hashlib
import io
from types import SimpleNamespace

import pytest

from forgeflow.documents import (
    EditOp,
    UnknownEditOpError,
    apply_edits,
    compute_diff,
    inspect_docx,
    resolve_intent,
    verify_docx,
)
from forgeflow.runtime import orchestrator as orch
from forgeflow.runtime.artifacts import (
    ARTIFACT_KIND_DOCX,
    artifacts_from_invocations,
)
from forgeflow.runtime.gate import PLATFORM_PLAN_TOOLS, TOOL_PERMISSION_MAP
from forgeflow.runtime.orchestrator import RequestContext, TaskCreate

pytestmark = pytest.mark.asyncio


# --------------------------------------------------------------------------- #
# Builders / helpers                                                           #
# --------------------------------------------------------------------------- #
def _docx_bytes(*, with_table: bool = True) -> bytes:
    """A known DOCX: two level-1 headings, body paragraphs with numbers, a table."""
    from docx import Document

    document = Document()
    document.add_heading("第一章 概述", level=1)
    document.add_paragraph("总金额 40 元。")
    document.add_paragraph("备注：请核对。")
    document.add_heading("第二章 明细", level=1)
    document.add_paragraph("单价 5 元。")
    if with_table:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _paragraph_texts(data: bytes) -> list[str]:
    from docx import Document

    return [p.text for p in Document(io.BytesIO(data)).paragraphs]


def _ctx() -> SimpleNamespace:
    return SimpleNamespace(
        run_id="run-inc43",
        step_id="run-inc43:0:0",
        tenant_id="t-inc43",
        user_id="u-inc43",
        role="manager",
        args={},
    )


@pytest.fixture
def tmp_doc_store(tmp_path, monkeypatch):
    """Point the handler's DOCX blob store at a tmp dir (keeps the suite hermetic)."""
    from forgeflow.documents import store as store_mod

    real = store_mod.DocArtifactStore

    def _factory(root=None):
        return real(root=root if root is not None else tmp_path / "docstore")

    import forgeflow.documents as docs_pkg

    monkeypatch.setattr(store_mod, "DocArtifactStore", _factory)
    monkeypatch.setattr(docs_pkg, "DocArtifactStore", _factory)
    return tmp_path


class _FakeModel:
    """A minimal chat model: returns a fixed JSON ``content`` (no ``reasoning``)."""

    def __init__(self, content: str) -> None:
        self._content = content
        self.calls = 0

    async def ainvoke(self, _messages):
        self.calls += 1
        return SimpleNamespace(content=self._content)


@pytest.fixture
def no_model(monkeypatch):
    """Force the LLM layer unavailable (so the handler's honest path is exercised)."""
    import forgeflow.models.provider as provider

    monkeypatch.setattr(provider, "get_model", lambda *a, **k: None)


# --------------------------------------------------------------------------- #
# 1. Layering                                                                  #
# --------------------------------------------------------------------------- #
async def test_apply_edits_really_changes_bytes_and_preserves_structure():
    old = _docx_bytes()
    structure = inspect_docx(old)

    new, changes = apply_edits(old, [EditOp(op="replace_text", match="40", replace="38")])

    # The bytes really changed …
    assert hashlib.sha256(new).hexdigest() != hashlib.sha256(old).hexdigest()
    assert changes == 1
    # … the new bytes reopen …
    after = inspect_docx(new)
    # … and the structure is preserved (heading levels + table count).
    assert [h["level"] for h in after.headings] == [h["level"] for h in structure.headings]
    assert [h["text"] for h in after.headings] == [h["text"] for h in structure.headings]
    assert after.tables == structure.tables == 1
    assert "38" in "\n".join(_paragraph_texts(new))


async def test_set_section_text_and_set_paragraph_write_the_body():
    old = _docx_bytes()
    new, changes = apply_edits(
        old,
        [EditOp(op="set_section_text", section="第一章 概述", text="新的正文甲\n新的正文乙")],
    )
    assert changes == 2
    texts = _paragraph_texts(new)
    assert "新的正文甲" in texts and "新的正文乙" in texts
    # The other section is untouched.
    assert "单价 5 元。" in texts

    new2, changes2 = apply_edits(
        old, [EditOp(op="set_paragraph", index=1, text="整段替换")]
    )
    assert changes2 == 1
    assert _paragraph_texts(new2)[1] == "整段替换"


async def test_unknown_edit_op_is_refused():
    with pytest.raises(UnknownEditOpError):
        EditOp(op="delete_document")
    with pytest.raises(UnknownEditOpError):
        EditOp.from_dict({"op": "nope"})
    with pytest.raises(UnknownEditOpError):
        # replace_text without a match is refused, never a silent no-op.
        EditOp.from_dict({"op": "replace_text", "replace": "x"})


async def test_resolve_intent_never_writes_bytes(monkeypatch):
    """The LLM layer returns ops and leaves the document bytes untouched."""
    data = _docx_bytes()
    structure = inspect_docx(data)
    fingerprint = hashlib.sha256(data).hexdigest()

    fake = _FakeModel('{"edits":[{"op":"replace_text","match":"40","replace":"38"}]}')
    monkeypatch.setattr(
        "forgeflow.models.provider.get_model", lambda *a, **k: fake
    )

    ops = await resolve_intent(data, "把总金额 40 改成 38", structure)

    assert ops == [EditOp(op="replace_text", match="40", replace="38")]
    assert fake.calls == 1
    # No write happened: the bytes and the structure are byte-for-byte unchanged.
    assert hashlib.sha256(data).hexdigest() == fingerprint
    assert _paragraph_texts(data) == _paragraph_texts(_docx_bytes())


async def test_resolve_intent_returns_none_when_model_unavailable(no_model):
    data = _docx_bytes()
    assert await resolve_intent(data, "改一下", inspect_docx(data)) is None


async def test_resolve_intent_returns_none_on_unusable_reply(monkeypatch):
    data = _docx_bytes()
    structure = inspect_docx(data)
    monkeypatch.setattr(
        "forgeflow.models.provider.get_model",
        lambda *a, **k: _FakeModel("这不是 JSON"),
    )
    assert await resolve_intent(data, "改一下", structure) is None

    monkeypatch.setattr(
        "forgeflow.models.provider.get_model",
        lambda *a, **k: _FakeModel('{"edits":[{"op":"delete_document"}]}'),
    )
    assert await resolve_intent(data, "改一下", structure) is None


# --------------------------------------------------------------------------- #
# 2. Diff is recomputable                                                      #
# --------------------------------------------------------------------------- #
async def test_compute_diff_exact_values_for_known_pair():
    old = _docx_bytes()
    new, _ = apply_edits(old, [EditOp(op="replace_text", match="40", replace="38")])

    report = compute_diff(old, new)
    assert report.to_dict() == {
        "modified": 1,
        "added": 0,
        "removed": 0,
        "numeric_changes": 1,
    }


async def test_compute_diff_counts_added_and_removed():
    from docx import Document

    document = Document()
    document.add_paragraph("第一行。")
    document.add_paragraph("第二行。")
    document.add_paragraph("第三行。")
    buffer = io.BytesIO()
    document.save(buffer)
    old = buffer.getvalue()

    # Drop the middle paragraph and append one → 1 removed + 1 added.
    document2 = Document(io.BytesIO(old))
    paragraphs = document2.paragraphs
    paragraphs[1]._p.getparent().remove(paragraphs[1]._p)
    document2.add_paragraph("第四行。")
    buffer2 = io.BytesIO()
    document2.save(buffer2)

    report = compute_diff(old, buffer2.getvalue())
    assert report.removed == 1
    assert report.added == 1
    assert report.modified == 0


async def test_compute_diff_numeric_only_change():
    old = _docx_bytes()
    # Same shape, only the number differs → modified == 1, numeric_changes == 1.
    new, _ = apply_edits(old, [EditOp(op="replace_text", match="5", replace="9")])
    report = compute_diff(old, new)
    assert report.modified == 1
    assert report.numeric_changes == 1
    assert report.added == 0 and report.removed == 0


async def test_compute_diff_is_empty_for_identical_documents():
    data = _docx_bytes()
    assert compute_diff(data, data).to_dict() == {
        "modified": 0,
        "added": 0,
        "removed": 0,
        "numeric_changes": 0,
    }


# --------------------------------------------------------------------------- #
# 3. Verification is honest (tri-state)                                        #
# --------------------------------------------------------------------------- #
async def test_verify_docx_requirement_ok_false_only_when_measured():
    data = _docx_bytes()
    over = verify_docx(data, {"max_chars": 1})
    assert over.openable is True
    assert over.requirement_ok is False
    assert over.structure_ok is None
    assert over.data_ok is None

    within = verify_docx(data, {"max_chars": 100000})
    assert within.requirement_ok is True


async def test_verify_docx_unmeasured_dimensions_are_none():
    report = verify_docx(_docx_bytes(), {})
    assert report.openable is True
    assert report.structure_ok is None
    assert report.data_ok is None
    assert report.requirement_ok is None


async def test_verify_docx_data_ok_flags_silent_number_change():
    old = _docx_bytes()
    new, _ = apply_edits(old, [EditOp(op="replace_text", match="40", replace="38")])

    # "40" vanished and the edit did not declare it → measured violation.
    silent = verify_docx(new, {"original_numbers": ["40"], "allowed_missing_numbers": []})
    assert silent.data_ok is False
    assert "40" in silent.notes

    # Same document, but the edit legitimately liberated "40" → consistent.
    declared = verify_docx(
        new, {"original_numbers": ["40"], "allowed_missing_numbers": ["40"]}
    )
    assert declared.data_ok is True


async def test_verify_docx_structure_preservation():
    data = _docx_bytes()
    structure = inspect_docx(data)
    ok = verify_docx(
        data, {"expected_sections": len(structure.sections), "expected_tables": structure.tables}
    )
    assert ok.structure_ok is True
    bad = verify_docx(data, {"expected_tables": 999})
    assert bad.structure_ok is False


async def test_verify_docx_unopenable_bytes_are_honest():
    report = verify_docx(b"not a docx", {"max_chars": 10})
    assert report.openable is False
    # Nothing else could be measured — never a False masquerading as "checked".
    assert report.structure_ok is None
    assert report.data_ok is None
    assert report.requirement_ok is None


# --------------------------------------------------------------------------- #
# 4. Handler honesty                                                           #
# --------------------------------------------------------------------------- #
async def test_document_edit_without_intent_and_without_model_fails_honestly(
    no_model, tmp_path
):
    from forgeflow.runtime.tool_handlers import document_edit

    source = tmp_path / "report.docx"
    source.write_bytes(_docx_bytes())

    result = await document_edit({"paths": [str(source)], "intent": "改一下这个东西"}, _ctx())

    assert result["ok"] is False
    assert result["unavailable"] is True
    assert "edits" in result["reason"]
    # No fabricated success fields.
    assert "modified" not in result
    assert "artifact_ref" not in result


async def test_document_edit_non_docx_path_fails_honestly(tmp_path):
    from forgeflow.runtime.tool_handlers import document_edit

    csv = tmp_path / "data.csv"
    csv.write_text("a,b\n1,2\n", encoding="utf-8")
    missing = tmp_path / "nope.docx"

    for paths in ([str(csv)], [str(missing)], []):
        result = await document_edit(
            {"paths": paths, "edits": [{"op": "replace_text", "match": "1", "replace": "2"}]},
            _ctx(),
        )
        assert result["ok"] is False
        assert result["not_executed"] is True
        assert "docx" in result["reason"]


async def test_document_edit_invalid_edit_op_fails_honestly(tmp_path):
    from forgeflow.runtime.tool_handlers import document_edit

    source = tmp_path / "report.docx"
    source.write_bytes(_docx_bytes())
    result = await document_edit(
        {"paths": [str(source)], "edits": [{"op": "delete_document"}]}, _ctx()
    )
    assert result["ok"] is False
    assert "非法" in result["reason"]
    assert "artifact_ref" not in result


async def test_document_edit_verification_exhaustion_fails_honestly(tmp_path, tmp_doc_store):
    from forgeflow.runtime.tool_handlers import document_edit

    source = tmp_path / "report.docx"
    source.write_bytes(_docx_bytes())

    # An explicit edit + an unsatisfiable constraint (≤1 字): retry is bounded and
    # then the handler reports the failure — it never pretends to have succeeded.
    result = await document_edit(
        {
            "paths": [str(source)],
            "edits": [{"op": "replace_text", "match": "40", "replace": "38"}],
            "max_chars": 1,
        },
        _ctx(),
    )
    assert result["ok"] is False
    assert "校验未通过" in result["reason"]
    assert result["modified"] == 0
    assert result["validation"]["requirement_ok"] is False
    assert "artifact_ref" not in result


async def test_document_edit_success_writes_bytes_and_returns_small_payload(tmp_path, tmp_doc_store):
    from forgeflow.runtime.tool_handlers import document_edit

    source = tmp_path / "report.docx"
    original = _docx_bytes()
    source.write_bytes(original)
    # An explicit numeric edit is legitimately allowed to change the number.
    result = await document_edit(
        {"paths": [str(source)], "edits": [{"op": "replace_text", "match": "40", "replace": "38"}]},
        _ctx(),
    )

    assert result["ok"] is True
    assert result["filename"] == "report.edited.docx"
    assert result["modified"] == 1
    assert result["numeric_changes"] == 1
    assert result["artifact_ref"].endswith(".docx")
    assert result["size_bytes"] > 0
    assert result["validation"]["structure_ok"] is True
    # The bytes are NOT in the payload (they live in the store) …
    assert "content" not in result
    # … and the very file at ``artifact_ref`` really is the edited document.
    from forgeflow.documents import DocArtifactStore

    blob = DocArtifactStore(root=tmp_doc_store / "docstore").get(result["artifact_ref"])
    assert hashlib.sha256(blob).hexdigest() != hashlib.sha256(original).hexdigest()
    assert "38" in "\n".join(_paragraph_texts(blob))


async def test_document_edit_with_explicit_edits_never_calls_the_model(
    tmp_path, tmp_doc_store, monkeypatch
):
    """The layer split in action: an explicit ``edits`` array bypasses the LLM."""
    from forgeflow.runtime.tool_handlers import document_edit

    class _Boom:
        async def ainvoke(self, *_a, **_k):  # pragma: no cover — must never run
            raise AssertionError("the LLM layer must not be called for explicit edits")

    monkeypatch.setattr("forgeflow.models.provider.get_model", lambda *a, **k: _Boom())

    source = tmp_path / "report.docx"
    source.write_bytes(_docx_bytes())
    result = await document_edit(
        {"paths": [str(source)], "edits": [{"op": "replace_text", "match": "40", "replace": "38"}]},
        _ctx(),
    )
    assert result["ok"] is True
    assert result["modified"] == 1


async def test_document_inspect_reports_real_structure(tmp_path):
    from forgeflow.runtime.tool_handlers import document_inspect

    source = tmp_path / "report.docx"
    source.write_bytes(_docx_bytes())

    ok = await document_inspect({"paths": [str(source)]}, _ctx())
    assert ok["ok"] is True
    assert ok["structure"]["tables"] == 1
    assert ok["structure"]["heading_count"] == 2
    assert ok["structure"]["paragraphs"] == 5

    bad = await document_inspect({"paths": [str(tmp_path / "x.txt")]}, _ctx())
    assert bad["ok"] is False
    assert bad["not_executed"] is True


async def test_artifact_save_registers_the_run_document(tmp_path):
    from forgeflow.runtime.tool_handlers import artifact_save

    trail = [
        {"tool": "document.inspect", "status": "ok", "payload": {"ok": True}},
        {
            "tool": "document.edit",
            "status": "ok",
            "payload": {"ok": True, "artifact_ref": "ab/abcd.docx", "filename": "r.edited.docx"},
        },
    ]
    saved = await artifact_save({"observations": trail}, _ctx())
    assert saved["ok"] is True
    assert saved["artifact_ref"] == "ab/abcd.docx"
    assert saved["kind"] == "document_docx"
    assert saved["filename"] == "r.edited.docx"

    # No document in the trail ⇒ honest failure (never an empty registration).
    empty = await artifact_save({"observations": []}, _ctx())
    assert empty["ok"] is False
    assert empty["not_executed"] is True


# --------------------------------------------------------------------------- #
# 5. Artifact projection                                                       #
# --------------------------------------------------------------------------- #
def _edit_success_invocation(ref: str = "ab/abcd1234.docx") -> dict:
    return {
        "tool": "document.edit",
        "status": "ok",
        "result_ref": "row/4210-131234",
        "started_at": "2026-10-02T00:00:00",
        "payload": {
            "ok": True,
            "filename": "report.edited.docx",
            "modified": 3,
            "added": 1,
            "removed": 2,
            "numeric_changes": 1,
            "changes": 6,
            "artifact_ref": ref,
            "size_bytes": 36857,
            "validation": {
                "openable": True,
                "structure_ok": True,
                "data_ok": True,
                "requirement_ok": None,
                "notes": "",
            },
        },
    }


async def test_document_edit_payload_projects_docx_artifact():
    artifacts = artifacts_from_invocations("run-x", [_edit_success_invocation()])

    assert len(artifacts) == 1
    artifact = artifacts[0]
    assert artifact["kind"] == ARTIFACT_KIND_DOCX == "document_docx"
    assert artifact["format"] == "docx"
    assert artifact["title"] == "report.edited.docx"
    assert artifact["content_ref"] == "ab/abcd1234.docx"
    assert artifact["content"] == ""  # the body is NOT in the payload
    assert artifact["diff"] == {
        "modified": 3,
        "added": 1,
        "removed": 2,
        "numeric_changes": 1,
    }
    assert artifact["source"] == "document.edit"


async def test_document_edit_and_artifact_save_emit_one_artifact():
    """The writer and the registrar carry the same ref — exactly one artifact."""
    trail = [
        _edit_success_invocation("ab/abcd1234.docx"),
        {
            "tool": "artifact.save",
            "status": "ok",
            "result_ref": "ab/abcd1234.docx",
            "started_at": "2026-10-02T00:00:01",
            "payload": {
                "ok": True,
                "artifact_ref": "ab/abcd1234.docx",
                "filename": "report.edited.docx",
                "kind": "document_docx",
            },
        },
    ]
    artifacts = artifacts_from_invocations("run-x", trail)
    assert [a["kind"] for a in artifacts] == ["document_docx"]


async def test_docx_artifact_without_ref_projects_nothing():
    inv = _edit_success_invocation()
    del inv["payload"]["artifact_ref"]
    assert artifacts_from_invocations("run-x", [inv]) == []


async def test_base64_payload_can_never_project_an_artifact():
    """Why the bytes must not ride in the payload (MAX_PAYLOAD_CHARS=4000).

    A naive "content = base64(DOCX)" design blows past the payload ceiling, so
    the executor replaces the whole payload with a ``truncated`` envelope: both
    the ``content`` and the ``artifact_ref`` keys vanish and the deliverable can
    never be projected (a permanent false empty state). The real design keeps the
    bytes in the store and the payload small — pinned by the contrast below.
    """
    from forgeflow.runtime.tool_executor import _bound_payload

    ref = "ab/abcd1234.docx"
    blob = base64.b64encode(_docx_bytes()).decode()
    assert len(blob) > 4000  # prerequisite: the naive payload really is oversized

    naive = {"ok": True, "filename": "report.docx", "artifact_ref": ref, "content": blob}
    bounded, truncated = _bound_payload(naive)
    assert truncated is True
    assert "content" not in bounded and "artifact_ref" not in bounded
    assert artifacts_from_invocations(
        "run-x", [{"tool": "document.edit", "status": "ok", "payload": bounded}]
    ) == []

    # The real (small) payload is NOT truncated and therefore DOES project.
    real = _edit_success_invocation(ref)["payload"]
    bounded_real, truncated_real = _bound_payload(real)
    assert truncated_real is False
    assert bounded_real["artifact_ref"] == ref
    assert len(artifacts_from_invocations("run-x", [{"tool": "document.edit", "status": "ok", "payload": bounded_real}])) == 1


# --------------------------------------------------------------------------- #
# 6. Routing pins                                                              #
# --------------------------------------------------------------------------- #
async def test_is_document_task_true_for_docx_false_for_csv(monkeypatch):
    ctx = RequestContext(tenant_id="t-doc", user_id="u", role="manager")

    by_tool = TaskCreate(intent="编辑文档", context={"declared_tools": ["document.edit"]})
    assert orch._is_document_task(by_tool, ctx) is True

    by_path = TaskCreate(intent="编辑文档", context={"paths": ["/srv/a/report.docx"]})
    assert orch._is_document_task(by_path, ctx) is True

    # A CSV task is not a document task (asserted before the resource stub below
    # so the stub cannot leak into it).
    csv = TaskCreate(intent="分析数据", context={"paths": ["/srv/a/data.csv"]})
    assert orch._is_document_task(csv, ctx) is False

    # Signal ② exactly as designed: a dereferenced resource path ending in .docx.
    monkeypatch.setattr(
        orch, "_resolve_resource_inputs", lambda _c: {"paths": ["/srv/a/report.docx"]}
    )
    resource_doc = TaskCreate(intent="编辑文档", context={"resources": ["r1"]})
    assert orch._is_document_task(resource_doc, ctx) is True


async def test_docx_task_is_not_an_analysis_task():
    ctx = RequestContext(tenant_id="t-doc", user_id="u", role="manager")
    task = TaskCreate(intent="编辑文档", context={"paths": ["/srv/a/report.docx"]})
    assert orch._is_analysis_task(task, ctx) is False
    # A CSV task is still an analysis task (no regression).
    csv = TaskCreate(intent="分析数据", context={"declared_tools": ["analysis.profile"]})
    assert orch._is_analysis_task(csv, ctx) is True


async def test_candidates_for_injects_document_plane_before_report():
    ctx = RequestContext(tenant_id="t-doc", user_id="u", role="manager")
    task = TaskCreate(intent="编辑文档", context={"paths": ["/srv/a/report.docx"]})
    tools = [c["tool"] for c in orch._candidates_for(task, ctx)]

    for tool in orch._DOCUMENT_TOOLS:
        assert tool in tools
    assert tools.index("document.inspect") < tools.index("document.edit")
    assert tools.index("document.edit") < tools.index("artifact.save")
    assert tools[-1] == "report.render"

    injected = orch._platform_injected_tools(task, ctx)
    assert injected == list(orch._DOCUMENT_TOOLS)
    # Injected steps count as declared, so an under-specified one stays ``blocked``
    # (visible) instead of being trimmed as not-applicable.
    declared = orch._declared_plan_tools(task, ctx)
    for tool in orch._DOCUMENT_TOOLS:
        assert tool in declared


async def test_document_tools_are_platform_tools_without_a_narrow_grant():
    for tool in ("document.inspect", "document.edit", "artifact.save"):
        assert tool in PLATFORM_PLAN_TOOLS
        assert tool not in TOOL_PERMISSION_MAP


async def test_blocked_document_edit_names_both_missing_inputs_in_order():
    """A declared document step missing its inputs blocks honestly and names them."""
    from forgeflow.runtime.planning import (
        CapabilityContext,
        applicability,
        blocked_reason,
        resolve_inputs,
    )

    cap = CapabilityContext(intent="编辑文档", declared_tools=["document.edit"])
    _args, missing = resolve_inputs("document.edit", cap)
    assert missing == ["paths", "edits"]
    assert applicability("document.edit", cap) == "blocked"
    reason = blocked_reason("document.edit", cap)
    assert "编辑意图(edits)" in reason  # the human label, not the raw key


# --------------------------------------------------------------------------- #
# 7. Store guard-rails                                                         #
# --------------------------------------------------------------------------- #
async def test_doc_store_roundtrip_and_refuses_in_project_root(tmp_path):
    from forgeflow.documents import DocArtifactStore, DocStorePathError

    store = DocArtifactStore(root=tmp_path / "blobs")
    data = _docx_bytes()
    ref = store.put(data)
    assert ref.endswith(".docx")
    assert store.get(ref) == data
    # Idempotent: same bytes ⇒ same content-addressed ref.
    assert store.put(data) == ref

    with pytest.raises(DocStorePathError):
        # A root inside the ForgeFlow checkout is refused (AC-11).
        import pathlib

        project = pathlib.Path(__file__).resolve().parents[2]
        DocArtifactStore(root=project)


# --------------------------------------------------------------------------- #
# 8. Full deterministic run (end to end)                                       #
# --------------------------------------------------------------------------- #
async def test_default_executor_document_task_produces_docx_artifact(tmp_path, tmp_doc_store):
    """A declared document task runs the plane and yields one downloadable DOCX."""
    from forgeflow.documents import DocArtifactStore

    source = tmp_path / "report.docx"
    source.write_bytes(_docx_bytes())

    class _Bus:
        def __init__(self):
            self.events = []

        async def emit(self, run_id, event_type, data):
            self.events.append((event_type, data))

    ctx = RequestContext(tenant_id="t-doc-e2e", user_id="u-1", role="manager")
    task = TaskCreate(
        intent="把文档里的总金额 40 元改成 38 元",
        context={
            "paths": [str(source)],
            "declared_tools": ["document.edit"],
            "edits": [{"op": "replace_text", "match": "40", "replace": "38"}],
        },
    )
    steps, errors = await orch._default_executor(task, ctx, _Bus(), "run-doc-e2e")

    assert errors == []
    by_tool = {s["tool"]: s["status"] for s in steps}
    assert by_tool["document.inspect"] == "ok"
    assert by_tool["document.edit"] == "ok"
    assert by_tool["artifact.save"] == "ok"

    artifacts = artifacts_from_invocations(
        "run-doc-e2e", list(task.context.get("tool_invocations") or []), fallback_created_at="t"
    )
    docx = [a for a in artifacts if a["kind"] == "document_docx"]
    assert len(docx) == 1
    assert docx[0]["diff"]["numeric_changes"] == 1
    blob = DocArtifactStore(root=tmp_doc_store / "docstore").get(docx[0]["content_ref"])
    assert "38" in "\n".join(_paragraph_texts(blob))
