"""The live acceptance oracle must reject untouched or damaged deliverables."""
import importlib.util
import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location(
    "complex_live_checks", Path(__file__).parents[2] / "scripts/verify_complex_tasks_live.py"
)
checks = importlib.util.module_from_spec(spec)
spec.loader.exec_module(checks)


@pytest.mark.parametrize("fixture", checks.fixtures())
def test_unmodified_input_is_not_task_completion(fixture):
    _, data, _, kind = fixture
    assert not all(value is True for value in checks.verify(kind, data).values())


def test_word_cleaning_does_not_hide_data_loss():
    _, data, _, _ = checks.fixtures()[1]
    doc = Document(io.BytesIO(data))
    doc.paragraphs[0].text = "项目验收"
    doc.tables[0].cell(0, 0).text = "项目"
    doc.tables[0].cell(1, 1).text = "12345"
    output = io.BytesIO()
    doc.save(output)
    result = checks.verify("word", output.getvalue())
    assert result["whitespace_cleaned"] is True
    assert result["amount_preserved"] is False
    from forgeflow.documents.whitespace import verify_office_cleanup
    assert verify_office_cleanup(data, output.getvalue(), checks.fixtures()[1][2], "docx") is False


def test_excel_cleaning_does_not_hide_formula_loss():
    _, data, _, _ = checks.fixtures()[2]
    book = load_workbook(io.BytesIO(data))
    book["销售"]["A2"] = "张三"
    book["销售"]["A3"] = "李四"
    book["销售"]["B4"] = 13235.67
    output = io.BytesIO()
    book.save(output)
    result = checks.verify("excel", output.getvalue())
    assert result["whitespace_cleaned"] is True
    assert result["formulas_preserved"] is False
    from forgeflow.documents.whitespace import verify_office_cleanup
    assert verify_office_cleanup(data, output.getvalue(), checks.fixtures()[2][2], "xlsx") is False



@pytest.mark.parametrize("index,kind", [(1, "docx"), (2, "xlsx")])
def test_office_cleanup_receipt_checks_actual_package(index, kind):
    from forgeflow.documents.whitespace import verify_office_cleanup

    _, original, intent, _ = checks.fixtures()[index]
    if kind == "docx":
        from forgeflow.documents.docx_edit import EditOp, apply_edits
        edited, _ = apply_edits(original, [EditOp(op="normalize_whitespace")])
    else:
        from forgeflow.documents.sheet_edit import SheetEditOp, apply_edits
        edited, _ = apply_edits(original, [SheetEditOp(op="normalize_whitespace")])
    assert verify_office_cleanup(original, edited, intent, kind) is True
    assert verify_office_cleanup(original, original, intent, kind) is False
    assert verify_office_cleanup(original, edited, intent + "并总结内容", kind) is None



async def test_workbook_request_cannot_be_narrowed_to_active_sheet(tmp_path):
    from types import SimpleNamespace

    from forgeflow.documents.store import DocArtifactStore
    from forgeflow.runtime.tool_handlers import sheet_edit

    _, original, intent, _ = checks.fixtures()[2]
    path = tmp_path / "compound.xlsx"
    path.write_bytes(original)
    result = await sheet_edit(
        {"sheet_paths": [str(path)], "edits": [{"op": "normalize_whitespace", "sheet": "销售"}]},
        SimpleNamespace(intent=intent, run_id="", tenant_id="qa-unit", role="manager"),
    )
    assert result["ok"] and result["validation"]["requirement_ok"] is True
    workbook = load_workbook(io.BytesIO(DocArtifactStore().get(result["artifact_ref"])))
    assert workbook["说明"]["A1"].value == "保留该工作表"
    assert workbook["销售"]["B4"].value == "=SUM(B2:B3)"
    assert path.read_bytes() == original
