"""INC44 T03 — the resource seam: PPTX + text / code FILEs (offline, no PG).

Nails the additive resource changes (design §1.2 / §1.3 / §7 T03):

  * ``.pptx`` joins the ``document`` kind (``SUPPORTED_FILE_EXTENSIONS``, the
    presentationml MIME, a real ``parsed`` summary with **measured** counts);
  * a registered ``.pptx`` dereferences to ``document_paths`` / ``document_names``
    (index-aligned) exactly like ``.docx`` — while ``paths`` is **item-for-item
    unchanged** (the code / analysis seams keep seeing every FILE path verbatim);
  * a registered ``.py`` (or any text / code FILE) dereferences to ``text_paths`` /
    ``text_names`` — likewise additive;
  * ``preview`` keeps the honest empty state for a document;
  * ``GET /resources/limits`` reads the one source of truth.
"""

from __future__ import annotations

import io
import os

import pytest

from forgeflow.repositories.memory.resource_repo import clear_resource_store
from forgeflow.resources import summaries
from forgeflow.resources.service import ResourceService, reset_resource_index

pytestmark = pytest.mark.asyncio


# --------------------------------------------------------------------------- #
# Builders / fixtures                                                          #
# --------------------------------------------------------------------------- #
def _pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    layout = presentation.slide_layouts[5]
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "总金额 40 元"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "备注 请核对"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def memory_resources(force_memory_backend):
    reset_resource_index()
    clear_resource_store()
    yield
    reset_resource_index()
    clear_resource_store()


@pytest.fixture
def blob_root(monkeypatch, tmp_path):
    from forgeflow.config import get_settings

    monkeypatch.setattr(get_settings(), "resource_store_root", str(tmp_path / "res-blobs"))
    return tmp_path / "res-blobs"


# --------------------------------------------------------------------------- #
# 1. Classification / MIME / limits                                            #
# --------------------------------------------------------------------------- #
async def test_pptx_is_a_supported_document_with_the_presentation_mime():
    assert summaries.content_kind("slides.pptx") == "document"
    assert summaries.content_kind("SLIDES.PPTX") == "document"  # case-insensitive
    assert summaries.is_supported_file("slides.pptx") is True
    assert ".pptx" in summaries.SUPPORTED_FILE_EXTENSIONS
    assert summaries.mime_for("slides.pptx") == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    # The sibling kinds are untouched.
    assert summaries.content_kind("book.xlsx") == "excel"
    assert summaries.content_kind("notes.md") == "text"
    assert summaries.content_kind("data.csv") == "table"


async def test_limits_reads_the_single_source_of_truth(memory_resources):
    extensions = ResourceService().limits()["supported_extensions"]
    assert extensions == list(summaries.SUPPORTED_FILE_EXTENSIONS)
    assert ".pptx" in extensions
    assert ".py" in extensions


# --------------------------------------------------------------------------- #
# 2. Summary is measured / degrades honestly                                   #
# --------------------------------------------------------------------------- #
async def test_summarize_bytes_parses_a_real_pptx_with_measured_counts():
    status, summary, detail = summaries.summarize_bytes(_pptx_bytes(), filename="slides.pptx")
    assert status == "parsed"
    assert detail == ""
    quality = summary.quality
    assert quality["slides"] == 1
    assert quality["text_frames"] >= 2
    assert summary.chars > 0


async def test_summarize_presentation_degrades_to_metadata_only_on_unreadable_bytes():
    for raw in (b"not a zip", b""):
        status, summary, detail = summaries.summarize_presentation(raw)
        assert status == "metadata_only"
        assert summary.note == "unparseable pptx"
        assert detail  # verbatim reason
        assert summary.chars is None  # unmeasured ⇒ None, never a fake 0


async def test_summarize_presentation_reports_unavailable_when_extra_missing(monkeypatch):
    """A missing ``python-pptx`` extra ⇒ metadata_only + verbatim (never 5xx)."""
    import forgeflow.documents.pptx_inspect as pptx_inspect

    def _boom(_data):
        raise ImportError("No module named 'pptx'")

    monkeypatch.setattr(pptx_inspect, "inspect_pptx", _boom)
    status, summary, detail = summaries.summarize_presentation(_pptx_bytes())
    assert status == "metadata_only"
    assert "pptx support unavailable" in summary.note
    assert detail


# --------------------------------------------------------------------------- #
# 3. Dereference is additive and index-aligned                                 #
# --------------------------------------------------------------------------- #
async def test_registered_pptx_dereferences_to_document_paths_additively(
    memory_resources, blob_root
):
    service = ResourceService()
    record = await service.register_file("t-inc44", name="slides.pptx", data=_pptx_bytes())
    assert record.status == "parsed"

    resolved = service.resolve_task_inputs({"resources": [record.id]})

    assert "document_paths" in resolved
    path = resolved["document_paths"][0]
    assert os.path.isfile(path)
    assert os.path.splitext(path)[1] == ""  # content-addressed ⇒ extensionless
    assert resolved["document_names"] == ["slides.pptx"]
    # Additive: ``paths`` still carries the file path verbatim.
    assert resolved["paths"] == [path]


async def test_registered_py_dereferences_to_text_paths_additively(memory_resources, blob_root):
    service = ResourceService()
    record = await service.register_file(
        "t-inc44", name="app.py", data=b"x = 40\nprint(x)\n"
    )
    assert record.status == "parsed"

    resolved = service.resolve_task_inputs({"resources": [record.id]})

    assert "text_paths" in resolved
    path = resolved["text_paths"][0]
    assert os.path.isfile(path)
    assert os.path.splitext(path)[1] == ""
    assert resolved["text_names"] == ["app.py"]
    # Additive: ``paths`` unchanged; a text FILE is NOT a document.
    assert resolved["paths"] == [path]
    assert "document_paths" not in resolved


async def test_all_dereferenced_lists_stay_item_for_item_aligned(
    memory_resources, blob_root
):
    """A mixed registration keeps every list index-aligned (same loop, same guard)."""
    service = ResourceService()
    csv = await service.register_file("t-inc44", name="data.csv", data=b"a,b\n1,2\n")
    py = await service.register_file("t-inc44", name="tool.py", data=b"x = 1\n")
    pptx = await service.register_file("t-inc44", name="deck.pptx", data=_pptx_bytes())

    resolved = service.resolve_task_inputs({"resources": [csv.id, py.id, pptx.id]})

    # ``paths`` carries every FILE path; the specialised lists are subsets.
    assert len(resolved["paths"]) == 3
    assert len(resolved["document_paths"]) == len(resolved["document_names"]) == 1
    assert len(resolved["text_paths"]) == len(resolved["text_names"]) == 1
    assert resolved["document_names"] == ["deck.pptx"]
    assert resolved["text_names"] == ["tool.py"]


async def test_resolve_task_inputs_is_empty_without_a_declaration(memory_resources):
    assert ResourceService().resolve_task_inputs({}) == {}


# --------------------------------------------------------------------------- #
# 4. Preview keeps the honest empty state                                      #
# --------------------------------------------------------------------------- #
async def test_preview_of_a_document_is_an_honest_empty_state(memory_resources, blob_root):
    service = ResourceService()
    record = await service.register_file("t-inc44", name="deck.pptx", data=_pptx_bytes())

    preview = await service.preview("t-inc44", record.id)

    assert preview["available"] is False
    assert preview["format"] == "document"  # the INC43 pin, unchanged
    assert preview["content"] == "" and preview["rows"] == []
    assert preview["note"]


async def test_preview_of_a_text_file_is_verbatim(memory_resources, blob_root):
    service = ResourceService()
    body = b"first line\nsecond line\n"
    record = await service.register_file("t-inc44", name="notes.txt", data=body)

    preview = await service.preview("t-inc44", record.id)

    assert preview["available"] is True
    assert preview["format"] == "text"
    # The preview joins the first ``n`` lines with LF (the existing text branch,
    # unchanged) — compare line-for-line, not byte-for-byte.
    assert preview["content"].splitlines() == body.decode("utf-8").splitlines()
