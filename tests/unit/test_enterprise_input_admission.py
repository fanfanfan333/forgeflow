"""Regression probes for upload delivery, restart admission and tenant isolation."""

from __future__ import annotations

import asyncio
import io
from types import SimpleNamespace
from unittest.mock import AsyncMock
from urllib.parse import quote

import pytest
from docx import Document
from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.testclient import TestClient

from forgeflow.api.dependencies import get_current_user
from forgeflow.api.hub_deps import resolve_tenant
from forgeflow.api.routers import resources, runs, tasks, workspace
from forgeflow.config import get_settings
from forgeflow.repositories.memory.resource_repo import clear_resource_store
from forgeflow.resources.service import ResourceService, reset_resource_index
from forgeflow.runtime.input_scope import (
    InputReferenceError,
    InputScope,
    current_input_scope,
    prepare_input_scope,
)
from forgeflow.runtime.orchestrator import (
    RequestContext,
    TaskCreate,
    get_run_store,
    register_running_run,
    reset_run_store,
)


@pytest.fixture(autouse=True)
def isolated(force_memory_backend, tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "resource_store_root", str(tmp_path / "inputs"))
    reset_resource_index()
    clear_resource_store()
    reset_run_store()
    yield
    reset_resource_index()
    clear_resource_store()
    reset_run_store()


def client_for(tenant="tenant-a"):
    app = FastAPI()
    app.include_router(resources.router, prefix="/resources")
    app.include_router(tasks.router, prefix="/tasks")
    app.include_router(workspace.router, prefix="/workspace")
    app.include_router(runs.router, prefix="/runs")
    app.dependency_overrides[resolve_tenant] = lambda: tenant
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
        user_id="user-a", role="admin"
    )
    return TestClient(app)


def docx_bytes():
    doc = Document()
    doc.add_paragraph("原文需要修复。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_uploaded_docx_is_edited_and_downloaded_after_cache_loss():
    original = docx_bytes()
    with client_for() as client:
        uploaded = client.post("/resources/files", files={"file": ("报告.docx", original)})
        assert uploaded.status_code == 201, uploaded.text
        rid = uploaded.json()["id"]
        reset_resource_index()  # Simulate submission to a fresh worker.
        response = client.post(
            "/tasks",
            json={
                "intent": "将原文替换为修复后的文字",
                "context": {
                    "resources": [rid],
                    "edits": [{"op": "replace_text", "match": "原文", "replace": "修复后的文字"}],
                },
            },
        )
        assert response.status_code == 200, response.text
        run_id = response.json()["run_id"]
        record = get_run_store().get(run_id)
        artifact = next(a for a in record.artifacts if a.get("format") == "docx")
        download = client.get(f"/runs/{run_id}/artifacts/{quote(artifact['id'], safe='')}")
        assert download.status_code == 200, download.text
        text = "\n".join(p.text for p in Document(io.BytesIO(download.content)).paragraphs)
        assert "修复后的文字需要修复。" in text
        assert download.content != original
        assert "原文" in Document(io.BytesIO(original)).paragraphs[0].text


@pytest.mark.asyncio
async def test_foreign_cached_resource_is_rejected_at_both_task_entries():
    record = await ResourceService().register_file("tenant-b", name="secret.txt", data=b"secret")
    with client_for() as client:
        for endpoint in ("/tasks", "/workspace/tasks"):
            response = client.post(
                endpoint, json={"intent": "read", "context": {"resources": [record.id]}}
            )
            assert response.status_code == 404, response.text
            assert response.json()["detail"] == "Resource not found"
    assert get_run_store().count() == 0


@pytest.mark.asyncio
async def test_empty_authoritative_snapshot_never_falls_back_to_cache():
    service = ResourceService()
    record = await service.register_file("tenant-b", name="secret.txt", data=b"secret")
    token = current_input_scope.set(InputScope("tenant-a", {}))
    try:
        assert service.resolve_task_inputs({"resources": [record.id]}) == {}
    finally:
        current_input_scope.reset(token)
    assert service.resolve_task_inputs({"resources": [record.id]}, records=[]) == {}


@pytest.mark.asyncio
async def test_parallel_scopes_do_not_leak_resources():
    service = ResourceService()
    a = await service.register_file("tenant-a", name="a.txt", data=b"a")
    b = await service.register_file("tenant-b", name="b.txt", data=b"b")

    async def resolve(record):
        token = current_input_scope.set(InputScope(record.tenant_id, {record.id: record}))
        try:
            await asyncio.sleep(0)
            return service.resolve_task_inputs({"resources": [a.id, b.id]})["text_names"]
        finally:
            current_input_scope.reset(token)

    assert await asyncio.gather(resolve(a), resolve(b)) == [["a.txt"], ["b.txt"]]
    assert current_input_scope.get() is None


@pytest.mark.asyncio
@pytest.mark.parametrize("key", ["parent_run_id", "continued_from_run_id"])
async def test_foreign_parent_is_rejected(key):
    register_running_run(
        run_id="foreign",
        thread_id="foreign-thread",
        tenant_id="tenant-b",
        intent="secret",
        workflow_type="generic",
    )
    with pytest.raises(InputReferenceError, match="Parent run not found"):
        await prepare_input_scope(
            TaskCreate(context={key: "foreign"}), RequestContext(tenant_id="tenant-a")
        )


@pytest.mark.asyncio
@pytest.mark.parametrize("key", ["parent_run_id", "continued_from_run_id"])
async def test_parent_from_another_user_in_same_tenant_is_rejected(key):
    register_running_run(run_id="private-run", thread_id="private-thread",
                         tenant_id="tenant-a", actor_user_id="alice",
                         intent="private", workflow_type="generic")
    with pytest.raises(InputReferenceError, match="Parent run not found"):
        await prepare_input_scope(TaskCreate(context={key: "private-run"}),
                                  RequestContext(tenant_id="tenant-a", user_id="bob"))


def test_other_user_cannot_read_or_stop_run_in_same_tenant():
    register_running_run(run_id="alice-run", thread_id="alice-thread",
                         tenant_id="tenant-a", actor_user_id="alice",
                         intent="private", workflow_type="generic")
    with client_for() as client:
        assert client.get("/runs/alice-run").status_code == 404
        assert client.post("/runs/alice-run/abort").status_code == 404
        listing = client.get("/runs").json()
        assert all(item["run_id"] != "alice-run" for item in listing["items"])


def test_request_key_lookup_cannot_find_another_users_run():
    from forgeflow.runtime.dispatcher import request_run_id
    key = "request-key-private"
    run_id = request_run_id("tenant-a", "alice", key)
    register_running_run(run_id=run_id, thread_id=run_id, tenant_id="tenant-a",
                         actor_user_id="alice", intent="private", workflow_type="generic")
    own_id = request_run_id("tenant-a", "user-a", "own-request-key")
    register_running_run(run_id=own_id, thread_id=own_id, tenant_id="tenant-a",
                         actor_user_id="user-a", intent="own", workflow_type="generic")
    with client_for() as client:
        assert client.get(f"/workspace/requests/{key}").status_code == 404
        assert client.get("/workspace/requests/own-request-key").json()["run_id"] == own_id


def test_owned_persisted_run_is_readable_when_not_in_history_cache(monkeypatch):
    from forgeflow.workspace.models import WorkspaceRunRecord
    from forgeflow.workspace.store import get_workspace_store

    asyncio.run(get_workspace_store().save(WorkspaceRunRecord(
        run_id="older-run", session_id="older-run", tenant_id="tenant-a",
        actor_user_id="user-a", intent="older")))
    monkeypatch.setattr(runs, "ensure_tenant_history", AsyncMock())
    with client_for() as client:
        assert client.get("/runs/older-run").status_code == 200


def test_other_user_cannot_read_or_delete_session_in_same_tenant():
    from forgeflow.workspace.models import WorkspaceRunRecord
    from forgeflow.workspace.store import get_workspace_store

    asyncio.run(get_workspace_store().save(WorkspaceRunRecord(
        run_id="alice-session", session_id="alice-session", tenant_id="tenant-a",
        actor_user_id="alice", intent="private")))
    with client_for() as client:
        assert client.get("/workspace/sessions/alice-session").status_code == 404
        assert client.delete("/workspace/sessions/alice-session").status_code == 404
        assert all(item["session_id"] != "alice-session"
                   for item in client.get("/workspace/sessions").json()["items"])


def test_deleting_own_session_does_not_evict_other_tenant_run():
    from forgeflow.workspace.models import WorkspaceRunRecord
    from forgeflow.workspace.store import get_workspace_store

    asyncio.run(get_workspace_store().save(WorkspaceRunRecord(
        run_id="own-run", session_id="shared-id", tenant_id="tenant-a",
        actor_user_id="user-a")))
    register_running_run(run_id="foreign-run", session_id="shared-id", thread_id="other",
                         tenant_id="tenant-b", actor_user_id="other", intent="private")
    with client_for() as client:
        assert client.delete("/workspace/sessions/shared-id").status_code == 200
    assert get_run_store().get("foreign-run") is not None


def test_other_user_cannot_approve_anothers_code_run():
    from forgeflow.api.routers import codeplane

    record = register_running_run(run_id="code-run", thread_id="thread",
        tenant_id="tenant-a", actor_user_id="alice", intent="private code")
    record.codeplane = {"workspace": {"workspace_id": "work"}}
    app = FastAPI()
    app.include_router(codeplane.router, prefix="/codeplane")
    app.dependency_overrides[resolve_tenant] = lambda: "tenant-a"
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(user_id="user-a", role="admin")
    with TestClient(app) as client:
        assert client.post("/codeplane/runs/code-run/approve", json={}).status_code == 404


def test_other_user_cannot_list_or_decide_private_code_approval():
    from forgeflow.api.routers import approvals_hub
    from forgeflow.codeplane.approval import request_code_approval
    from forgeflow.repositories import get_policy_repository

    register_running_run(run_id="private-code", thread_id="thread", tenant_id="tenant-a",
                         actor_user_id="alice", intent="private")
    approval = asyncio.run(request_code_approval(get_policy_repository(), tenant_id="tenant-a",
                           run_id="private-code", requester="alice"))
    app = FastAPI()
    app.include_router(approvals_hub.router, prefix="/approvals")
    app.dependency_overrides[resolve_tenant] = lambda: "tenant-a"
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(user_id="user-a", role="admin")
    with TestClient(app) as client:
        assert approval.id not in {item["id"] for item in client.get("/approvals").json()["items"]}
        assert client.post(f"/approvals/{approval.id}/decision", json={"decision": "approve"}).status_code == 404


def test_other_user_cannot_list_or_decide_private_run_approval():
    from forgeflow.api.routers import approvals_hub
    from forgeflow.governance.models import ApprovalRecord
    from forgeflow.repositories import get_policy_repository

    register_running_run(run_id="private-generic", thread_id="thread", tenant_id="tenant-a",
                         actor_user_id="alice", intent="private")
    approval = ApprovalRecord(tenant_id="tenant-a", run_id="private-generic",
                              requested_action="private", requester="alice")
    asyncio.run(get_policy_repository().save_approval(approval))
    app = FastAPI()
    app.include_router(approvals_hub.router, prefix="/approvals")
    app.dependency_overrides[resolve_tenant] = lambda: "tenant-a"
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(user_id="user-a", role="admin")
    with TestClient(app) as client:
        assert approval.id not in {item["id"] for item in client.get("/approvals").json()["items"]}
        assert client.post(f"/approvals/{approval.id}/decision", json={"decision": "approve"}).status_code == 404


def test_other_user_cannot_read_anothers_artifact_version(monkeypatch, tmp_path):
    from forgeflow.api.routers import artifact_review
    from forgeflow.documents.review_store import InMemoryArtifactReviewStore
    from forgeflow.documents.store import DocArtifactStore

    register_running_run(run_id="artifact-run", thread_id="thread", tenant_id="tenant-a",
                         actor_user_id="alice", intent="private")
    artifact_id = "artifact-private-test"
    store = InMemoryArtifactReviewStore(blob_store=DocArtifactStore(tmp_path / "artifacts"))
    monkeypatch.setattr(artifact_review, "get_artifact_review_store", lambda: store)
    version, _ = store.create_pending_version(
        "tenant-a", artifact_id=artifact_id, content=b"private", fmt="text", run_id="artifact-run")
    app = FastAPI()
    app.include_router(artifact_review.router, prefix="/artifacts")
    app.dependency_overrides[resolve_tenant] = lambda: "tenant-a"
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(user_id="user-a", role="admin")
    with TestClient(app) as client:
        assert client.get(f"/artifacts/{artifact_id}/versions/{version.version}/content").status_code == 404


def test_revert_keeps_original_run_ownership(monkeypatch, tmp_path):
    from forgeflow.api.routers import artifact_review
    from forgeflow.documents.review_store import InMemoryArtifactReviewStore
    from forgeflow.documents.store import DocArtifactStore

    register_running_run(run_id="owned-run", thread_id="thread", tenant_id="tenant-a",
                         actor_user_id="user-a", intent="task")
    store = InMemoryArtifactReviewStore(blob_store=DocArtifactStore(tmp_path / "artifacts"))
    monkeypatch.setattr(artifact_review, "get_artifact_review_store", lambda: store)
    store.create_pending_version("tenant-a", artifact_id="owned", content=b"data",
                                 fmt="text", run_id="owned-run")
    def fake_revert(tenant, **kwargs):
        assert kwargs["run_id"] == "owned-run"
        return store.create_pending_version(tenant, artifact_id=kwargs["artifact_id"],
                                            content=b"data", fmt="text", run_id=kwargs["run_id"])
    monkeypatch.setattr(artifact_review.vc, "revert", fake_revert)
    app = FastAPI()
    app.include_router(artifact_review.router, prefix="/artifacts")
    app.dependency_overrides[resolve_tenant] = lambda: "tenant-a"
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(user_id="user-a", role="admin")
    with TestClient(app) as client:
        response = client.post("/artifacts/owned/revert", json={"to_version": 1})
        assert response.status_code == 200, response.text
        assert client.get("/artifacts/owned/versions/2/content").status_code == 200


@pytest.mark.asyncio
async def test_task_cannot_join_another_users_session():
    from forgeflow.workspace.models import WorkspaceRunRecord
    from forgeflow.workspace.store import get_workspace_store

    await get_workspace_store().save(WorkspaceRunRecord(
        run_id="alice-run", session_id="alice-session", tenant_id="tenant-a",
        actor_user_id="alice"))
    with pytest.raises(InputReferenceError, match="Session not found"):
        await prepare_input_scope(TaskCreate(context={"session_id": "alice-session"}),
                                  RequestContext(tenant_id="tenant-a", user_id="bob"))


def test_experience_cannot_reference_another_users_run():
    from forgeflow.api.routers import experiences

    register_running_run(run_id="private-experience-run", thread_id="thread",
                         tenant_id="tenant-a", actor_user_id="alice", intent="secret")
    app = FastAPI()
    app.include_router(experiences.router, prefix="/experiences")
    app.dependency_overrides[resolve_tenant] = lambda: "tenant-a"
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(user_id="user-a", role="admin")
    with TestClient(app) as client:
        assert client.post("/experiences", json={"run_id": "private-experience-run"}).status_code == 404


@pytest.mark.asyncio
async def test_upload_reads_only_limit_plus_one_and_closes(monkeypatch):
    monkeypatch.setattr(get_settings(), "multimodal_max_bytes", 16)
    file = UploadFile(filename="large.txt", file=io.BytesIO(b"x" * 100_000))
    file.read = AsyncMock(wraps=file.read)
    file.close = AsyncMock(wraps=file.close)
    with pytest.raises(HTTPException) as exc:
        await resources.register_file(file, SimpleNamespace(user_id="user-a"), "tenant-a")
    assert exc.value.status_code == 413
    file.read.assert_awaited_once_with(17)
    file.close.assert_awaited_once()


@pytest.mark.asyncio
async def test_dispatch_capacity_is_atomic_under_concurrent_submission(monkeypatch):
    from forgeflow.runtime.dispatcher import RunCapacityError, RunDispatcher

    dispatcher = RunDispatcher(max_pending_runs=1)
    blocker = asyncio.Event()

    async def blocked_run(*args):
        await blocker.wait()

    monkeypatch.setattr(dispatcher, "_run_and_finalize", blocked_run)
    results = await asyncio.gather(
        dispatcher.dispatch(TaskCreate(intent="a"), RequestContext(tenant_id="tenant-a")),
        dispatcher.dispatch(TaskCreate(intent="b"), RequestContext(tenant_id="tenant-a")),
        return_exceptions=True,
    )
    assert sum(isinstance(result, RunCapacityError) for result in results) == 1
    assert get_run_store().count() == 1
    pending = list(dispatcher._tasks.values())
    blocker.set()
    await asyncio.gather(*pending)


@pytest.mark.asyncio
async def test_dispatch_does_not_accept_task_when_run_header_cannot_be_persisted(monkeypatch):
    from forgeflow.runtime.dispatcher import RunDispatcher
    from forgeflow.workspace.store import get_workspace_store

    store = get_workspace_store()
    async def fail_save(record):
        raise OSError("storage unavailable")
    monkeypatch.setattr(store, "save", fail_save)
    dispatcher = RunDispatcher(max_pending_runs=1)
    with pytest.raises(Exception, match="storage unavailable"):
        await dispatcher.dispatch(TaskCreate(intent="test"),
                                  RequestContext(tenant_id="tenant-a", user_id="user-a"))
    assert get_run_store().count("tenant-a") == 0
    assert not dispatcher._tasks


@pytest.mark.asyncio
async def test_persisted_parent_is_loaded_by_id_after_restart():
    from forgeflow.workspace.models import WorkspaceRunRecord
    from forgeflow.workspace.store import get_workspace_store

    resource = await ResourceService().register_file("tenant-a", name="parent.txt", data=b"parent")
    await get_workspace_store().save(
        WorkspaceRunRecord(
            run_id="old-parent",
            tenant_id="tenant-a",
            session_id="old-parent",
            declared_inputs={"resources": [resource.id]},
        )
    )
    reset_run_store()
    reset_resource_index()
    scope = await prepare_input_scope(
        TaskCreate(context={"continued_from_run_id": "old-parent"}),
        RequestContext(tenant_id="tenant-a"),
    )
    assert resource.id in scope.resources
    assert get_run_store().get("old-parent").tenant_id == "tenant-a"


@pytest.mark.asyncio
async def test_multi_turn_history_inherits_latest_file_and_answers_after_restart():
    from forgeflow.documents.store import DocArtifactStore
    from forgeflow.runtime.orchestrator import (
        _resolve_continued_context,
        _resolve_inherited_resources,
    )
    from forgeflow.workspace.models import WorkspaceRunRecord
    from forgeflow.workspace.store import get_workspace_store

    resource = await ResourceService().register_file(
        "tenant-a", name="original.txt", data=b"original"
    )
    ref = DocArtifactStore().put(b"edited value 120", ".txt")
    store = get_workspace_store()
    await store.save(
        WorkspaceRunRecord(
            run_id="turn-one",
            tenant_id="tenant-a",
            session_id="chain",
            intent="remember invoice 123",
            declared_inputs={"resources": [resource.id]},
            artifacts=[
                {"id": "file/slash", "format": "text", "content_ref": ref, "title": "edited.txt"}
            ],
            llm={"final_answer": {"text": "Changed to 120"}},
        )
    )
    await store.save(
        WorkspaceRunRecord(
            run_id="turn-two",
            tenant_id="tenant-a",
            session_id="chain",
            parent_run_id="turn-one",
            intent="another question",
            llm={"final_answer": {"text": "Second reply"}},
        )
    )
    reset_run_store()
    reset_resource_index()
    task = TaskCreate(context={"continued_from_run_id": "turn-two"})
    scope = await prepare_input_scope(task, RequestContext(tenant_id="tenant-a"))
    token = current_input_scope.set(scope)
    try:
        history = _resolve_continued_context(task.context)
        assert "remember invoice 123" in history
        assert "Changed to 120" in history
        assert "Second reply" in history
        inputs, _, _ = _resolve_inherited_resources(task.context)
        assert inputs["text_names"] == ["edited.txt"]
        from pathlib import Path

        assert Path(inputs["text_paths"][0]).read_bytes() == b"edited value 120"
    finally:
        current_input_scope.reset(token)


def test_preview_reads_file_blob_and_rejects_foreign_run():
    from forgeflow.documents.store import DocArtifactStore

    ref = DocArtifactStore().put("修改后的真实内容：120万元".encode(), ".txt")
    record = register_running_run(
        run_id="preview",
        thread_id="thread",
        tenant_id="tenant-a",
        actor_user_id="user-a",
        intent="edit",
        workflow_type="generic",
    )
    record.artifacts = [
        {
            "id": "artifact/slash",
            "format": "text",
            "content": "",
            "content_ref": ref,
            "title": "edited.txt",
        }
    ]
    with client_for() as client:
        response = client.get(
            "/runs/preview/artifact-preview", params={"artifact_id": "artifact/slash"}
        )
        assert response.status_code == 200, response.text
        assert response.json()["text"] == "修改后的真实内容：120万元"
    with client_for("tenant-b") as client:
        assert (
            client.get(
                "/runs/preview/artifact-preview", params={"artifact_id": "artifact/slash"}
            ).status_code
            == 404
        )


def test_workspace_capacity_error_returns_retryable_503(monkeypatch):
    from forgeflow.runtime.dispatcher import RunCapacityError

    monkeypatch.setattr(
        workspace,
        "get_run_dispatcher",
        lambda: SimpleNamespace(
            dispatch=AsyncMock(side_effect=RunCapacityError("任务队列已满，请稍后重试"))
        ),
    )
    with client_for() as client:
        response = client.post("/workspace/tasks", json={"intent": "检查"})
    assert response.status_code == 503
    assert response.headers["Retry-After"] == "5"
