"""INC46 T24 — 修复循环与失败路径（Repair & Escalation）钉子测试。

What this file nails down (task book §DoD, 逐条)
-----------------------------------------------
1. **共享状态常量** —— ``failed_validation`` 是**加性**新增（不在既有
   ``TERMINAL_STATUSES`` 里），且与 T16 labeler 的本地常量**逐字相等**。
2. **可修复类（真编辑器 + 真验证栈）**：
   * 格式丢失（恢复 run 样式）⇒ ≤3 轮内修复，终态 ``pass``；
   * 区间外改动（回滚）⇒ 修复；
   * 不变量被改（回填）⇒ 修复。
3. **不可修复** —— 结构损坏（不可读）⇒ **不空转** 3 轮：``rounds_used == 0``、
   ``failed_validation``、报告含 **verbatim** 证据、无 committed 版本。
4. **终止条件** —— 轮数 ≤ 3（预算）；失败项数**单调不增**（递减用例为
   ``[2, 1, 0]``）；「修复使失败项增加」⇒ **立即终止**（振荡用例 ``rounds_used == 1``）。
5. **失败报告** —— 含 verbatim 失败信息、失败层、尝试次数、建议用户操作；
   run 状态 = ``failed_validation``。
6. **诚实** —— 未测量层（``status=None``）不得折算 pass；验证失败的文档**绝不**
   作为成功返回（:class:`RepairOutcome` 不携带任何文档字节）。
7. **反事实（真跑）** —— 去掉轮数上限 ⇒ 卡死用例在外层超时保护下转红；去掉
   「单调不增」判定 ⇒ 振荡用例跑到上限、断言转红。
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from forgeflow.documents.repair import (
    MAX_ROUNDS,
    RepairAttempt,
    default_editor,
    diagnose,
    failure_count,
    is_pass,
    repair_document,
)
from forgeflow.documents.status import FAILED_VALIDATION
from forgeflow.documents.validation import validate_document
from forgeflow.documents.validation.verdict import (
    FAIL,
    LAYER_FORMAT,
    LAYER_INVARIANTS,
    LAYER_ORDER,
    LAYER_RENDER,
    LAYER_SEMANTIC,
    LAYER_STRUCTURAL,
    NEEDS_REVIEW,
    PASS,
    LayerVerdict,
    ValidationVerdict,
)

_MARKER = "[[TGT::repair]]"
_REPLACEMENT = "已编辑::repair"


# --------------------------------------------------------------------------- #
# builders (real DOCX)                                                         #
# --------------------------------------------------------------------------- #
def _save(document: Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sample_doc(*, bold_target: bool = False) -> bytes:
    """Heading + a target paragraph (amount + date) + a table (mirrors T23)."""
    document = Document()
    document.add_heading("报告标题", level=1)
    document.add_paragraph("文档编号：UNIT；用途：修复循环。")
    target = document.add_paragraph(
        f"占位 {_MARKER}，金额 1,000.00 元，日期 2026-10-03，只应改这一段。"
    )
    if bold_target:
        target.runs[0].bold = True
    document.add_heading("正文小节", level=2)
    document.add_paragraph("正文内容，不应被改动。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    return _save(document)


def _target_range(data: bytes) -> tuple[int, int]:
    from forgeflow.documents.docx_inspect import open_docx

    texts = [p.text for p in open_docx(data).paragraphs]
    index = next(i for i, text in enumerate(texts) if _MARKER in text)
    return index, index + 1


def _compliant_edit(data: bytes) -> bytes:
    from forgeflow.documents.docx_edit import EditOp, apply_edits

    edited, changes = apply_edits(
        data, [EditOp(op="replace_text", match=_MARKER, replace=_REPLACEMENT)]
    )
    assert changes == 1, changes
    return edited


def _read_parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _repack(parts: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in sorted(parts):
            archive.writestr(name, parts[name])
    return buffer.getvalue()


def _mutate_document_xml(data: bytes, old: bytes, new: bytes) -> bytes:
    parts = _read_parts(data)
    blob = parts["word/document.xml"]
    assert old in blob, old
    parts["word/document.xml"] = blob.replace(old, new)
    return _repack(parts)


def _format_lost_after(before: bytes) -> bytes:
    """A compliant in-range edit that **loses the leading run's bold** (L3 fail)."""
    document = Document(io.BytesIO(before))
    index, _ = _target_range(before)
    run = document.paragraphs[index].runs[0]
    rpr = run._r.find(qn("w:rPr"))
    if rpr is not None:
        bold = rpr.find(qn("w:b"))
        if bold is not None:
            rpr.remove(bold)
    run.text = run.text.replace(_MARKER, _REPLACEMENT)
    return _save(document)


def _out_of_range_bolden_after(before: bytes) -> bytes:
    """A compliant in-range edit that **also reformats paragraph 0** (L3 fail)."""
    edited = _compliant_edit(before)
    document = Document(io.BytesIO(edited))
    document.paragraphs[0].runs[0].bold = True
    return _save(document)


def _amount_dropped_after(before: bytes) -> bytes:
    """An in-range edit that **drops the protected amount** (L2 fail)."""
    return _mutate_document_xml(before, b"1,000.00", b"")


# --------------------------------------------------------------------------- #
# synthetic verdicts (for injected editor/verifier counterfactuals)             #
# --------------------------------------------------------------------------- #
def _verdict(lost_count: int) -> ValidationVerdict:
    """A verdict whose only non-pass layer is L2, with ``lost_count`` lost items."""
    overall = PASS if lost_count == 0 else FAIL
    layers: dict[str, LayerVerdict] = {}
    for name in LAYER_ORDER:
        if name == LAYER_INVARIANTS:
            if lost_count == 0:
                layers[name] = LayerVerdict(name, PASS, "invariants:ok", {})
            else:
                layers[name] = LayerVerdict(
                    name,
                    FAIL,
                    f"invariants:FAIL[n={lost_count}]",
                    {"lost": [f"amount={i}元@paragraph[{i}]" for i in range(lost_count)]},
                )
        elif name in (LAYER_STRUCTURAL, LAYER_FORMAT):
            layers[name] = LayerVerdict(name, PASS, f"{name}:ok", {})
        else:
            layers[name] = LayerVerdict(name, None, f"{name}:not measured", {})
    return ValidationVerdict(
        layers=layers,
        overall=overall,
        unmeasured=[LAYER_RENDER, LAYER_SEMANTIC],
        notes="synthetic",
    )


def _needs_review_verdict() -> ValidationVerdict:
    layers: dict[str, LayerVerdict] = {}
    for name in LAYER_ORDER:
        if name == LAYER_SEMANTIC:
            layers[name] = LayerVerdict(name, NEEDS_REVIEW, "semantic:low", {"advisory_only": True})
        else:
            layers[name] = LayerVerdict(name, PASS, f"{name}:ok", {})
    return ValidationVerdict(layers=layers, overall=NEEDS_REVIEW, unmeasured=[], notes="advisory")


# --------------------------------------------------------------------------- #
# 1. the shared status constant                                                #
# --------------------------------------------------------------------------- #
def test_failed_validation_constant_is_additive_and_matches_the_labeler():
    from forgeflow.outcomes.labeler import FAILED_VALIDATION_STATUS
    from forgeflow.workspace.store import TERMINAL_STATUSES

    assert FAILED_VALIDATION == "failed_validation"
    # 加性：既有终态枚举一字不改，failed_validation 是追加的独立取值。
    assert FAILED_VALIDATION not in TERMINAL_STATUSES
    assert TERMINAL_STATUSES == frozenset(
        {"completed", "failed", "aborted", "interrupted", "rejected"}
    )
    # 与 T16 labeler 的本地同值常量逐字一致（唯一共享来源）。
    assert FAILED_VALIDATION_STATUS == FAILED_VALIDATION


# --------------------------------------------------------------------------- #
# 2. failure-count + success predicate                                          #
# --------------------------------------------------------------------------- #
def test_failure_count_sums_measured_items_across_failing_layers():
    assert failure_count(_verdict(0)) == 0
    assert failure_count(_verdict(3)) == 3
    l3 = ValidationVerdict(
        layers={
            LAYER_STRUCTURAL: LayerVerdict(LAYER_STRUCTURAL, PASS, "ok", {}),
            LAYER_INVARIANTS: LayerVerdict(LAYER_INVARIANTS, PASS, "ok", {}),
            LAYER_FORMAT: LayerVerdict(
                LAYER_FORMAT,
                FAIL,
                "format:FAIL",
                {"out_of_range_changes": ["paragraph[0]"], "style_inheritance_failures": ["paragraph[2]"]},
            ),
            LAYER_RENDER: LayerVerdict(LAYER_RENDER, None, "n/a", {}),
            LAYER_SEMANTIC: LayerVerdict(LAYER_SEMANTIC, None, "n/a", {}),
        },
        overall=FAIL,
        unmeasured=[LAYER_RENDER, LAYER_SEMANTIC],
    )
    assert failure_count(l3) == 2  # 1 out-of-range + 1 style-lost


def test_is_pass_follows_overall_and_never_upgrades_advisory_or_unmeasured():
    assert is_pass(_verdict(0)) is True
    assert is_pass(_verdict(1)) is False
    # needs_review（advisory）绝不算 pass。
    assert is_pass(_needs_review_verdict()) is False
    # 未测量层（status=None）不是 pass，也不制造成功声明。
    verdict = _verdict(0)
    assert verdict.layers[LAYER_RENDER].status is None
    assert verdict.layers[LAYER_RENDER].status != PASS


# --------------------------------------------------------------------------- #
# 3. 阳性探针 —— 真编辑器 + 真验证栈                                             #
# --------------------------------------------------------------------------- #
def test_positive_format_loss_is_repaired_within_budget_real_stack():
    before = _sample_doc(bold_target=True)
    after = _format_lost_after(before)
    start, end = _target_range(before)

    initial = validate_document(before, after, start=start, end=end)
    assert initial.layer(LAYER_FORMAT).status == FAIL  # the defect really exists

    rows: list[dict] = []
    outcome = repair_document(
        before, after, initial, max_rounds=MAX_ROUNDS, trace_sink=rows.append
    )

    assert outcome.status == PASS
    assert outcome.failed_layer is None
    assert 1 <= outcome.rounds_used <= MAX_ROUNDS
    assert outcome.attempts and all(isinstance(a, RepairAttempt) for a in outcome.attempts)
    # 每轮落一条 run_steps
    assert len(rows) == outcome.rounds_used
    assert rows[0]["step_type"] == "repair"
    assert rows[0]["failure_count"] == 0  # repaired ⇒ no measured failures
    # 失败项数单调不增
    counts = [a.failure_count for a in outcome.attempts]
    assert all(b <= a for a, b in zip(counts, counts[1:]))


def test_positive_out_of_range_change_is_rolled_back_real_stack():
    before = _sample_doc()
    after = _out_of_range_bolden_after(before)
    start, end = _target_range(before)

    initial = validate_document(before, after, start=start, end=end)
    assert initial.layer(LAYER_FORMAT).status == FAIL
    assert "paragraph[0]" in initial.layer(LAYER_FORMAT).detail["out_of_range_changes"]

    outcome = repair_document(before, after, initial)
    assert outcome.status == PASS
    assert outcome.rounds_used == 1
    assert outcome.attempts[0].actions == ["rollback_paragraph(paragraph[0])"]


def test_positive_lost_invariant_is_backfilled_real_stack():
    before = _sample_doc()
    after = _amount_dropped_after(before)
    start, end = _target_range(before)

    initial = validate_document(before, after, start=start, end=end)
    assert initial.layer(LAYER_INVARIANTS).status == FAIL
    assert initial.layer(LAYER_INVARIANTS).detail["lost"]  # a real amount was lost

    outcome = repair_document(before, after, initial)
    assert outcome.status == PASS
    assert outcome.rounds_used == 1
    assert any("backfill_paragraph" in action for action in outcome.attempts[0].actions)


# --------------------------------------------------------------------------- #
# 4. 阴性探针 —— 不可修复 ⇒ 诚实失败，绝不成功返回                              #
# --------------------------------------------------------------------------- #
def test_negative_structural_damage_fails_immediately_with_verbatim_report():
    before = _sample_doc()
    after = b"this is not a docx at all"
    start, end = _target_range(before)
    initial = validate_document(before, after, start=start, end=end)
    assert initial.layer(LAYER_STRUCTURAL).status == FAIL

    rows: list[dict] = []
    outcome = repair_document(before, after, initial, trace_sink=rows.append)

    assert outcome.status == FAILED_VALIDATION
    assert outcome.failed_layer == LAYER_STRUCTURAL
    # 不空转 3 轮：不可修复 ⇒ 立即终止。
    assert outcome.rounds_used == 0
    assert outcome.attempts == []
    assert rows == []  # no round ran ⇒ no run_steps
    # 报告含 verbatim 证据（L1 的 evidence_ref 逐字出现）
    assert initial.layer(LAYER_STRUCTURAL).evidence_ref in outcome.failure_report
    assert "建议用户操作" in outcome.failure_report
    # 无 committed 版本；RepairOutcome 不携带任何文档字节。
    assert "已提交版本：无" in outcome.failure_report
    for attribute in ("repaired_bytes", "committed_bytes", "artifact", "bytes", "content"):
        assert not hasattr(outcome, attribute)


def test_negative_failed_document_is_never_returned_as_success():
    before = _sample_doc()
    after = _amount_dropped_after(before)
    start, end = _target_range(before)
    initial = validate_document(before, after, start=start, end=end)

    # 一个每次修复都「没有进展」的编辑器 + 恒定失败的验证器。
    def stuck_editor(before_bytes, after_bytes, diagnosis):
        return after_bytes

    def stuck_verifier(before_bytes, after_bytes):
        return initial

    outcome = repair_document(
        before, after, initial, editor=stuck_editor, verifier=stuck_verifier, max_rounds=3
    )
    assert outcome.status == FAILED_VALIDATION
    assert outcome.status != PASS
    assert outcome.rounds_used == MAX_ROUNDS  # bounded by the budget
    assert outcome.failure_report is not None


def test_negative_advisory_needs_review_is_not_upgraded_to_success():
    outcome = repair_document(b"B", b"A", _needs_review_verdict())
    assert outcome.status == FAILED_VALIDATION
    assert outcome.rounds_used == 0
    assert outcome.attempts == []
    assert "已提交版本：无" in outcome.failure_report


def test_default_editor_returns_none_when_after_is_unreadable():
    assert default_editor(_sample_doc(), b"not a docx", diagnose(_verdict(1))) is None


# --------------------------------------------------------------------------- #
# 5. 诊断                                                                       #
# --------------------------------------------------------------------------- #
def test_diagnose_marks_structural_failure_unrepairable():
    verdict = ValidationVerdict(
        layers={
            LAYER_STRUCTURAL: LayerVerdict(LAYER_STRUCTURAL, FAIL, "structural:OPEN_FAILED", {}),
            LAYER_INVARIANTS: LayerVerdict(LAYER_INVARIANTS, FAIL, "invariants:ERROR", {}),
            LAYER_FORMAT: LayerVerdict(LAYER_FORMAT, FAIL, "format:ERROR", {}),
            LAYER_RENDER: LayerVerdict(LAYER_RENDER, None, "n/a", {}),
            LAYER_SEMANTIC: LayerVerdict(LAYER_SEMANTIC, None, "n/a", {}),
        },
        overall=FAIL,
        unmeasured=[LAYER_RENDER, LAYER_SEMANTIC],
    )
    diagnosis = diagnose(verdict)
    assert diagnosis.repairable is False
    assert diagnosis.layer == LAYER_STRUCTURAL
    assert diagnosis.actions == []


# --------------------------------------------------------------------------- #
# 6. 终止条件（注入 editor/verifier）                                            #
# --------------------------------------------------------------------------- #
def test_decreasing_failure_sequence_is_monotonic_then_passes():
    sequence = [3, 2, 1, 0]
    state = {"round": 0}

    def editor(before_bytes, after_bytes, diagnosis):
        state["round"] += 1
        return after_bytes + b"#"

    def verifier(before_bytes, after_bytes):
        return _verdict(sequence[state["round"]])

    outcome = repair_document(b"B", b"A", _verdict(3), editor=editor, verifier=verifier, max_rounds=3)
    assert outcome.status == PASS
    assert [a.failure_count for a in outcome.attempts] == [2, 1, 0]  # 单调不增
    assert outcome.rounds_used == 3


def test_oscillation_makes_the_loop_stop_early():
    state = {"round": 0}

    def growing_editor(before_bytes, after_bytes, diagnosis):
        state["round"] += 1
        return after_bytes + b"#"

    def growing_verifier(before_bytes, after_bytes):
        return _verdict(1 + 2 * state["round"])  # 3, 5, 7 → 每轮都更糟

    outcome = repair_document(b"B", b"A", _verdict(1), editor=growing_editor, verifier=growing_verifier, max_rounds=3)
    assert outcome.status == FAILED_VALIDATION
    assert outcome.rounds_used == 1  # 一发现「变糟」立即终止
    assert outcome.attempts[-1].failure_count == 3
    assert outcome.failure_report is not None


# --------------------------------------------------------------------------- #
# 7. 反事实（真跑）                                                              #
# --------------------------------------------------------------------------- #
def test_counterfactual_removing_the_round_bound_makes_the_stuck_loop_unbounded(monkeypatch):
    import forgeflow.documents.repair as repair_mod

    calls = {"n": 0}

    def stuck_editor(before_bytes, after_bytes, diagnosis):
        calls["n"] += 1
        if calls["n"] > 50:  # 外层超时保护（circuit breaker）
            raise TimeoutError("外层超时保护：修复未在轮数预算内终止")
        return after_bytes  # 恒定无进展

    def stuck_verifier(before_bytes, after_bytes):
        return _verdict(1)  # 恒定失败

    # green：有轮数上限 ⇒ 在预算内终止。
    bounded = repair_document(
        b"B", b"A", _verdict(1), editor=stuck_editor, verifier=stuck_verifier, max_rounds=3
    )
    assert bounded.rounds_used == 3
    assert calls["n"] == 3

    # 反事实：去掉轮数上限 ⇒ 卡死循环永不终止，外层超时保护触发（转红）。
    monkeypatch.setattr(repair_mod, "_effective_max_rounds", lambda _m: 10**9)
    with pytest.raises(TimeoutError):
        repair_document(
            b"B", b"A", _verdict(1), editor=stuck_editor, verifier=stuck_verifier, max_rounds=3
        )


def test_counterfactual_removing_the_monotonic_guard_makes_oscillation_run_to_the_cap(
    monkeypatch,
):
    import forgeflow.documents.repair as repair_mod

    state = {"round": 0}

    def growing_editor(before_bytes, after_bytes, diagnosis):
        state["round"] += 1
        return after_bytes + b"#"

    def growing_verifier(before_bytes, after_bytes):
        return _verdict(1 + 2 * state["round"])

    def _assert_early_stop(outcome):
        assert outcome.rounds_used == 1, (
            f"振荡未提前终止（rounds_used={outcome.rounds_used}）"
        )

    _assert_early_stop(
        repair_document(b"B", b"A", _verdict(1), editor=growing_editor, verifier=growing_verifier, max_rounds=3)
    )  # green

    # 反事实：摘掉「单调不增」判定 ⇒ 振荡用例跑到上限，断言转红。
    state["round"] = 0
    monkeypatch.setattr(repair_mod, "_oscillation_violated", lambda _prev, _cur: False)
    with pytest.raises(AssertionError) as excinfo:
        _assert_early_stop(
            repair_document(b"B", b"A", _verdict(1), editor=growing_editor, verifier=growing_verifier, max_rounds=3)
        )
    assert "振荡未提前终止" in str(excinfo.value)


# --------------------------------------------------------------------------- #
# 8. 预算边界                                                                   #
# --------------------------------------------------------------------------- #
def test_max_rounds_zero_never_repairs():
    outcome = repair_document(b"B", b"A", _verdict(1), max_rounds=0)
    assert outcome.status == FAILED_VALIDATION
    assert outcome.rounds_used == 0
    assert outcome.attempts == []


def test_negative_max_rounds_is_rejected():
    with pytest.raises(ValueError):
        repair_document(b"B", b"A", _verdict(1), max_rounds=-1)
