"""INC46 T19 — 种子 Skill 与冷启动（seed skills + cold start）.

判据（任务书 T19 阳性/阴性/反事实探针）:

* **阳性（承重）** — 空经验库租户 + 文档编辑任务 ⇒ T09 Discovery 命中种子 ⇒ 经**真实编排路径**
  (:func:`forgeflow.runtime.orchestrator._resolve_injected_skills` →
  :func:`forgeflow.skills.runtime.load_procedure` →
  :func:`~forgeflow.skills.runtime.to_plan_candidates`) 逐步交给真实
  :class:`~forgeflow.runtime.tool_executor.ToolExecutor` 真执行 ⇒ 真实
  ``run_id → skill_id → skill_version`` 链 + T01 ``run_steps`` 落盘。种子是**可被编排器自动驱动**
  的，不是「建议文本」。
* **阴性** — ① 无该 Skill 权限的角色**不可见**（T09「权限先于一切」）；② evolution 尝试**原地覆盖**
  种子被拒（复用既有 ``versioning.create_version`` 的「只新增 semver」契约，不另造闸门）；
  ③ 种子缺**任一**七段 ⇒ 显式装载失败（禁止静默补全）。
* **反事实（真跑）** — 移除种子装载 ⇒ 阳性用例转红（Discovery 无命中）。
* **同构** — 每个种子走同一套七段契约（T07）+ 同一物化路径（T11）+ T18 校验；
  ``origin=seed`` 用**既有列**表达（无迁移）。

引文纪律：一律 ``file.py::symbol``，不写行号。
"""

from __future__ import annotations

import json
import uuid
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from forgeflow.repositories import get_skill_repository
from forgeflow.repositories.memory.skill_repo import clear_skill_store
from forgeflow.repositories.skill_schema_repo import SkillPermission
from forgeflow.runtime import orchestrator as orch
from forgeflow.runtime import tool_registry
from forgeflow.runtime.tool_executor import ToolCallContext, ToolExecutor
from forgeflow.runtime.trace_store import TraceStep, reset_trace_sink, set_trace_sink
from forgeflow.skills.models import SkillRecord
from forgeflow.skills.registry import SkillRegistry
from forgeflow.skills.retrieval import retrieve_skills
from forgeflow.skills.runtime import (
    SkillRuntime,
    load_procedure,
    to_plan_candidates,
)
from forgeflow.skills.schemas import validate_contract_document
from forgeflow.skills.seeds import loader as seed_loader
from forgeflow.skills.seeds.loader import (
    ORIGIN_AUTHORED,
    ORIGIN_SEED,
    SEED_ORIGIN_TAG,
    SEED_OWNER,
    SEED_VERSION,
    SeedError,
    apply_evolution,
    install_seeds,
    is_seed,
    load_seed_dir,
    load_seeds,
    origin_of,
    spec_sha256,
)
from forgeflow.skills.segments import REQUIRED_SEGMENTS, SEGMENT_ORDER
from forgeflow.skills.spec_validator import validate_skill_md

TENANT_A = "tenant-cold-start-a"
TENANT_B = "tenant-cold-start-b"
#: A document-editing intent — the seeds' whole reason to exist.
_QUERY = "改写 Word 文档章节并保持格式"
#: The four slugs the task book names, verbatim.
TASK_BOOK_SLUGS = frozenset(
    {
        "docx-section-rewrite",
        "docx-exact-replace",
        "pptx-slide-text-edit",
        "xlsx-cell-edit",
    }
)


# --------------------------------------------------------------------------- #
# fixtures / helpers                                                           #
# --------------------------------------------------------------------------- #
class _RecordingSink:
    """A minimal in-process :class:`TraceStore` (the T01 test seam)."""

    def __init__(self) -> None:
        self.steps: list[TraceStep] = []

    async def insert_step(self, step: TraceStep) -> None:
        self.steps.append(step)

    async def list_for_run(
        self, tenant_id: str | None, run_id: str, *, limit: int = 500
    ) -> list[TraceStep]:
        return [s for s in self.steps if s.run_id == run_id and s.tenant_id == tenant_id][:limit]

    async def list_for_tenant(
        self, tenant_id: str | None, *, limit: int = 1000
    ) -> list[TraceStep]:
        return [s for s in self.steps if s.tenant_id == tenant_id][:limit]


@pytest.fixture(autouse=True)
def _clean_state(force_memory_backend):
    """Every test starts from default tool bindings, no sink, empty skill store."""
    tool_registry.reset_registry()
    tool_registry.load_default_bindings()
    reset_trace_sink()
    clear_skill_store()
    yield
    tool_registry.reset_registry()
    tool_registry.load_default_bindings()
    reset_trace_sink()
    clear_skill_store()


def _valid_spec(slug: str, *, domain: str = "文档编辑") -> dict:
    """A fresh, spec-compliant seven-segment contract for ``slug``."""
    return {
        "manifest": {
            "name": slug,
            "display_name": "示例种子",
            "version": "1.0.0",
            "description": "示例种子：用于测试的合法七段契约。",
            "metadata": {"domain": domain},
        },
        "knowledge": {"references": [{"file": "guide.md", "summary": "指南"}]},
        "procedure": {"steps": ["第一步", "第二步"]},
        "policies": {"constraints": ["约束一"], "allowed_tools": []},
        "tool_bindings": {"tools": ["document.inspect"], "scripts": []},
        "evaluation": {"ref": "assets/evals/", "note": "DB"},
        "examples": {"examples": ["示例一"]},
    }


def _write_seed(root: Path, slug: str, spec: dict, procedure: dict | None = None) -> Path:
    directory = root / slug
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "spec.json").write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    if procedure is not None:
        (directory / "procedure.json").write_text(
            json.dumps(procedure, ensure_ascii=False), encoding="utf-8"
        )
    return directory


def _make_docx(path: str) -> None:
    doc = Document()
    doc.add_heading("付款条款", level=1)
    doc.add_paragraph("甲方应在 2024-01-01 前支付首款。")
    doc.save(path)


def _make_pptx(path: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "旧标题"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "旧正文"
    prs.save(path)


def _make_xlsx(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "标题"
    ws["B2"] = 1
    wb.save(path)


def _seed_args(tool: str, tmp: str, prior: dict) -> dict:
    """Per-tool args so each real handler actually runs and returns ``ok``."""
    docx = str(Path(tmp) / "a.docx")
    pptx = str(Path(tmp) / "a.pptx")
    xlsx = str(Path(tmp) / "a.xlsx")
    if tool == "document.inspect":
        return {"document_paths": [docx, pptx]}
    if tool == "document.edit":
        if prior.get("_format") == "pptx" or prior.get("format") == "pptx":
            edits = [{"op": "replace_text", "match": "旧", "text": "新"}]
            return {"document_paths": [pptx], "edits": edits, "intent": _QUERY}
        edits = [{"op": "replace_text", "match": "甲方", "text": "采购方"}]
        return {"document_paths": [docx], "edits": edits, "intent": _QUERY}
    if tool == "sheet.inspect":
        return {"sheet_paths": [xlsx]}
    if tool == "sheet.edit":
        return {
            "sheet_paths": [xlsx],
            "edits": [{"op": "set_cell", "cell": "B2", "value": 42}],
            "intent": _QUERY,
        }
    if tool == "artifact.save":
        return {"artifact_ref": prior.get("artifact_ref") or "run://t19/artifact"}
    return {}


async def _installed(tenant: str):
    repo = get_skill_repository()
    return await install_seeds(tenant, repo=repo), repo


# --------------------------------------------------------------------------- #
# 1. 阳性 —— 四枚任务书种子 + T07 契约 + T11 物化 + T18 校验                      #
# --------------------------------------------------------------------------- #
def test_catalogue_is_exactly_the_four_task_book_seeds() -> None:
    slugs = {seed.slug for seed in load_seeds()}
    assert slugs == TASK_BOOK_SLUGS, f"种子集合必须是任务书四枚：缺 {TASK_BOOK_SLUGS - slugs}，多 {slugs - TASK_BOOK_SLUGS}"


def test_every_contract_passes_t07_and_materialises_a_t18_valid_skill_md() -> None:
    for seed in load_seeds():
        report = validate_contract_document(seed.spec, parent_dir=seed.slug)
        assert report.ok, f"{seed.slug} 的契约必须通过 T07：{report.errors}"

        assert "SKILL.md" in seed.bundle.files, f"{seed.slug} 必须物化出 SKILL.md"
        assert seed.bundle.validation["ok"] is True
        skmd = seed.bundle.files["SKILL.md"]
        assert skmd.startswith("---")
        skill_md_report = validate_skill_md(skmd, parent_dir=seed.slug)
        assert skill_md_report.ok, f"{seed.slug} 的 SKILL.md 必须通过 T18：{skill_md_report.errors}"

        # The shipped materialised product must equal the freshly rendered one.
        shipped = (seed.path / seed_loader.SKILL_MD_FILENAME).read_text(encoding="utf-8")
        assert shipped == skmd, f"{seed.slug} 随附的 SKILL.md 必须等于契约的物化结果"


def test_seeds_are_isomorphic_seven_segment() -> None:
    """种子与自研 Skill 同构：七段齐全，不是简化格式。"""
    for seed in load_seeds():
        present = set(seed.spec)
        for segment in REQUIRED_SEGMENTS:
            assert segment in present, f"{seed.slug} 缺少段 {segment}"
        assert tuple(s for s in SEGMENT_ORDER if s in present) == SEGMENT_ORDER


async def test_origin_seed_is_expressed_with_existing_columns() -> None:
    installed, _repo = await _installed(TENANT_A)
    seed_skill = installed[0].skill
    assert seed_skill.owner == SEED_OWNER
    assert SEED_ORIGIN_TAG in seed_skill.tags
    assert is_seed(seed_skill) is True
    assert origin_of(seed_skill) == ORIGIN_SEED

    authored = SkillRecord(tenant_id=TENANT_A, name="自研技能", owner="alice", tags=["x"])
    assert is_seed(authored) is False
    assert origin_of(authored) == ORIGIN_AUTHORED


# --------------------------------------------------------------------------- #
# 2. 阳性 —— 空经验库冷启动命中（T09 Discovery）                                 #
# --------------------------------------------------------------------------- #
async def test_cold_start_empty_tenant_discovers_a_seed() -> None:
    installed, repo = await _installed(TENANT_A)
    assert installed, "冷启动必须安装种子"
    # Empty experience library by construction — nothing was learned yet.
    _skills, total = await repo.list_skills(TENANT_A, limit=100)
    assert total == len(installed)

    registry = SkillRegistry(repo=repo)
    hits = await registry.retrieve(TENANT_A, _QUERY, k=3)
    assert hits, "空经验库首次检索必须命中种子 —— 空结果会让阳性断言变成空集永真"
    assert hits[0].skill.name == "改写 Word 指定章节"
    assert hits[0].lexical_score is not None and hits[0].lexical_score > 0.0


# --------------------------------------------------------------------------- #
# 3. 阳性（承重）—— 经真实编排路径可驱动 + 真执行 + 真实 ID 链                    #
# --------------------------------------------------------------------------- #
async def test_seed_is_really_orchestratable_through_the_stored_spec() -> None:
    """承重：DB 里存的 ``version.spec["procedure"]`` 必须是**运行时形状**（list）。

    这是「种子可被编排器自动驱动」与「退化成建议文本」的分水岭：走
    :func:`forgeflow.runtime.orchestrator._resolve_injected_skills`（真实路径）从**仓储**解析，
    而不是加载器的内存投影。
    """
    installed, repo = await _installed(TENANT_A)

    registry = SkillRegistry(repo=repo)
    hits = await registry.retrieve(TENANT_A, _QUERY, k=3)
    assert hits

    resolved = await orch._resolve_injected_skills(TENANT_A, [h.skill.id for h in hits])
    by_id = {entry["id"]: entry for entry in resolved}
    target = next(item for item in installed if item.slug == "docx-section-rewrite")
    entry = by_id[target.skill.id]

    assert isinstance(entry.get("procedure"), list) and entry["procedure"], (
        "退化为建议文本：存储的 spec 未被解析出结构化 procedure"
    )
    assert entry["steps"] == list(target.seed.steps)

    procedure = load_procedure(entry)
    assert [s.tool for s in procedure] == ["document.inspect", "document.edit", "artifact.save"]
    assert SkillRuntime().degrade_reason(procedure, role="manager") is None
    assert [c["tool"] for c in to_plan_candidates(entry, procedure)] == [
        "document.inspect",
        "document.edit",
        "artifact.save",
    ]


async def test_seed_procedure_really_executes_and_persists_run_steps(tmp_path: Path) -> None:
    """阳性承重 —— 命中种子后逐句经真实 ToolExecutor 执行，T01 落 ``run_steps``。"""
    sink = _RecordingSink()
    set_trace_sink(sink)

    _make_docx(str(tmp_path / "a.docx"))
    _make_pptx(str(tmp_path / "a.pptx"))
    _make_xlsx(str(tmp_path / "a.xlsx"))

    installed, repo = await _installed(TENANT_A)
    registry = SkillRegistry(repo=repo)
    hits = await registry.retrieve(TENANT_A, _QUERY, k=3)
    assert hits, "必须先命中种子才能谈执行"

    run_ids: list[str] = []
    for item in installed:
        # ① 真实解析路径：仓储 → _resolve_injected_skills
        resolved = await orch._resolve_injected_skills(TENANT_A, [item.skill.id])
        assert len(resolved) == 1
        skill = resolved[0]
        procedure = load_procedure(skill)
        plan = to_plan_candidates(skill, procedure)

        # ② 每步都交给真实 ToolExecutor（handler 真跑）
        run_id = str(uuid.uuid4())
        run_ids.append(run_id)
        prior: dict = {"_format": "pptx" if "pptx" in item.slug else "docx"}
        for index, candidate in enumerate(plan):
            inv = await ToolExecutor().execute(
                candidate["tool"],
                ctx=ToolCallContext(
                    run_id=run_id,
                    step_id=f"{run_id}:0:{index}",
                    tenant_id=TENANT_A,
                    user_id="u-cold-start",
                    role="manager",
                    intent=_QUERY,
                    args=_seed_args(candidate["tool"], str(tmp_path), prior),
                ),
            )
            assert inv.status == "ok", f"{item.slug}/{inv.tool} 未真执行成功：{inv.error}"
            assert inv.executed is True and inv.invoked is True
            assert inv.development_stub is False, f"{inv.tool} 不得是开发态 stub"
            assert inv.provider != "development-stub"
            assert inv.latency_ms is not None and inv.latency_ms > 0  # 真测量
            result = getattr(inv, "result", None)
            if isinstance(result, dict):
                prior.update({k: result.get(k) for k in ("artifact_ref", "format")})
                if result.get("format"):
                    prior["_format"] = result["format"]

        # ③ 真实 run_steps（T01 落盘，tenant 作用域）
        rows = [s for s in sink.steps if s.run_id == run_id]
        assert len(rows) == len(plan)
        assert {r.tool for r in rows} == set(item.seed.procedure_tools())
        assert all(r.status == "ok" for r in rows)
        assert all(r.tenant_id == TENANT_A for r in rows)

    # ④ 真实 ID 链：run_id → skill_id → skill_version（全部来自真实记录）
    assert len(run_ids) == len(installed)
    for item, run_id in zip(installed, run_ids):
        assert any(s.run_id == run_id and s.tenant_id == TENANT_A for s in sink.steps)
        stored = await repo.get_version(TENANT_A, item.skill.id, SEED_VERSION)
        assert stored is not None, "skill_id → skill_version 必须能由仓储回读"
        assert item.skill.current_version == SEED_VERSION


# --------------------------------------------------------------------------- #
# 4. 阴性 —— 无该 Skill 权限的角色不可见（T09 权限先于一切）                      #
# --------------------------------------------------------------------------- #
async def test_negative_unentitled_permission_cannot_see_the_seed() -> None:
    installed, repo = await _installed(TENANT_A)
    skills, _total = await repo.list_skills(TENANT_A, limit=100)

    # 阳性对照：无任何权限行 ⇒ 未声明限制，命中。
    open_hits = retrieve_skills(TENANT_A, _QUERY, skills, permissions={}, k=5)
    assert open_hits, "无权限行时应当命中（阳性对照）"

    # 阴性：为每个命中 skill 加一条 granted=False ⇒ 候选池阶段即剔除。
    deny = {
        str(hit.skill.id): [
            SkillPermission(skill_id=str(hit.skill.id), tool="run", granted=False)
        ]
        for hit in open_hits
    }
    blocked = retrieve_skills(TENANT_A, _QUERY, skills, permissions=deny, k=5)
    assert blocked == [], "无权限身份的元数据不得进入候选池（权限先于一切）"


# --------------------------------------------------------------------------- #
# 5. 阴性 —— evolution 尝试原地覆盖种子被拒（复用 create_version 只新增 semver）  #
# --------------------------------------------------------------------------- #
async def test_negative_evolution_never_overwrites_the_seed_in_place() -> None:
    """原地覆盖是不被允许的：演化只**新增** semver，incumbent 本体 sha256 恒定。"""
    installed, repo = await _installed(TENANT_A)
    target = next(item for item in installed if item.slug == "docx-exact-replace")
    incumbent_semver = target.version.semver
    before = spec_sha256(target.version.spec)
    taken = {v.semver for v in await repo.list_versions(TENANT_A, target.skill.id)}
    assert incumbent_semver in taken

    evolved = target.seed.runtime_spec()
    new_version = await apply_evolution(target, evolved, repo=repo)

    # 只新增 semver：绝不复用已占用的 semver（即不可能原地覆盖）。
    assert new_version.semver not in taken
    assert new_version.semver != incumbent_semver
    # incumbent 本体一字未改。
    stored = await repo.get_version(TENANT_A, target.skill.id, incumbent_semver)
    assert spec_sha256(stored.spec) == before


# --------------------------------------------------------------------------- #
# 6. 阴性 —— 缺任一七段 ⇒ 显式装载失败（禁止静默补全）                            #
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("segment", SEGMENT_ORDER)
def test_negative_missing_any_segment_fails_to_load(tmp_path: Path, segment: str) -> None:
    spec = _valid_spec("demo-store")
    del spec[segment]
    _write_seed(tmp_path, "demo-store", spec)
    with pytest.raises(SeedError) as exc:
        load_seed_dir(tmp_path / "demo-store")
    assert segment in str(exc.value), f"缺段 {segment} 的报错必须点名该段"


def test_negative_illegal_slug_and_off_catalogue_tool_fail_loudly(tmp_path: Path) -> None:
    _write_seed(tmp_path, "Bad Slug", _valid_spec("Bad Slug"))
    with pytest.raises(SeedError) as exc:
        load_seed_dir(tmp_path / "Bad Slug")
    assert "Bad Slug" in str(exc.value)

    bad_proc = {"domain": "general", "procedure": [{"purpose": "越权", "tool": "evil.tool"}]}
    _write_seed(tmp_path, "demo-store", _valid_spec("demo-store"), bad_proc)
    with pytest.raises(SeedError) as exc2:
        load_seed_dir(tmp_path / "demo-store")
    assert "evil.tool" in str(exc2.value) and "白名单" in str(exc2.value)


def test_a_seed_without_procedure_degrades_honestly(tmp_path: Path) -> None:
    from forgeflow.skills.runtime import UNDECLARED_PROCEDURE_REASON

    _write_seed(tmp_path, "demo-store", _valid_spec("demo-store"))
    seed = load_seed_dir(tmp_path / "demo-store")
    assert seed.procedure == ()
    assert SkillRuntime().degrade_reason([], role="manager") == UNDECLARED_PROCEDURE_REASON


# --------------------------------------------------------------------------- #
# 7. 反事实（真跑）—— 移除种子装载 ⇒ 阳性转红                                     #
# --------------------------------------------------------------------------- #
async def test_counterfactual_removing_seed_loading_turns_the_positive_red(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """把「种子装载」关掉 ⇒ Discovery 无命中 ⇒ 阳性断言必须转红。"""

    async def _no_install(*args, **kwargs):  # noqa: ANN002, ANN003
        return []

    monkeypatch.setattr(seed_loader, "install_seeds", _no_install)

    installed = await seed_loader.install_seeds(TENANT_A, repo=get_skill_repository())
    assert installed == [], "反事实前提：卸载装载后没有任何种子"

    skills, _total = await get_skill_repository().list_skills(TENANT_A, limit=100)
    hits = retrieve_skills(TENANT_A, _QUERY, skills, k=3)
    # 这正是阳性用例 ``assert hits`` 的那一行 —— 它现在会失败。
    assert hits == [], "移除种子装载后，Discovery 必须无命中（阳性转红）"


async def test_install_seeds_is_idempotent_and_tenant_scoped() -> None:
    installed, repo = await _installed(TENANT_A)
    again = await install_seeds(TENANT_A, repo=repo)
    assert [i.skill.id for i in again] == [i.skill.id for i in installed]
    _skills, total = await repo.list_skills(TENANT_A, limit=100)
    assert total == len(installed), "重复安装不得产生重复种子"

    # 跨租户隔离（红线 5）
    skills_b, total_b = await repo.list_skills(TENANT_B, limit=100)
    assert skills_b == [] and total_b == 0
    registry = SkillRegistry(repo=repo)
    assert await registry.retrieve(TENANT_B, _QUERY, k=3) == []
