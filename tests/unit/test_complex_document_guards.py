import io
from types import SimpleNamespace

import pytest
from docx import Document

from forgeflow.documents import DocArtifactStore
from forgeflow.documents.content_read import read_content
from forgeflow.documents.requested_edits import requested_replacements
from forgeflow.runtime.tool_handlers import document_edit, document_inspect


def fixture_bytes():
    doc = Document()
    doc.add_paragraph("ForgeFlow INC46 工程主控任务簿 v3", "Title")
    doc.add_paragraph("其他正文仍提到工程主控任务簿")
    doc.add_paragraph("实现了函数不等于完成")
    doc.add_paragraph("文末声明：不是「读起来合理」")
    doc.add_heading("任务依赖表", 1)
    table = doc.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["T28", "T12 T21 T29", "L6", "M3"], strict=True):
        cell.text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


async def test_original_three_replacements_override_destructive_model_title(tmp_path):
    data = fixture_bytes()
    source = tmp_path / "tasks.docx"
    source.write_bytes(data)
    result = await document_edit(
        {
            "document_paths": [str(source)],
            "_user_intent": "1. 将标题中的「工程主控任务簿」替换为「工程验收任务簿」；"
            "2. 将「实现了函数」替换为「实现了单个函数」；"
            "3. 将「不是「读起来合理」」替换为「不是「读起来合理」或仅有口头说明」。"
            "保留其余内容和数字不变。",
            "edits": [{"op": "set_slide_title", "text": "工程验收任务簿"}],
        },
        SimpleNamespace(run_id="complex-test", tenant_id="test", role="manager"),
    )
    assert result["ok"] and result["validation"]["requirement_ok"]
    new = Document(io.BytesIO(DocArtifactStore().get(result["artifact_ref"])))
    assert new.paragraphs[0].text == "ForgeFlow INC46 工程验收任务簿 v3"
    assert new.paragraphs[1].text == "其他正文仍提到工程主控任务簿"
    assert new.paragraphs[2].text == "实现了单个函数不等于完成"
    assert new.paragraphs[3].text == "文末声明：不是「读起来合理」或仅有口头说明"
    assert source.read_bytes() == data


async def test_model_cannot_relax_original_number_constraint(tmp_path):
    source = tmp_path / "tasks.docx"
    source.write_bytes(fixture_bytes())
    result = await document_edit(
        {
            "document_paths": [str(source)],
            "_user_intent": "改写标题，保留所有数字不变",
            "edits": [{"op": "set_paragraph", "index": 0, "text": "工程验收任务簿"}],
        },
        SimpleNamespace(run_id="bad-number", tenant_id="test", role="manager"),
    )
    assert result["ok"] is False
    assert "artifact_ref" not in result


def test_literal_replacement_missing_target_is_not_success():
    with pytest.raises(ValueError, match="未找到"):
        requested_replacements(fixture_bytes(), "将「不存在」替换为「新文字」")


def test_real_table_rows_have_locations_and_support_paging():
    page = read_content(fixture_bytes(), "T28 前置 层 里程碑", limit=1)
    assert "table:0/row:0" in page["text"]
    assert "T12 T21 T29" in page["text"]
    assert page["next_offset"] == 1
    later = read_content(fixture_bytes(), "T28 前置 层 里程碑", offset=1, limit=1)
    assert later["text"] != page["text"]


async def test_inspection_content_survives_payload_ceiling(tmp_path):
    from forgeflow.runtime.tool_executor import _bound_payload

    source = tmp_path / "tasks.docx"
    source.write_bytes(fixture_bytes())
    result = await document_inspect(
        {"document_paths": [str(source)]}, SimpleNamespace(intent="T28 的前置任务是什么")
    )
    bounded, truncated = _bound_payload(result)
    assert not truncated
    assert "T12 T21 T29" in bounded["content"]
    assert "UNTRUSTED_TOOL_OUTPUT" in bounded["content"]


async def test_model_short_query_cannot_omit_original_questions(tmp_path):
    doc = Document(io.BytesIO(fixture_bytes()))
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "每阶段最小样本 / 时长"
    table.rows[0].cells[1].text = "至少 30 个标记样本且至少 24h"
    source = tmp_path / "tasks.docx"
    doc.save(source)
    result = await document_inspect(
        {"document_paths": [str(source)], "query": "T28"},
        SimpleNamespace(intent="T28 前置是什么？每阶段最小样本和时长是多少？"),
    )
    assert "30" in result["content"] and "24h" in result["content"]


async def test_writer_noop_cannot_claim_literal_request_completed(tmp_path, monkeypatch):
    import forgeflow.documents

    data = fixture_bytes()
    source = tmp_path / "tasks.docx"
    source.write_bytes(data)
    monkeypatch.setattr(forgeflow.documents, "apply_edits", lambda data, ops: (data, 0))
    result = await document_edit(
        {
            "document_paths": [str(source)],
            "_user_intent": "将标题中的「工程主控任务簿」替换为「工程验收任务簿」，保留其余内容。",
        },
        SimpleNamespace(run_id="noop-writer", tenant_id="test", role="manager"),
    )
    assert result["ok"] is False
    assert result["validation"]["requirement_ok"] is False
    assert "artifact_ref" not in result


def test_scoped_title_replacement_preserves_surrounding_runs():
    from forgeflow.documents.docx_edit import apply_edits

    doc = Document()
    title = doc.add_paragraph(style="Title")
    title.add_run("INC46 ").italic = True
    title.add_run("工程主控任务簿").bold = True
    title.add_run(" v3").italic = True
    doc.add_paragraph("工程主控任务簿正文")
    buffer = io.BytesIO()
    doc.save(buffer)
    original = buffer.getvalue()
    ops = requested_replacements(original, "将标题中的「工程主控任务簿」替换为「工程验收任务簿」")
    edited, _ = apply_edits(original, ops)
    result = Document(io.BytesIO(edited))
    assert result.paragraphs[1].text == "工程主控任务簿正文"
    assert [(r.text, r.bold, r.italic) for r in result.paragraphs[0].runs] == [
        ("INC46 ", None, True),
        ("工程验收任务簿", True, None),
        (" v3", None, True),
    ]


def test_long_paragraph_resumes_without_skipping_unread_tail():
    doc = Document()
    doc.add_paragraph("正文" * 1500 + "末尾证据")
    buffer = io.BytesIO()
    doc.save(buffer)
    page = read_content(buffer.getvalue(), limit=1)
    assert page["next_offset"] == 0 and page["next_char_offset"] > 0
    following = read_content(
        buffer.getvalue(),
        offset=page["next_offset"],
        limit=1,
        char_offset=page["next_char_offset"],
    )
    assert "末尾证据" in following["text"]
    assert following["next_offset"] is None
