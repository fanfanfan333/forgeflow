"""INC46 T27 — format capability matrix + PPTX/XLSX limited edit + PDF honest downgrade.

This file pins the **explicit** statement of "读 4 种、改 1 种" into a testable
contract: a capability matrix (format × capability × honest reason), the *limited*
edits each editable format really performs, and the *honest downgrade* PDF gets.

Coverage map (red lines 14 / 17):

* **Capability matrix** — derived from the real writers, served by both
  ``GET /documents/capabilities`` and the pure ``render_capability_text`` (which
  must contain the degradation statements).
* **Byte sniffing** — the format is read from the bytes (never a suffix), because
  a registered blob is content-addressed / extensionless.
* **Positive probes** — editing one PPTX slide leaves the *other* slides' XML
  (canonicalised) unchanged; editing an XLSX text cell leaves the formula +
  formats + data-validation + conditional-formatting intact and the formula is
  **not** hard-coded into a value.
* **Negative probes** — an in-place PDF request ⇒ ``unsupported_inplace`` with
  alternatives (never a claimed success); an XLSX constant-over-formula without an
  explicit formula intent ⇒ refused; an over-capability PPTX request ⇒
  ``unsupported``.
* **Counterfactual (真跑)** — removing the formula guard turns the negative case
  red (named red point).

No model, no PG: every fixture is synthesised in-process.
"""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import forgeflow.documents as D
from forgeflow.api.main import app as forgeflow_app
from forgeflow.api.routers import documents as documents_router
from forgeflow.auth.jwt import create_access_token
from forgeflow.middleware.auth import RBACMiddleware

# --------------------------------------------------------------------------- #
# Fixtures / builders                                                         #
# --------------------------------------------------------------------------- #
def _pptx_bytes(*, slides: int = 2) -> bytes:
    """A known PPTX: slide 0 has a title (with a number) + a body text box."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    layout = presentation.slide_layouts[5]  # "Title Only"
    first = presentation.slides.add_slide(layout)
    first.shapes.title.text = "总金额 40 元"
    box = first.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "备注 请核对"
    for index in range(1, slides):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = f"第 {index + 1} 章 明细"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pptx_styled_bytes() -> bytes:
    """A PPTX whose slide-0 title run is explicitly **bold** (style probe)."""
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[5]
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "总金额 40 元"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.bold = True
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _slide_xml(data: bytes) -> list[str]:
    """Each slide's XML, canonicalised (whitespace / attr-order independent)."""
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    return [ET.canonicalize(slide._element.xml) for slide in presentation.slides]


def _slide_title(data: bytes, index: int = 0) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    return str(presentation.slides[index].shapes.title.text or "")


def _title_run_bold(data: bytes, index: int = 0) -> bool:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    paragraph = presentation.slides[index].shapes.title.text_frame.paragraphs[0]
    return bool(paragraph.runs[0].font.bold)


def _xlsx_bytes() -> bytes:
    """A workbook: A2 text, B2 formula, C2 formatted + validated + conditionally formatted."""
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "name"
    sheet["A2"] = "apple"
    sheet["B2"] = "=SUM(1,2)"
    sheet["C2"] = 10
    sheet["C2"].font = Font(bold=True, color="FF0000")
    sheet["C2"].fill = PatternFill("solid", fgColor="FFFF00")
    validation = DataValidation(type="whole", operator="between", formula1="1", formula2="100")
    sheet.add_data_validation(validation)
    validation.add("C2")
    sheet.conditional_formatting.add(
        "C2",
        CellIsRule(operator="greaterThan", formula=["5"], fill=PatternFill("solid", fgColor="FFC7CE")),
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    from forgeflow.documents.pdf_generate import PdfGenerateSpec, apply_spec

    return apply_spec(
        PdfGenerateSpec(title="合同摘要", paragraphs=["金额 40 元", "请核对"], page_size="A4")
    )


def _constant_over_formula_is_refused(data: bytes) -> bool:
    """Try to clobber the formula cell B2 with a constant; report refusal."""
    try:
        D.apply_xlsx_edits(data, [{"op": "set_cell", "cell": "B2", "value": 42}])
    except D.FormulaProtectionError:
        return True
    return False


# --------------------------------------------------------------------------- #
# 1. Capability matrix                                                         #
# --------------------------------------------------------------------------- #
def test_sniff_format_reads_bytes_not_suffix():
    assert D.sniff_format(_pptx_bytes()) == "pptx"
    assert D.sniff_format(_xlsx_bytes()) == "xlsx"
    assert D.sniff_format(_pdf_bytes()) == "pdf"
    import docx as _docx

    doc = _docx.Document()
    doc.add_paragraph("hello")
    buffer = io.BytesIO()
    doc.save(buffer)
    assert D.sniff_format(buffer.getvalue()) == "docx"
    # a bare zip that is neither OOXML family, and junk, both fall to "other"
    assert D.sniff_format(b"PK\x03\x04 not really ooxml") == "other"
    assert D.sniff_format(b"") == "other"


def test_capability_matrix_shape_and_honest_reasons():
    matrix = D.capability_matrix()
    assert list(matrix) == ["docx", "pptx", "xlsx", "pdf", "other"]

    docx = matrix["docx"]
    assert docx["editable"] is True and docx["inplace"] is True and docx["supported"]

    pptx = matrix["pptx"]
    assert pptx["editable"] is True
    assert "set_notes_text" in pptx["supported"]
    assert {"layout", "master", "image"} <= set(pptx["protected"])
    assert {"animation", "complex_shape"} <= set(pptx["unsupported"])

    xlsx = matrix["xlsx"]
    assert "formula" in xlsx["protected"]
    assert {"data_validation", "conditional_formatting", "chart_reference"} <= set(xlsx["preserved"])

    pdf = matrix["pdf"]
    assert pdf["editable"] is False and pdf["inplace"] is False
    assert pdf["supported"] == []
    assert set(pdf["strategies"]) == {"rebuild", "annotate", "refuse"}

    other = matrix["other"]
    assert other["editable"] is False and other["unsupported"] == ["*"]


def test_unsupported_reason_is_honest():
    # a supported op on a supported format has no reason to give
    assert D.unsupported_reason("docx", D.SUPPORTED_OPS[0]) == ""
    # an unknown format is declared unsupported, with a reason
    assert "unsupported" in D.unsupported_reason("png")
    # PDF in-place is the canonical honest refusal
    assert "unsupported_inplace" in D.unsupported_reason("pdf")
    # a format that exists but cannot do the op
    assert D.unsupported_reason("pptx", "delete_slide") != ""
    assert D.format_supports("pptx", "replace_text") is True
    assert D.format_supports("pptx", "delete_slide") is False


def test_render_capability_text_is_pure_and_states_degradation():
    text = D.render_capability_text()
    # the three mandatory honest statements (red line 17)
    assert "不承诺原位编辑" in text
    assert "公式单元格默认受保护" in text
    assert "不改版式/母版/图片" in text
    assert "unsupported" in text
    # pure: calling twice yields the identical string (no side effects)
    assert D.render_capability_text() == text
    # single-format rendering only mentions that format
    only_pdf = D.render_capability_text("pdf")
    assert "PDF" in only_pdf and "重建" in only_pdf


# --------------------------------------------------------------------------- #
# 2. PPTX — limited edit; run style preserved; other slides untouched          #
# --------------------------------------------------------------------------- #
def test_pptx_replace_text_preserves_run_style():
    data = _pptx_styled_bytes()
    assert _title_run_bold(data) is True
    new, changes = D.apply_pptx_edits(data, [{"op": "replace_text", "match": "40", "replace": "38"}])
    assert changes == 1
    assert _slide_title(new) == "总金额 38 元"
    # the run's formatting survived the text edit (INC46 T27 「保留 run 样式」)
    assert _title_run_bold(new) is True


def test_pptx_editing_one_slide_leaves_other_slides_xml_unchanged():
    data = _pptx_bytes(slides=3)
    before = _slide_xml(data)
    new, changes = D.apply_pptx_edits(
        data, [{"op": "set_slide_title", "index": 0, "text": "新标题"}]
    )
    assert changes == 1 and _slide_title(new, 0) == "新标题"
    after = _slide_xml(new)
    # the edited slide changed …
    assert after[0] != before[0]
    # … and **every other** slide is byte-identical after canonicalisation
    assert after[1:] == before[1:]


def test_pptx_notes_edit_is_supported_and_round_trips():
    from pptx import Presentation

    data = _pptx_bytes()
    new, changes = D.apply_pptx_edits(data, [{"op": "set_notes_text", "index": 0, "text": "新的备注"}])
    assert changes == 1
    presentation = Presentation(io.BytesIO(new))
    assert presentation.slides[0].notes_slide.notes_text_frame.text == "新的备注"


def test_pptx_over_capability_is_declared_unsupported():
    animation = D.pptx_unsupported_result("把这一页的动画删掉")
    assert animation is not None and animation.status == "unsupported"
    assert animation.to_dict()["kind"] == "unsupported"
    # layout / master / image are preserved (text edit may proceed) — not edited
    image = D.pptx_unsupported_result("把 logo 图片换掉")
    assert image is not None and image.status == "preserved"
    # an in-scope request yields no boundary result at all
    assert D.pptx_unsupported_result("把标题改成 2026 年总结") is None


# --------------------------------------------------------------------------- #
# 3. XLSX — limited edit; formulas protected; format preserved                 #
# --------------------------------------------------------------------------- #
def test_xlsx_text_edit_keeps_formula_and_formatting_intact():
    import openpyxl

    data = _xlsx_bytes()
    new, changes = D.apply_xlsx_edits(data, [{"op": "set_cell", "cell": "A2", "value": "pear"}])
    assert changes == 1

    sheet = D.open_workbook(new)["Sheet1"]
    assert sheet["A2"].value == "pear"
    # the formula is still a formula (never hard-coded into a value) …
    assert sheet["B2"].value == "=SUM(1,2)"
    # … and a data-only read shows no cached literal was baked in
    cached = openpyxl.load_workbook(io.BytesIO(new), data_only=True)["Sheet1"]["B2"].value
    assert cached is None
    # formatting / data validation / conditional formatting all survived
    assert sheet["C2"].font.bold is True
    assert sheet["C2"].fill.fgColor.rgb.endswith("FFFF00")
    assert len(sheet.data_validations.dataValidation) == 1
    assert list(sheet.conditional_formatting)


def test_xlsx_constant_over_formula_without_intent_is_refused():
    assert _constant_over_formula_is_refused(_xlsx_bytes()) is True, "常量覆盖公式未被拒绝（红线 17）"


def test_xlsx_explicit_formula_intent_allows_the_change():
    data = _xlsx_bytes()
    new, changes = D.apply_xlsx_edits(
        data, [{"op": "set_cell", "cell": "B2", "value": "=SUM(1,3)", "intent": "formula"}]
    )
    assert changes == 1
    assert D.open_workbook(new)["Sheet1"]["B2"].value == "=SUM(1,3)"
    # the global escape hatch works too
    new2, _ = D.apply_xlsx_edits(
        data, [{"op": "set_cell", "cell": "B2", "value": 7}], allow_formula_overwrite=True
    )
    assert D.open_workbook(new2)["Sheet1"]["B2"].value == 7


def test_xlsx_formula_cells_are_reported():
    assert D.formula_cells(_xlsx_bytes()) == [("Sheet1", "B2", "=SUM(1,2)")]


def test_counterfactual_removing_the_formula_guard_turns_the_negative_case_red(monkeypatch):
    import forgeflow.documents.xlsx_edit as xlsx_edit

    data = _xlsx_bytes()
    # wiring in place ⇒ the guard refuses (green)
    assert _constant_over_formula_is_refused(data) is True

    # remove the formula protection: the guard becomes a no-op
    monkeypatch.setattr(xlsx_edit, "_guard_formula_overwrite", lambda *a, **k: None)
    with pytest.raises(AssertionError) as exc:
        assert _constant_over_formula_is_refused(data) is True, "常量覆盖公式未被拒绝（红线 17）"
    # the red point is exactly the refusal assertion
    assert "未被拒绝" in str(exc.value), str(exc.value)


# --------------------------------------------------------------------------- #
# 4. PDF — never an in-place success; honest alternatives                      #
# --------------------------------------------------------------------------- #
def test_pdf_inplace_request_is_unsupported_with_alternatives():
    decision = D.evaluate_pdf_request("请原位编辑这个 PDF，把第一段改掉")
    assert decision.status == D.PDF_INPLACE_UNSUPPORTED
    assert decision.inplace is False
    assert decision.alternatives == list(D.PDF_STRATEGIES)
    # never a claimed success: no strategy is attached to the refusal, and the
    # reason names the alternatives instead of pretending the edit happened
    assert decision.strategy is None
    assert "重建" in decision.reason and "批注" in decision.reason
    assert D.classify_pdf_request("原位修改") == "inplace"


def test_pdf_rebuild_is_a_new_document_not_in_place():
    from forgeflow.documents.pdf_inspect import inspect_pdf

    source = _pdf_bytes()
    new, provenance = D.rebuild_from_pdf(source)
    assert provenance["inplace"] is False
    assert provenance["label"] == "重建（非原位）"
    assert new[:5] == b"%PDF-"
    assert inspect_pdf(new).page_count >= 1
    # the source bytes are never mutated
    assert inspect_pdf(source).page_count == 1


def test_pdf_annotate_overlay_is_non_inplace_and_guards_inputs():
    from forgeflow.documents.pdf_inspect import inspect_pdf

    source = _pdf_bytes()
    annotated = D.annotate_overlay(source, "此处需复核", page_index=0)
    assert annotated[:5] == b"%PDF-"
    assert inspect_pdf(annotated).page_count == 1
    with pytest.raises(D.PdfPolicyError):
        D.annotate_overlay(source, "x", page_index=99)
    with pytest.raises(D.PdfPolicyError):
        D.annotate_overlay(source, "   ")


# --------------------------------------------------------------------------- #
# 5. API surface — the matrix is served AND user-visible                       #
# --------------------------------------------------------------------------- #
def _rbac_client() -> TestClient:
    minimal = FastAPI()
    minimal.add_middleware(RBACMiddleware)
    minimal.include_router(documents_router.router, prefix="/documents")
    return TestClient(minimal)


def test_capabilities_route_is_registered():
    paths = forgeflow_app.openapi().get("paths", {})
    assert "/documents/capabilities" in paths
    assert "get" in paths["/documents/capabilities"]


def test_unauthenticated_capabilities_is_401():
    assert _rbac_client().get("/documents/capabilities").status_code == 401


def test_authenticated_capabilities_returns_matrix_and_text():
    token = create_access_token(user_id="admin-1", role="admin")
    response = _rbac_client().get(
        "/documents/capabilities", headers={"Authorization": f"Bearer {token}"}
    )
    assert response.status_code == 200
    body = response.json()
    assert set(body["matrix"]) == {"docx", "pptx", "xlsx", "pdf", "other"}
    assert body["matrix"]["pdf"]["inplace"] is False
    assert "不承诺原位编辑" in body["text"]
    # the API matrix and the pure user-visible text describe the SAME truth
    assert body["text"] == D.render_capability_text()
