from __future__ import annotations

from tests.business_auth import authenticate_business_app

import uuid
from unittest.mock import AsyncMock

import httpx
import pytest
from fastapi import FastAPI
from langchain_core.messages import AIMessage

from forgeflow.api.routers import skills
from forgeflow.auth.jwt import create_access_token
from forgeflow.config import get_settings
from forgeflow.documents.store import DocArtifactStore
from forgeflow.middleware.auth import RBACMiddleware
from forgeflow.repositories import get_skill_candidate_repository, get_skill_repository
from forgeflow.repositories.memory.skill_repo import clear_skill_store
from forgeflow.resources.service import ResourceService
from forgeflow.runtime import orchestrator as orch
from forgeflow.runtime.input_scope import InputReferenceError, prepare_input_scope
from forgeflow.skills.models import SkillCandidateRecord

pytestmark = pytest.mark.asyncio
INPUT = {
    "name": "文档清洗",
    "description": "清理文档但保留数字",
    "instructions": "读取实际文件，将 DIRTY 替换成 CLEAN，保留数字100，并返回文件。",
    "example": "清洗这份文档",
}


@pytest.fixture(autouse=True)
def isolated(force_memory_backend, tmp_path, monkeypatch):
    clear_skill_store()
    orch.reset_run_store()
    monkeypatch.setattr(get_settings(), "resource_store_root", str(tmp_path))
    yield
    clear_skill_store()
    orch.reset_run_store()


def headers(tenant, role="manager"):
    return {
        "Authorization": "Bearer " + create_access_token("skill-user", role, workspace_id=tenant)
    }


@pytest.fixture
async def client():
    app = FastAPI()
    authenticate_business_app(app)
    app.add_middleware(RBACMiddleware)
    app.include_router(skills.router, prefix="/skills")
    app.include_router(skills.candidates_router, prefix="/skill-candidates")
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app, raise_app_exceptions=False), base_url="http://test"
    ) as c:
        yield c


async def test_resolve_skill_beyond_list_limit_and_reject_foreign_tenant(client):
    registry = skills.SkillRegistry()
    tenant = str(uuid.uuid4())
    for number in range(201):
        await registry.create(tenant, f"older-{number}", "test", description="test")
    target = await registry.create(tenant, "目标清洗技能", "documents", description="保留数字")
    for ref in (target.id, target.name):
        response = await client.get("/skills/resolve", params={"ref": ref}, headers=headers(tenant))
        assert response.status_code == 200, response.text
        assert response.json()["id"] == target.id
        assert (await client.get("/skills/resolve", params={"ref": ref}, headers=headers("foreign"))).status_code == 404
    assert (await client.get("/skills/resolve", params={"ref": "missing"}, headers=headers(tenant))).status_code == 404


async def test_instruction_creation_edit_export_and_tenant_scope(client):
    tenant = str(uuid.uuid4())
    h = headers(tenant)
    created = await client.post("/skills/instructions", headers=h, json=INPUT)
    assert created.status_code == 201, created.text
    skill = created.json()
    assert skill["status"] == "published"
    response = await client.get(f"/skills/{skill['id']}/instructions", headers=h)
    assert response.json()["instructions"] == INPUT["instructions"]
    assert response.json()["source"].startswith("---\nname:")
    assert (
        await client.get(f"/skills/{skill['id']}/instructions", headers=headers("other"))
    ).status_code == 404
    edited = {**INPUT, "instructions": "清理指定内容并保留数据，更新后的工作指令。"}
    update = await client.post(f"/skills/{skill['id']}/instructions", headers=h, json=edited)
    assert update.status_code == 200, update.text
    assert update.json()["current_version"] == "1.0.1"
    versions = (await client.get(f"/skills/{skill['id']}/versions", headers=h)).json()
    assert len(versions) == 2
    assert (
        await client.post("/skills/instructions", headers=headers(tenant, "viewer"), json=INPUT)
    ).status_code == 403
    assert (
        await client.post(
            "/skills/instructions", headers=h, json={**INPUT, "instructions": " " * 20}
        )
    ).status_code == 400


async def test_cannot_use_another_users_trial_run(client):
    tenant = str(uuid.uuid4())
    record = orch.register_running_run(run_id="other-trial", thread_id="trial-thread",
        tenant_id=tenant, actor_user_id="someone-else", intent="test")
    record.status, record.outcome = "completed", "success"
    record.declared_inputs = {"instruction_draft": {"instructions": INPUT["instructions"]}}
    response = await client.post("/skills/instructions", headers=headers(tenant),
                                 json={**INPUT, "trial_run_id": record.run_id})
    assert response.status_code == 409


async def test_creation_failure_does_not_leave_an_empty_skill(client, monkeypatch):
    repo = get_skill_repository()
    monkeypatch.setattr(
        repo, "add_version", AsyncMock(side_effect=RuntimeError("database write failed"))
    )
    h = headers(str(uuid.uuid4()))
    assert (await client.post("/skills/instructions", headers=h, json=INPUT)).status_code == 500
    assert (await client.get("/skills", headers=h)).json()["items"] == []


async def test_selected_instructions_reach_model_and_real_file_tools(client, monkeypatch):
    tenant = str(uuid.uuid4())
    skill = (await client.post("/skills/instructions", headers=headers(tenant), json=INPUT)).json()
    resource = await ResourceService().register_file(tenant, name="report.txt", data=b"DIRTY 100")

    class Model:
        model = "skill-regression"
        _llm_type = "test"
        calls = 0

        def bind_tools(self, tools):
            return self

        def bind(self, **kwargs):
            return self

        async def ainvoke(self, messages, **kwargs):
            assert any(INPUT["instructions"] in str(m.content) for m in messages)
            self.calls += 1
            if self.calls == 1:
                return AIMessage(
                    content="",
                    tool_calls=[{"name": "textfile.inspect", "args": {}, "id": "inspect"}],
                )
            if self.calls == 2:
                return AIMessage(
                    content="",
                    tool_calls=[
                        {
                            "name": "textfile.edit",
                            "args": {
                                "edits": [
                                    {"op": "replace_text", "match": "DIRTY", "replace": "CLEAN"}
                                ]
                            },
                            "id": "edit",
                        }
                    ],
                )
            return AIMessage(content="已按文档清洗技能处理并保留100。")

    model = Model()
    monkeypatch.setattr(get_settings(), "agent_runtime_mode", "react")
    monkeypatch.setattr(
        orch,
        "_build_planner_models",
        lambda: (model, model, [{"slot": "strong", "model": model.model}]),
    )
    task = orch.TaskCreate(
        intent="处理这份文档", context={"resources": [resource.id], "skill_ids": [skill["id"]]}
    )
    handle = await orch.run_task(
        task, orch.RequestContext(tenant_id=tenant, role="manager", user_id="skill-user")
    )
    record = orch.get_run_store().get(handle.run_id)
    assert record.status == "completed", record.errors
    assert record.skills_used == [skill["id"]]
    assert record.declared_inputs["skill_ids"] == [skill["id"]]
    artifact = next(a for a in record.artifacts if a.get("format") == "text")
    assert DocArtifactStore().get(artifact["content_ref"]).decode("utf-8-sig") == "CLEAN 100"
    model.calls = 0
    trial = await client.post(
        "/skills/instructions/try",
        headers=headers(tenant),
        json={"skill": INPUT, "resources": [resource.id], "intent": "清洗文档，返回文件"},
    )
    assert trial.status_code == 200, trial.text
    import asyncio

    for _ in range(100):
        trial_record = orch.get_run_store().get(trial.json()["run_id"])
        if trial_record and trial_record.status not in {"running", "queued"}:
            break
        await asyncio.sleep(0.01)
    assert trial_record.status == "completed", trial_record.errors
    assert (
        trial_record.declared_inputs["instruction_draft"]["instructions"] == INPUT["instructions"]
    )
    saved_trial = await client.post(
        "/skills/instructions",
        headers=headers(tenant),
        json={**INPUT, "name": "试运行后的技能", "trial_run_id": trial_record.run_id},
    )
    # The edit ran, but its generic cleanup requirement has no measured acceptance.
    assert trial_record.outcome == "partial"
    assert saved_trial.status_code == 409, saved_trial.text
    altered = await client.post(
        "/skills/instructions",
        headers=headers(tenant),
        json={
            **INPUT,
            "name": "变化后的技能",
            "instructions": "更换指令后不能复用先前验证证据。",
            "trial_run_id": trial_record.run_id,
        },
    )
    assert altered.status_code == 409
    refused = await client.post(
        "/skills/instructions/try",
        headers=headers(tenant, "viewer"),
        json={"skill": INPUT, "resources": [resource.id], "intent": "清洗文档"},
    )
    assert refused.status_code == 403
    with pytest.raises(InputReferenceError):
        await prepare_input_scope(task, orch.RequestContext(tenant_id="foreign"))
    await client.post(f"/skills/{skill['id']}/archive", headers=headers(tenant))
    with pytest.raises(InputReferenceError):
        await prepare_input_scope(task, orch.RequestContext(tenant_id=tenant))


async def test_candidate_deletion_is_authorized_and_durable(client):
    tenant = str(uuid.uuid4())
    repo = get_skill_candidate_repository()
    record = SkillCandidateRecord(tenant_id=tenant, name="candidate", domain="general")
    await repo.save_candidate(record)
    url = f"/skill-candidates/{record.id}"
    assert (await client.delete(url, headers=headers("other"))).status_code == 404
    assert (await client.delete(url, headers=headers(tenant, "viewer"))).status_code == 403
    assert (await client.delete(url, headers=headers(tenant))).status_code == 200
    assert (await client.get("/skill-candidates", headers=headers(tenant))).json()["items"] == []
    assert (await repo.get_candidate(tenant, record.id)).status == "archived"


async def test_unmatched_literal_edit_cannot_claim_delivery(tmp_path):
    from forgeflow.runtime.tool_handlers import textfile_edit

    path = tmp_path / "sample.txt"
    path.write_text("客户：  张三\n金额：  100万元\n", encoding="utf-8")
    result = await textfile_edit(
        {
            "text_paths": [str(path)],
            "edits": [{"op": "replace_text", "match": r"\s+", "replace": " "}],
        },
        orch.RequestContext(tenant_id="default"),
    )
    assert result["ok"] is False
    assert result["changes"] == 0
    assert "artifact_ref" not in result
    assert "100万元" in path.read_text(encoding="utf-8")


async def test_natural_draft_is_model_generated_and_failures_are_explicit(client, monkeypatch):
    import json
    from forgeflow.models import provider

    class Model:
        _llm_type = "test"

        async def ainvoke(self, messages):
            assert "清洗文档" in str(messages[-1].content)
            return AIMessage(content=json.dumps(INPUT, ensure_ascii=False))

    monkeypatch.setattr(provider, "get_model", lambda **kwargs: Model())
    h = headers("draft-author")
    response = await client.post(
        "/skills/instructions/draft", headers=h, json={"goal": "清洗文档，保留数字并返回修改文件"}
    )
    assert response.status_code == 200, response.text
    assert response.json()["instructions"] == INPUT["instructions"]
    assert (await client.get("/skills", headers=h)).json()["items"] == []
    model = Model()
    model._llm_type = "mock"
    monkeypatch.setattr(provider, "get_model", lambda **kwargs: model)
    response = await client.post(
        "/skills/instructions/draft", headers=h, json={"goal": "清洗文档，保留数字并返回修改文件"}
    )
    assert response.status_code == 503


async def test_sheet_inspection_exposes_actual_cells_and_formulas(tmp_path):
    from openpyxl import Workbook
    from forgeflow.runtime.tool_handlers import sheet_inspect

    path = tmp_path / "actual.xlsx"
    book = Workbook()
    book.active["A1"] = "客户"
    book.active["A2"] = "  张三  "
    book.active["B2"] = 100
    book.active["B3"] = "=SUM(B2)"
    book.save(path)
    result = await sheet_inspect(
        {"sheet_paths": [str(path)]}, orch.RequestContext(tenant_id="default")
    )
    assert result["ok"] is True
    values = {value["cell"]: value["value"] for value in result["cell_values"]}
    assert values["A2"] == "  张三  "
    assert values["B2"] == 100
    assert values["B3"] == "=SUM(B2)"
    assert result["cell_values_truncated"] is False


async def test_skill_metrics_use_actual_scoped_runs_and_unknown_is_not_zero(client):
    from forgeflow.workspace.store import get_workspace_store
    from forgeflow.workspace.models import WorkspaceRunRecord

    tenant = str(uuid.uuid4())
    h = headers(tenant)
    skill = (await client.post("/skills/instructions", headers=h, json=INPUT)).json()
    endpoint = f"/skills/{skill['id']}/usage"
    empty = (await client.get(endpoint, headers=h)).json()
    assert empty["attempts"] == 0 and empty["success_rate"] is None
    store = get_workspace_store()
    for index, outcome in enumerate(["success", "partial"]):
        await store.save(
            WorkspaceRunRecord(
                run_id=str(uuid.uuid4()),
                session_id=str(uuid.uuid4()),
                tenant_id=tenant,
                status="completed",
                outcome=outcome,
                declared_inputs={"skill_ids": [skill["id"]]},
                created_at="2026-10-08T00:00:00+00:00",
                completed_at=f"2026-10-08T00:00:0{2 + 4 * index}+00:00",
            )
        )
    measured = (await client.get(endpoint, headers=h)).json()
    assert measured["attempts"] == 2 and measured["success_rate"] == 0.5
    assert measured["average_seconds"] == 4 and measured["average_cost"] is None
    assert (await client.get(endpoint, headers=headers("foreign"))).status_code == 404


async def test_deterministic_cleaning_preserves_word_tables_excel_formulas_and_code():
    import io
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from forgeflow.documents.docx_edit import apply_edits as edit_word
    from forgeflow.documents.sheet_edit import apply_edits as edit_sheet
    from forgeflow.documents.textfile_edit import apply_edits as edit_text

    edits = [{"op": "normalize_whitespace"}]
    text, changes = edit_text(
        " 客户：  张三\r\n金额：  100万元\r\n".encode(), edits, filename="sample.txt"
    )
    assert text.decode() == "客户： 张三\r\n金额： 100万元\r\n" and changes == 2
    code, _ = edit_text(b'def f():\n    return "a  b"  \n', edits, filename="sample.py")
    assert code == b'def f():\n    return "a  b"\n'
    word = Document()
    run = word.add_paragraph().add_run("客户：  张三")
    run.bold = True
    word.add_table(rows=1, cols=1).cell(0, 0).text = "金额：  100万元"
    original = io.BytesIO()
    word.save(original)
    result, changes = edit_word(original.getvalue(), edits)
    reopened = Document(io.BytesIO(result))
    assert reopened.paragraphs[0].text == "客户： 张三"
    assert reopened.paragraphs[0].runs[0].bold is True
    assert reopened.tables[0].cell(0, 0).text == "金额： 100万元"
    assert changes == 2
    workbook = Workbook()
    workbook.active["A2"] = "  张三  "
    workbook.active["B2"] = 100
    workbook.active["B3"] = "=SUM(B2)"
    original = io.BytesIO()
    workbook.save(original)
    result, changes = edit_sheet(original.getvalue(), edits)
    reopened = load_workbook(io.BytesIO(result))
    assert reopened.active["A2"].value == "张三"
    assert reopened.active["B2"].value == 100 and reopened.active["B3"].value == "=SUM(B2)"
    assert changes == 1
