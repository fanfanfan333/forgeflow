"""INC44 T04 — the PPTX plane through the runtime handlers + artifact projection.

The format is **sniffed from the OOXML package** (a registered document path is
content-addressed and extensionless), so ``document.inspect`` / ``document.edit``
must dispatch to ``python-pptx`` for a ``.pptx`` and report the honest provider —
with no new tool name (design §1.2).
"""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest

from forgeflow.runtime.artifacts import (
    ARTIFACT_KIND_PPTX,
    artifacts_from_invocations,
)


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[5]
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "总金额 40 元"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _ctx() -> SimpleNamespace:
    return SimpleNamespace(run_id="run-pptx", tenant_id="t-pptx", role="manager", args={})


@pytest.fixture
def tmp_doc_store(tmp_path, monkeypatch):
    from forgeflow.documents import store as store_mod

    real = store_mod.DocArtifactStore

    def _factory(root=None):
        return real(root=root if root is not None else tmp_path / "docstore")

    import forgeflow.documents as docs_pkg

    monkeypatch.setattr(store_mod, "DocArtifactStore", _factory)
    monkeypatch.setattr(docs_pkg, "DocArtifactStore", _factory)
    return tmp_path


async def test_document_inspect_dispatches_pptx_to_python_pptx(tmp_path):
    from forgeflow.runtime.tool_handlers import document_inspect

    source = tmp_path / "deck.pptx"
    source.write_bytes(_pptx_bytes())

    ok = await document_inspect({"paths": [str(source)]}, _ctx())

    assert ok["ok"] is True
    assert ok["provider"] == "python-pptx"
    assert ok["format"] == "pptx"
    assert ok["structure"]["slides"] == 1


async def test_document_edit_writes_a_real_pptx(tmp_path, tmp_doc_store):
    from pptx import Presentation

    from forgeflow.documents import DocArtifactStore
    from forgeflow.runtime.tool_handlers import document_edit

    source = tmp_path / "deck.pptx"
    source.write_bytes(_pptx_bytes())

    result = await document_edit(
        {
            "paths": [str(source)],
            "edits": [{"op": "replace_text", "match": "40", "replace": "38"}],
        },
        _ctx(),
    )

    assert result["ok"] is True
    assert result["provider"] == "python-pptx"
    assert result["format"] == "pptx"
    assert result["filename"] == "deck.edited.pptx"
    assert result["artifact_ref"].endswith(".pptx")
    assert result["numeric_changes"] == 1
    assert "content" not in result  # the bytes live in the store, not the payload

    blob = DocArtifactStore(root=tmp_doc_store / "docstore").get(result["artifact_ref"])
    reopened = Presentation(io.BytesIO(blob))
    assert "38" in str(reopened.slides[0].shapes.title.text)


async def test_document_edit_exhaustion_fails_honestly_for_pptx(tmp_path, tmp_doc_store):
    from forgeflow.runtime.tool_handlers import document_edit

    source = tmp_path / "deck.pptx"
    source.write_bytes(_pptx_bytes())
    result = await document_edit(
        {
            "paths": [str(source)],
            "edits": [{"op": "replace_text", "match": "40", "replace": "38"}],
            "max_chars": 1,  # unsatisfiable ⇒ bounded retry then honest failure
        },
        _ctx(),
    )
    assert result["ok"] is False
    assert "校验未通过" in result["reason"]
    assert result["modified"] == 0
    assert "artifact_ref" not in result


async def test_pptx_payload_projects_the_document_pptx_artifact():
    inv = {
        "tool": "document.edit",
        "status": "ok",
        "result_ref": "ab/cd",
        "started_at": "2026-10-03T00:00:00",
        "payload": {
            "ok": True,
            "format": "pptx",
            "filename": "deck.edited.pptx",
            "modified": 1,
            "added": 0,
            "removed": 0,
            "numeric_changes": 1,
            "artifact_ref": "ab/cd.pptx",
        },
    }
    artifacts = artifacts_from_invocations("run-p", [inv])
    assert len(artifacts) == 1
    artifact = artifacts[0]
    assert artifact["kind"] == ARTIFACT_KIND_PPTX == "document_pptx"
    assert artifact["format"] == "pptx"
    assert artifact["title"] == "deck.edited.pptx"
    assert artifact["content_ref"] == "ab/cd.pptx"
    assert artifact["source"] == "document.edit"


async def test_pptx_media_type_is_served():
    from forgeflow.api.routers.runs import _ARTIFACT_MEDIA

    media_type, ext = _ARTIFACT_MEDIA["pptx"]
    assert media_type.endswith("presentationml.presentation")
    assert ext == "pptx"
