"""INC17 T03 — the ReAct closed-loop executor (scripted model + fake tool).

This file pins the core behaviour of :mod:`forgeflow.runtime.react_executor` with
**no daemon and no real provider**: a scripted chat-model stand-in returns
``AIMessage``s with ``tool_calls``, and a fake ``research.search`` handler is
registered so a round really executes through the single honest entry point.

Pinned contract:

  * the loop runs **≥2 rounds** and feeds each tool result back verbatim;
  * the fed-back ``ToolMessage`` content is exactly
    ``json.dumps(record["payload"], sort_keys=True, ensure_ascii=False)`` — i.e. it
    is taken from the (PI-sanitised) execution record, never the raw handler output;
  * the loop ends when the model stops emitting tool calls (``terminated_by
    == "model"``); the run ceiling cuts it off otherwise (``"max_iterations"``)
    and a HITL/RBAC halt cuts it short (``"halted"``);
  * ``report.render`` is never model-visible, always runs as the forced closing
    step, and carries the model's final answer;
  * usage accumulates one entry per model call.
"""

from __future__ import annotations

import json

import pytest
from langchain_core.messages import AIMessage, ToolMessage

import forgeflow.runtime.orchestrator as orch
from forgeflow.config import get_settings
from forgeflow.runtime.orchestrator import (
    RequestContext,
    TaskCreate,
    reset_run_store,
    run_task,
)
from forgeflow.runtime.tool_registry import (
    ToolBinding,
    load_default_bindings,
    register,
    reset_registry,
)

pytestmark = pytest.mark.asyncio


# --------------------------------------------------------------------------- #
# Test doubles                                                                 #
# --------------------------------------------------------------------------- #
class _RecordingBus:
    """Minimal bus stub — run_task / the executor only call ``emit``."""

    def __init__(self) -> None:
        self.events: list[tuple[str, dict]] = []

    async def emit(self, run_id: str, event_type: str, data: dict) -> None:
        self.events.append((event_type, data))


class _ScriptedBound:
    def __init__(self, model: _ScriptedModel) -> None:
        self._model = model

    async def ainvoke(self, messages, **kwargs):  # noqa: ANN001, ANN003
        return self._model._next(messages)


class _ScriptedModel:
    """Returns scripted ``AIMessage``s (with ``tool_calls``) in order.

    Records every message list it is handed (so a test can assert what the model
    *saw*), and repeats its last scripted reply once the script is exhausted so a
    ceiling / replan loop stays deterministic.
    """

    def __init__(self, script: list[dict], usage: tuple[int, int] = (10, 5)) -> None:
        self._script = list(script)
        self._usage = usage
        self._last: dict = {"content": "", "tool_calls": []}
        self.calls: list[list] = []
        self.bound_tools: list | None = None
        self.model = "fake-react-model"
        self._llm_type = "fake"

    def bind_tools(self, tools):  # noqa: ANN001
        self.bound_tools = tools
        return _ScriptedBound(self)

    def bind(self, **kwargs):  # noqa: ANN003
        return _ScriptedBound(self)

    async def ainvoke(self, messages, **kwargs):  # noqa: ANN001, ANN003
        return self._next(messages)

    def _next(self, messages) -> AIMessage:  # noqa: ANN001
        self.calls.append(list(messages))
        if self._script:
            self._last = self._script.pop(0)
        return AIMessage(
            content=self._last.get("content", ""),
            tool_calls=self._last.get("tool_calls", []),
            usage_metadata={
                "input_tokens": self._usage[0],
                "output_tokens": self._usage[1],
                "total_tokens": sum(self._usage),
            },
        )


_UNSET = object()


class _SettingsProxy:
    def __init__(self, real, *, mode: object = _UNSET, provider: object = _UNSET) -> None:
        object.__setattr__(self, "_real", real)
        object.__setattr__(self, "_mode", mode)
        object.__setattr__(self, "_provider", provider)

    def __getattr__(self, item: str):
        if item == "agent_runtime_mode" and self._mode is not _UNSET:
            return self._mode
        if item == "llm_provider" and self._provider is not _UNSET:
            return self._provider
        return getattr(self._real, item)


def _patch_settings(monkeypatch, *, mode="react", provider="mock") -> None:
    proxy = _SettingsProxy(get_settings(), mode=mode, provider=provider)
    monkeypatch.setattr(orch, "get_settings", lambda: proxy)


def _patch_models(monkeypatch, fake: _ScriptedModel) -> None:
    monkeypatch.setattr(
        orch,
        "_build_planner_models",
        lambda: (fake, fake, [{"slot": "strong", "class": "_ScriptedModel", "model": fake.model}]),
    )


async def _fake_research(args, ctx):  # noqa: ANN001
    return {
        "ok": True,
        "provider": "test",
        "query": args.get("query"),
        "results": [{"title": "t", "url": "https://example.org/x"}],
        "summary": "fake search ok",
    }


@pytest.fixture(autouse=True)
def _clean_state():
    reset_run_store()
    yield
    reset_run_store()
    reset_registry()


@pytest.fixture
def fake_research():
    """Register a deterministic fake ``research.search`` binding (real kind)."""
    load_default_bindings()
    register(
        ToolBinding(
            tool_id="research.search",
            handler=_fake_research,
            kind="real",
            provider="test",
            description="scripted test search",
        )
    )
    yield


def _ctx(role: str = "admin", tenant: str = "t-inc17") -> RequestContext:
    return RequestContext(tenant_id=tenant, user_id="u-1", role=role)


async def _drive(monkeypatch, script, *, context=None, role="admin", usage=(10, 5), intent="检索并总结"):
    """Patch settings+models, run a react task, and hand back the artefacts."""
    fake = _ScriptedModel(script, usage=usage)
    _patch_settings(monkeypatch, mode="react", provider="mock")
    _patch_models(monkeypatch, fake)
    task = TaskCreate(intent=intent, workflow_type="generic", context=context or {})
    bus = _RecordingBus()
    handle = await run_task(task, _ctx(role=role), bus=bus)
    return fake, task, bus, handle


async def test_functional_package_uses_real_skill_reader_in_react_loop(monkeypatch, force_memory_backend):
    from types import SimpleNamespace

    from forgeflow.skills.functional_evaluation import execute_case
    from forgeflow.skills.package_import import package_spec

    spec = package_spec({
        "SKILL.md": "---\nname: review-test\ndescription: Review the attached document.\nallowed-tools: skills.read\n---\nRead references/check.md to obtain the review checklist.",
        "references/check.md": "Check that all requirements have execution evidence.",
    })
    fake = _ScriptedModel([
        {"tool_calls": [{"name": "skills.read", "args": {"skill_id": "instruction-draft", "version": "draft"}, "id": "s1"}]},
        {"tool_calls": [{"name": "skills.read", "args": {"skill_id": "instruction-draft", "ref": "references/check.md"}, "id": "s2"}]},
        {"content": "Review checklist read; no document was supplied for review."},
    ])
    _patch_settings(monkeypatch)
    _patch_models(monkeypatch, fake)
    record, _ = await execute_case(SimpleNamespace(
        instruction="Read the selected skill review checklist without claiming to review a document.",
        expected_summary=json.dumps({"assertions": [{"kind": "report_answer_contains", "value": "checklist"}]}),
    ), spec, _ctx(), "package-reader-test")
    payloads = [r["payload"] for r in record.tool_invocations if r["tool"] == "skills.read"]
    assert "上述正文已加载" in fake.calls[0][1].content
    assert "使用技能前调用 skills.read" not in fake.calls[0][1].content
    assert len(payloads) == 2
    assert "references/check.md" in json.dumps(payloads[0])
    assert "all requirements have execution evidence" in json.dumps(payloads[1])
    assert "script_executed" in json.dumps(payloads[1])


# --------------------------------------------------------------------------- #
# 1. ≥2 rounds, verbatim feedback, model convergence                            #
# --------------------------------------------------------------------------- #
async def test_multi_round_loop_feeds_tool_results_back(monkeypatch, force_memory_backend, fake_research):
    script = [
        {"content": "先查资料", "tool_calls": [
            {"name": "research.search", "args": {"query": "q1"}, "id": "c1"}]},
        {"content": "再查一次", "tool_calls": [
            {"name": "research.search", "args": {"query": "q2"}, "id": "c2"}]},
        {"content": "最终答复：已完成检索与分析", "tool_calls": []},
    ]
    fake, task, bus, handle = await _drive(monkeypatch, script)

    meta = task.context["llm_runtime"]
    records = task.context["tool_invocations"]

    # 3 model calls (2 tool rounds + 1 convergence) ⇒ 3 usage entries.
    assert len(fake.calls) == 3
    assert len(task.context["llm_usage"]) == 3
    # Two tool rounds + the platform-forced report.render (duplicates preserved).
    assert [r["tool"] for r in records] == [
        "research.search", "research.search", "report.render"
    ]
    assert len(meta["rounds"]) == 2
    assert meta["terminated_by"] == "model"
    assert meta["final_answer"]["text"] == "最终答复：已完成检索与分析"

    # The 2nd model call's input contains the 1st round's tool result, VERBATIM.
    second = fake.calls[1]
    tool_msgs = [m for m in second if isinstance(m, ToolMessage)]
    assert tool_msgs, "第二轮模型输入必须含上一轮工具返回"
    assert tool_msgs[0].content == json.dumps(
        records[0]["payload"], sort_keys=True, ensure_ascii=False
    )
    assert tool_msgs[0].content  # non-empty
    # research.search output is PI-sanitised *before* it reaches the record/model.
    assert "<UNTRUSTED_TOOL_OUTPUT" in records[0]["payload"]

    # The loop ended on the model's own term (no more calls after convergence).
    assert len(fake.calls) == 3


async def test_final_answer_is_carried_by_report_render(monkeypatch, force_memory_backend, fake_research):
    script = [{"content": "结论：一切正常", "tool_calls": []}]
    _fake, task, _bus, handle = await _drive(monkeypatch, script)

    records = task.context["tool_invocations"]
    assert records[-1]["tool"] == "report.render"
    report_content = records[-1]["payload"]["content"]
    assert "## 最终答案" in report_content
    assert "结论：一切正常" in report_content
    # The deliverable is produced ONLY by report.render.
    assert any(a["source"] == "report.render" for a in handle.detail["artifacts"])
    assert "结论：一切正常" in handle.detail["artifacts"][0]["content"]


# --------------------------------------------------------------------------- #
# 2. report.render is never model-visible but is always the forced closer       #
# --------------------------------------------------------------------------- #
async def test_report_render_excluded_from_allowed_tools_but_forced(monkeypatch, force_memory_backend, fake_research):
    script = [{"content": "答", "tool_calls": []}]
    _fake, task, _bus, _handle = await _drive(monkeypatch, script)

    meta = task.context["llm_runtime"]
    assert "report.render" not in meta["allowed_tools"]
    assert meta["platform_forced_tools"] == ["report.render"]
    # The schema handed to the model also excludes it.
    assert all(s["function"]["name"] != "report.render" for s in _fake.bound_tools)


# --------------------------------------------------------------------------- #
# 3. An unknown tool is rejected (no record) and explained to the model          #
# --------------------------------------------------------------------------- #
async def test_unknown_tool_is_rejected_without_a_record(monkeypatch, force_memory_backend, fake_research):
    script = [
        {"content": "试着调用未知工具", "tool_calls": [
            {"name": "bogus.tool", "args": {}, "id": "x1"}]},
        {"content": "最终答复", "tool_calls": []},
    ]
    fake, task, bus, _handle = await _drive(monkeypatch, script)

    records = task.context["tool_invocations"]
    # bogus.tool produced NO record — only the forced report.render remains.
    assert [r["tool"] for r in records] == ["report.render"]
    # The rejection is fed back so the model can correct itself.
    tm = [m for m in fake.calls[1] if isinstance(m, ToolMessage)]
    assert tm and "bogus.tool" in tm[0].content
    # And it is surfaced as an honest warning.
    warns = [d for t, d in bus.events if t == "run.warning" and d.get("reason") == "tool_not_allowed"]
    assert warns and warns[0]["tool"] == "bogus.tool"
    # No round was recorded for the rejected call.
    assert len(task.context["llm_runtime"]["rounds"]) == 0


# --------------------------------------------------------------------------- #
# 4. The round ceiling cuts the loop off (never a fabricated answer)             #
# --------------------------------------------------------------------------- #
async def test_max_iterations_cutoff(monkeypatch, force_memory_backend, fake_research):
    script = [{"content": "继续查", "tool_calls": [
        {"name": "research.search", "args": {"query": "q"}, "id": "c"}]}]
    _fake, task, bus, _handle = await _drive(monkeypatch, script, context={"max_react_iterations": 2})

    meta = task.context["llm_runtime"]
    assert meta["terminated_by"] == "max_iterations"
    assert meta["final_answer"] is None
    assert meta["iterations"] == 2
    assert _handle.detail["errors"]
    assert _handle.status != "completed"
    # Exactly 2 tool rounds were executed, then it stopped (no 3rd model call).
    assert len(meta["rounds"]) == 2
    # The ceiling is announced, never disguised as the answer.
    warns = [d for t, d in bus.events if t == "run.warning" and d.get("reason") == "max_iterations"]
    assert warns
    content = task.context["tool_invocations"][-1]["payload"]["content"]
    assert "因达到轮次上限" in content


async def test_failed_tool_cannot_be_hidden_by_models_success_answer(
    monkeypatch, force_memory_backend, fake_research,
):
    async def failure(args, ctx):
        return {"ok": False, "reason": "真实校验失败"}

    register(ToolBinding(tool_id="research.search", handler=failure, kind="real", provider="test", description="failure probe"))
    script = [
        {"content": "检查", "tool_calls": [{"name": "research.search", "args": {}, "id": "c"}]},
        {"content": "所有要求都已完成", "tool_calls": []},
    ]
    _, task, _, handle = await _drive(monkeypatch, script)
    assert handle.status == "failed"
    meta = task.context["llm_runtime"]
    assert meta["model_final_answer_text"] == "所有要求都已完成"
    assert "未完成验收" in meta["final_answer"]["text"]
    assert "真实校验失败" in meta["final_answer"]["text"]


@pytest.mark.parametrize("complete", [True, False])
async def test_readonly_document_questions_require_complete_cited_answers(
    monkeypatch, force_memory_backend, complete,
):
    import io

    from docx import Document

    from forgeflow.resources.service import ResourceService

    doc = Document()
    doc.add_paragraph("T28 前置 T12")
    doc.add_paragraph("预算 25 步")
    buffer = io.BytesIO()
    doc.save(buffer)
    resource = await ResourceService().register_file("t-inc17", name="questions.docx", data=buffer.getvalue())
    answers = [
        {"index": 1, "text": "前置 T12", "source": "paragraph:0", "quote": "T28 前置 T12"},
        {"index": 2, "text": "预算 25 步", "source": "paragraph:1", "quote": "预算 25 步"},
    ]
    script = [
        {"content": "读取原文", "tool_calls": [{"name": "document.inspect", "args": {}, "id": "d1"}]},
        {"content": json.dumps({"answers": answers if complete else answers[:1]}), "tool_calls": []},
    ]
    _, task, _, handle = await _drive(
        monkeypatch, script, context={"resources": [resource.id]},
        intent="只读取文档，1. T28 前置是什么？2. 预算是多少？",
    )
    assert handle.status == ("completed" if complete else "failed")
    assert task.context["llm_runtime"]["document_questions"]["coverage_verified"] is complete


# --------------------------------------------------------------------------- #
# 5. Ceiling default == 6 and fail-closed on an invalid override                 #
# --------------------------------------------------------------------------- #
async def test_default_iterations_value_is_six(force_memory_backend):
    from forgeflow.runtime import react_executor as rx

    assert rx.MAX_REACT_ITERATIONS == 6


@pytest.mark.parametrize(
    "raw,expected",
    [
        (6, 6),
        (2, 2),
        (0, 6),          # zero ⇒ ignored (fail-closed)
        (-3, 6),         # negative ⇒ ignored
        ("3", 6),        # non-int ⇒ ignored
        (1.5, 6),        # float ⇒ ignored
        (True, 6),       # bool is an int subclass ⇒ explicitly ignored
        (None, 6),       # missing ⇒ constant
    ],
)
async def test_max_react_iterations_override_is_fail_closed(monkeypatch, force_memory_backend, fake_research, raw, expected):
    script = [{"content": "答", "tool_calls": []}]
    context = {} if raw is None else {"max_react_iterations": raw}
    _fake, task, _bus, _handle = await _drive(monkeypatch, script, context=context)
    assert task.context["llm_runtime"]["max_iterations"] == expected


# --------------------------------------------------------------------------- #
# 6. HITL / RBAC halt: terminated_by == "halted", skip report.render, 1:1        #
# --------------------------------------------------------------------------- #
async def test_hitl_high_risk_halts_and_skips_report(monkeypatch, force_memory_backend, fake_research):
    script = [{"content": "发起转账", "tool_calls": [
        {"name": "payment.transfer", "args": {"amount": 1}, "id": "p1"}]}]
    _fake, task, bus, handle = await _drive(monkeypatch, script)

    meta = task.context["llm_runtime"]
    assert meta["terminated_by"] == "halted"
    assert meta["final_answer"] is None
    # The blocked (high-risk) call produced NO record, and report.render was
    # skipped — so the trail is empty and steps/records stay 1:1 (0 == 0).
    records = task.context["tool_invocations"]
    assert records == []
    assert not any(r["tool"] == "report.render" for r in records)
    assert handle.detail["steps"] == []
    assert len(handle.detail["steps"]) == len(records)
    assert handle.detail["errors"], "HITL 拦截必须留下错误"
    assert any(t == "run.error" for t, _ in bus.events)


# --------------------------------------------------------------------------- #
# 7. steps / records 1:1 (cumulative) and the plan is the react projection       #
# --------------------------------------------------------------------------- #
async def test_steps_records_and_plan_are_1to1(monkeypatch, force_memory_backend, fake_research):
    script = [
        {"content": "查", "tool_calls": [{"name": "research.search", "args": {"query": "q"}, "id": "c1"}]},
        {"content": "答", "tool_calls": []},
    ]
    _fake, task, _bus, handle = await _drive(monkeypatch, script)

    records = task.context["tool_invocations"]
    steps = handle.detail["steps"]
    assert len(steps) == len(records) == 2
    # 1:1 step_id alignment between the plan and the records.
    plan = task.context["plan"]
    assert len(plan["steps"]) == len(records)
    assert [s["step_id"] for s in plan["steps"]] == [r["step_id"] for r in records]
    assert plan["source"] == "react"
    # Last plan step / last record is the deliverable.
    assert plan["steps"][-1]["tool"] == "report.render"
    assert records[-1]["tool"] == "report.render"



async def test_durable_resume_keeps_feedback_without_repeating_tools(monkeypatch, force_memory_backend, fake_research, tmp_path):
    import asyncio

    import forgeflow.runtime.recovery as recovery
    from forgeflow.runtime.execution_budget import current_budget

    store = recovery.RecoveryStore(tmp_path / "recover.sqlite3")
    monkeypatch.setattr(recovery, "RecoveryStore", lambda: store)
    first = _ScriptedModel([{"content": "read", "tool_calls": [
        {"name": "research.search", "args": {"query": "q1"}, "id": "c1"}]}])
    original_next = first._next
    def cancel_after_tool(messages):
        if first.calls:
            raise asyncio.CancelledError()
        return original_next(messages)
    first._next = cancel_after_tool
    _patch_settings(monkeypatch)
    _patch_models(monkeypatch, first)
    task = TaskCreate(intent="读取来源", context={})
    task._enable_recovery = True
    with pytest.raises(asyncio.CancelledError):
        await run_task(task, _ctx(), run_id="original-recovery", bus=_RecordingBus())
    saved = store.get("t-inc17", "u-1", "original-recovery")
    assert saved["status"] == "paused"
    assert saved["state"]["inflight"] is False
    assert saved["state"]["budget"]["model_calls"] == 2
    assert saved["state"]["budget"]["tool_calls"] == 1
    assert current_budget.get() is None
    resumed = store.claim("t-inc17", "u-1", "original-recovery", saved["version"])
    second = _ScriptedModel([{"content": "已基于原工具结果完成", "tool_calls": []}])
    _patch_models(monkeypatch, second)
    next_task = TaskCreate(**resumed["task"])
    next_task._enable_recovery = True
    next_task._recovery_state = resumed
    next_task._recovery_source = "original-recovery"
    await run_task(next_task, _ctx(), run_id="resumed-recovery", bus=_RecordingBus())
    calls = next_task.context["tool_invocations"]
    assert sum(call["tool"] == "research.search" for call in calls) == 1
    assert any(isinstance(message, ToolMessage) for message in second.calls[0])
    assert next_task.context["llm_runtime"]["execution_budget"]["model_calls"] == 3



async def test_functional_trial_does_not_recall_or_learn_from_holdouts(monkeypatch, force_memory_backend, fake_research):
    from unittest.mock import AsyncMock
    from forgeflow.experience.extractor import ExperienceExtractor
    from forgeflow.skills.learning import verified_run
    import forgeflow.context.viking_sync as sync

    _patch_settings(monkeypatch)
    _patch_models(monkeypatch, _ScriptedModel([{"content":"已完成", "tool_calls":[]}]))
    extracted=AsyncMock(side_effect=AssertionError("holdout must not be learned"))
    queued=AsyncMock()
    monkeypatch.setattr(ExperienceExtractor,"extract",extracted)
    monkeypatch.setattr(sync,"queue_run",queued)
    task=TaskCreate(intent="试运行",context={})
    task._functional_eval=True
    handle=await run_task(task,_ctx(),bus=_RecordingBus())
    record=orch.get_run_store().get(handle.run_id)
    assert record.experience_id is None and record.llm["evaluation_only"] is True
    assert not verified_run(record)
    extracted.assert_not_awaited()
    queued.assert_not_awaited()
