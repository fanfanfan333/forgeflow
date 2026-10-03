"""INC46 T23 — five-layer validation stack pins (+ the module→package migration).

What this file nails down
-------------------------
0. **Migration** — ``forgeflow.documents.validation`` is now a *package* but every
   public import path is preserved verbatim: ``from forgeflow.documents.validation
   import verify_docx`` **and** ``from forgeflow.documents import verify_docx`` are
   the **same object**, and the package ``__all__`` is exactly the legacy six.
1. **L1** — ``pass`` on a real package; ``fail`` on a non-ZIP and on a package
   missing ``[Content_Types].xml``; the absent second engine is reported
   "not measured" (never gating).
2. **L2** — ``pass`` on a compliant edit; ``fail`` on a lost amount, an
   out-of-range change, a new number, a heading-count change, a table merge
   change (each a real mutation).
3. **L3** — ``pass`` when the edited text inherits the original run style;
   ``fail`` when out-of-range formatting changes and when in-range style is lost.
4. **L4** — ``None`` (not ``pass``) whenever the engine/rasteriser is unavailable
   (red line 15); a real injected render diff ⇒ ``fail``; identical rasters ⇒
   ``pass``.
5. **L5** — no judge ⇒ ``None``; a good injected judge ⇒ ``pass``; a low injected
   judge ⇒ ``needs_review`` (advisory).
6. **Stack** — aggregation precedence, ``unmeasured`` surfacing, the
   ``disabled_layers`` counterfactual hook, and an unreadable input becoming an
   honest ``fail`` (never a crash, never a pass).

Non-empty evidence (red line 4 / 15): every measured verdict carries a
reproducible ``evidence_ref``; an unmeasured verdict is ``status=None``.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from forgeflow.documents import (
    VerifyReport,
    verify_docx,
    verify_pdf,
    verify_pptx,
    verify_sheet,
    verify_textfile,
)
from forgeflow.documents import validation as validation_pkg
from forgeflow.documents.validation import (
    LayerVerdict,
    ValidationConfig,
    ValidationVerdict,
    validate_document,
    verify_docx as package_verify_docx,
)
from forgeflow.documents.validation import (
    format as format_layer,
    invariants as invariants_layer,
    render as render_layer,
    semantic as semantic_layer,
    structural as structural_layer,
)
from forgeflow.documents.validation.legacy import verify_docx as legacy_verify_docx
from forgeflow.documents.validation.verdict import (
    FAIL,
    LAYER_FORMAT,
    LAYER_INVARIANTS,
    LAYER_ORDER,
    LAYER_RENDER,
    LAYER_SEMANTIC,
    LAYER_STRUCTURAL,
    NEEDS_REVIEW,
    PASS,
)

_MARKER = "[[TGT::unit]]"
_REPLACEMENT = "已编辑::unit"


# --------------------------------------------------------------------------- #
# Fixtures / builders                                                          #
# --------------------------------------------------------------------------- #
def _save(document: Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sample_doc() -> bytes:
    """A body with a heading, a target paragraph (amount+date), and a table."""
    document = Document()
    document.add_heading("报告标题", level=1)
    document.add_paragraph("文档编号：UNIT；用途：五层验证。")
    document.add_paragraph(
        f"占位 {_MARKER}，金额 1,000.00 元，日期 2026-10-03，只应改这一段。"
    )
    document.add_heading("正文小节", level=2)
    document.add_paragraph("正文内容，不应被改动。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    return _save(document)


def _target_range(data: bytes) -> tuple[int, int]:
    from forgeflow.documents.docx_inspect import open_docx

    texts = [p.text for p in open_docx(data).paragraphs]
    index = next(i for i, text in enumerate(texts) if _MARKER in text)
    return index, index + 1


def _compliant_edit(data: bytes) -> bytes:
    from forgeflow.documents.docx_edit import EditOp, apply_edits

    edited, changes = apply_edits(
        data, [EditOp(op="replace_text", match=_MARKER, replace=_REPLACEMENT)]
    )
    assert changes == 1, changes
    return edited


def _repack(parts: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in sorted(parts):
            archive.writestr(name, parts[name])
    return buffer.getvalue()


def _read_parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _drop_part(data: bytes, name: str) -> bytes:
    parts = _read_parts(data)
    assert name in parts, name
    del parts[name]
    return _repack(parts)


def _mutate_text(data: bytes, old: bytes, new: bytes) -> bytes:
    parts = _read_parts(data)
    blob = parts["word/document.xml"]
    assert old in blob, old
    parts["word/document.xml"] = blob.replace(old, new)
    return _repack(parts)


def _bolden_out_of_range(data: bytes, index: int = 0) -> bytes:
    document = Document(io.BytesIO(data))
    document.paragraphs[index].runs[0].bold = True
    return _save(document)


# --------------------------------------------------------------------------- #
# 0. module → package migration (public import paths preserved)                 #
# --------------------------------------------------------------------------- #
def test_validation_is_a_package_and_all_import_paths_are_the_same_objects():
    assert validation_pkg.__file__.endswith("validation\\__init__.py") or validation_pkg.__file__.endswith(
        "validation/__init__.py"
    )
    assert verify_docx is package_verify_docx is legacy_verify_docx
    for name in ("VerifyReport", "verify_docx", "verify_pptx", "verify_textfile", "verify_sheet", "verify_pdf"):
        assert hasattr(validation_pkg, name), name


def test_validation_package_all_is_exactly_the_legacy_surface():
    assert validation_pkg.__all__ == [
        "VerifyReport",
        "verify_docx",
        "verify_pptx",
        "verify_textfile",
        "verify_sheet",
        "verify_pdf",
    ]


@pytest.mark.parametrize(
    "func",
    [verify_docx, verify_pptx, verify_sheet, verify_textfile, verify_pdf],
)
def test_all_verify_planes_still_return_the_tri_state_report(func):
    report = func(b"definitely not a real document of any kind")
    assert isinstance(report, VerifyReport)
    payload = report.to_dict()
    assert set(payload) == {"openable", "structure_ok", "data_ok", "requirement_ok", "notes"}
    # the tri-state honesty rule: every dimension is True / False / None
    for key in ("structure_ok", "data_ok", "requirement_ok"):
        assert payload[key] in (True, False, None)


# --------------------------------------------------------------------------- #
# 1. L1 structural                                                              #
# --------------------------------------------------------------------------- #
def test_L1_passes_on_a_real_package_and_reports_the_absent_second_engine():
    verdict = structural_layer.structural_verdict(_sample_doc())
    assert verdict.layer == LAYER_STRUCTURAL
    assert verdict.status == PASS
    assert verdict.evidence_ref
    # the optional second engine is honestly reported (never silently claimed)
    assert "second_engine=" in verdict.evidence_ref


def test_L1_fails_on_a_non_zip_blob():
    verdict = structural_layer.structural_verdict(b"this is not a zip at all")
    assert verdict.status == FAIL
    assert verdict.evidence_ref


def test_L1_fails_when_content_types_is_missing():
    broken = _drop_part(_sample_doc(), "[Content_Types].xml")
    verdict = structural_layer.structural_verdict(broken)
    assert verdict.status == FAIL


def test_L1_package_integrity_flags_an_unresolved_relationship():
    parts = _read_parts(_sample_doc())
    rels = parts["word/_rels/document.xml.rels"]
    parts["word/_rels/document.xml.rels"] = rels.replace(
        b"Target=", b"Target=missing_", 1
    )
    ok, evidence, _detail = structural_layer.package_integrity(_repack(parts))
    assert ok is False
    assert "unresolved" in evidence or "FAIL" in evidence


# --------------------------------------------------------------------------- #
# 2. L2 invariants                                                              #
# --------------------------------------------------------------------------- #
def test_L2_passes_on_a_compliant_edit():
    data = _sample_doc()
    start, end = _target_range(data)
    verdict = invariants_layer.invariants_verdict(data, _compliant_edit(data), start=start, end=end)
    assert verdict.status == PASS
    assert verdict.evidence_ref
    assert verdict.detail["out_of_range_changes"] == []


def test_L2_fails_when_an_in_range_amount_is_lost():
    data = _sample_doc()
    start, end = _target_range(data)
    broken = _mutate_text(data, b"1,000.00", b"9,999.99")
    verdict = invariants_layer.invariants_verdict(data, broken, start=start, end=end)
    assert verdict.status == FAIL
    assert any("amount" in failure for failure in verdict.detail["lost"])
    assert "FAIL" in verdict.evidence_ref


def test_L2_fails_when_an_out_of_range_paragraph_changes():
    data = _sample_doc()
    start, end = _target_range(data)
    # the amendment lands on the *receipt* text of paragraph 0, outside [start, end)
    broken = _mutate_text(data, "报告标题".encode(), "报告标题（改）".encode())
    verdict = invariants_layer.invariants_verdict(data, broken, start=start, end=end)
    assert verdict.status == FAIL
    assert verdict.detail["out_of_range_changes"] == ["paragraph[0]"]


def test_L2_fails_when_a_new_number_appears():
    data = _sample_doc()
    start, end = _target_range(data)
    broken = _mutate_text(data, "只应改这一段".encode(), "只应改这一段，新增 7 项".encode())
    verdict = invariants_layer.invariants_verdict(data, broken, start=start, end=end)
    assert verdict.status == FAIL
    assert verdict.detail["new_numbers"]


def test_L2_fails_when_the_heading_count_changes():
    data = _sample_doc()
    start, end = _target_range(data)
    document = Document(io.BytesIO(data))
    document.add_heading("新增一级标题", level=1)
    verdict = invariants_layer.invariants_verdict(data, _save(document), start=start, end=end)
    assert verdict.status == FAIL
    assert verdict.detail["heading_delta"] == 1


def test_L2_fails_when_a_table_merge_structure_changes():
    data = _sample_doc()
    start, end = _target_range(data)
    document = Document(io.BytesIO(data))
    table = document.tables[0]
    table.cell(0, 0).merge(table.cell(0, 1))  # now a single horizontally-merged cell
    verdict = invariants_layer.invariants_verdict(data, _save(document), start=start, end=end)
    assert verdict.status == FAIL
    assert verdict.detail["table_changes"]


# --------------------------------------------------------------------------- #
# 3. L3 format fidelity                                                         #
# --------------------------------------------------------------------------- #
def test_L3_passes_when_edited_text_inherits_the_original_run_style():
    data = _sample_doc()
    start, end = _target_range(data)
    verdict = format_layer.format_verdict(data, _compliant_edit(data), start=start, end=end)
    assert verdict.status == PASS
    assert verdict.evidence_ref
    assert verdict.detail["style_inheritance_failures"] == []


def test_L3_fails_when_an_out_of_range_paragraph_is_reformatted():
    data = _sample_doc()
    start, end = _target_range(data)
    broken = _bolden_out_of_range(data, index=0)
    verdict = format_layer.format_verdict(data, broken, start=start, end=end)
    assert verdict.status == FAIL
    assert "paragraph[0]" in verdict.detail["out_of_range_changes"]


def test_L3_fails_when_the_edited_run_loses_its_formatting():
    data = _sample_doc()
    start, end = _target_range(data)
    # make the in-range run bold in the *before* document …
    before_doc = Document(io.BytesIO(data))
    before_doc.paragraphs[start].runs[0].bold = True
    before = _save(before_doc)
    # … then edit in a way that rebuilds the run *without* the bold
    after_doc = Document(io.BytesIO(before))
    run = after_doc.paragraphs[start].runs[0]
    rpr = run._r.find(qn("w:rPr"))
    if rpr is not None:
        bold = rpr.find(qn("w:b"))
        if bold is not None:
            rpr.remove(bold)
    run.text = run.text.replace(_MARKER, _REPLACEMENT)
    after = _save(after_doc)
    verdict = format_layer.format_verdict(before, after, start=start, end=end)
    assert verdict.status == FAIL
    assert verdict.detail["style_inheritance_failures"] == [f"paragraph[{start}]"]


# --------------------------------------------------------------------------- #
# 4. L4 render (unmeasured ⇒ None, never pass)                                  #
# --------------------------------------------------------------------------- #
def test_L4_is_none_not_pass_when_no_engine_is_available(monkeypatch):
    monkeypatch.setattr(render_layer, "probe_render_engine", lambda: None)
    verdict = render_layer.render_verdict(_sample_doc(), _sample_doc())
    assert verdict.status is None
    assert verdict.status != PASS
    assert verdict.measured is False


def test_L4_is_none_when_conversion_fails():
    verdict = render_layer.render_verdict(
        _sample_doc(), _sample_doc(), engine="no-such-engine", converter=lambda src, eng: None
    )
    assert verdict.status is None


def test_L4_is_none_when_rasterisation_is_unavailable():
    verdict = render_layer.render_verdict(
        _sample_doc(),
        _sample_doc(),
        engine="fake",
        converter=lambda src, eng: b"%PDF-1.4 fake",
        rasterizer=lambda pdf, tool, dpi: None,
    )
    assert verdict.status is None


def test_L4_fails_on_a_real_injected_pixel_difference_and_passes_when_identical():
    grid_a = [[(0, 0, 0)] * 4 for _ in range(4)]
    grid_b = [[(255, 255, 255)] * 4 for _ in range(4)]

    def raster(pdf: bytes, tool: str, dpi: int):
        return grid_a if pdf == b"PDF-A" else grid_b

    def converter(src: bytes, eng: str):
        return b"PDF-A" if src == b"DOC-A" else b"PDF-B"

    different = render_layer.render_verdict(
        b"DOC-A", b"DOC-B", engine="fake", converter=converter, rasterizer=raster
    )
    assert different.status == FAIL
    assert different.detail["diff_ratio"] == 1.0

    same = render_layer.render_verdict(
        b"DOC-A", b"DOC-A", engine="fake", converter=converter, rasterizer=raster
    )
    assert same.status == PASS
    assert same.detail["diff_ratio"] == 0.0


# --------------------------------------------------------------------------- #
# 5. L5 semantic (advisory; unmeasured ⇒ None)                                  #
# --------------------------------------------------------------------------- #
def _score(value: float, judge: str) -> semantic_layer.SemanticScore:
    return semantic_layer.SemanticScore(
        score=value,
        dimensions={name: value for name in semantic_layer.DIMENSIONS},
        rationale="unit",
        judge=judge,
    )


def test_L5_is_none_not_pass_when_no_judge_is_wired(monkeypatch):
    monkeypatch.setattr(semantic_layer, "_DEFAULT_JUDGE", None)
    data = _sample_doc()
    start, end = _target_range(data)
    verdict = semantic_layer.semantic_verdict(data, _compliant_edit(data), start=start, end=end)
    assert verdict.status is None
    assert verdict.status != PASS


def test_L5_passes_and_needs_review_with_injected_judges():
    data = _sample_doc()
    start, end = _target_range(data)
    edited = _compliant_edit(data)

    good = semantic_layer.semantic_verdict(
        data, edited, start=start, end=end, judge=lambda req: _score(0.95, "good")
    )
    assert good.status == PASS
    assert "good" in good.evidence_ref

    low = semantic_layer.semantic_verdict(
        data, edited, start=start, end=end, judge=lambda req: _score(0.1, "low")
    )
    assert low.status == NEEDS_REVIEW


def test_L5_effective_score_is_the_minimum_of_dims():
    score = semantic_layer.SemanticScore(
        score=0.9, dimensions={"风格达标": 0.3, "事实未变": 0.9, "无新增数字/名称": 0.9, "语言通顺": 0.9}
    )
    assert score.effective() == 0.3


# --------------------------------------------------------------------------- #
# 6. Stack aggregation / disable / error                                        #
# --------------------------------------------------------------------------- #
def test_layer_order_and_verdict_dict_shape():
    assert LAYER_ORDER == (LAYER_STRUCTURAL, LAYER_INVARIANTS, LAYER_FORMAT, LAYER_RENDER, LAYER_SEMANTIC)
    verdict = LayerVerdict(LAYER_STRUCTURAL, PASS, "evidence", {"k": 1})
    payload = verdict.to_dict()
    assert set(payload) == {"layer", "status", "evidence_ref", "detail"}
    assert verdict.measured is True


def test_stack_aggregates_pass_and_surfaces_unmeasured_layers():
    data = _sample_doc()
    start, end = _target_range(data)
    verdict = validate_document(data, _compliant_edit(data), start=start, end=end)
    assert isinstance(verdict, ValidationVerdict)
    assert verdict.layer(LAYER_STRUCTURAL).status == PASS
    assert verdict.layer(LAYER_INVARIANTS).status == PASS
    assert verdict.layer(LAYER_FORMAT).status == PASS
    # no engine / no judge on this box ⇒ both unmeasured, and that is surfaced
    assert verdict.layer(LAYER_RENDER).status is None
    assert verdict.layer(LAYER_SEMANTIC).status is None
    assert set(verdict.unmeasured) >= {LAYER_RENDER, LAYER_SEMANTIC}
    assert verdict.overall == PASS
    assert all(v.status != PASS for v in verdict.layers.values() if v.status is None)


def test_stack_disable_hook_turns_a_measured_layer_into_none():
    data = _sample_doc()
    start, end = _target_range(data)
    broken = _mutate_text(data, b"1,000.00", b"9,999.99")

    enabled = validate_document(data, broken, start=start, end=end)
    assert enabled.layer(LAYER_INVARIANTS).status == FAIL
    assert enabled.overall == FAIL

    disabled = validate_document(
        data,
        broken,
        start=start,
        end=end,
        config=ValidationConfig(disabled_layers=frozenset({LAYER_INVARIANTS})),
    )
    assert disabled.layer(LAYER_INVARIANTS).status is None
    assert LAYER_INVARIANTS in disabled.unmeasured


def test_stack_L5_low_score_is_needs_review_never_pass(monkeypatch):
    from forgeflow.documents.validation import semantic as semantic_mod

    semantic_mod.set_default_judge(lambda req: _score(0.05, "default-low"))
    try:
        data = _sample_doc()
        start, end = _target_range(data)
        verdict = validate_document(data, _compliant_edit(data), start=start, end=end)
        assert verdict.layer(LAYER_SEMANTIC).status == NEEDS_REVIEW
        assert verdict.overall == NEEDS_REVIEW
        assert verdict.overall != PASS
    finally:
        semantic_mod.set_default_judge(None)


def test_stack_turns_an_unreadable_after_into_an_honest_fail():
    data = _sample_doc()
    start, end = _target_range(data)
    verdict = validate_document(data, b"not a docx", start=start, end=end)
    assert verdict.layer(LAYER_INVARIANTS).status == FAIL
    assert verdict.layer(LAYER_FORMAT).status == FAIL
    assert verdict.overall == FAIL


def test_stack_resolves_a_selector_into_a_range():
    data = _sample_doc()
    verdict = validate_document(data, _compliant_edit(data), selector="标题含 报告标题")
    assert verdict.layer(LAYER_INVARIANTS).status == PASS
    assert "located" in verdict.notes or "range:" in verdict.notes


# --------------------------------------------------------------------------- #
# 7. F1 — L4 page-count change ⇒ stack ``warn`` (never L4 fail; None ⇒ no warn) #
# --------------------------------------------------------------------------- #
def _page_render_kwargs(
    before: bytes, after: bytes, pages: dict[bytes, int | None]
) -> dict:
    """A render fixture with **identical rasters** (diff 0) but **distinct PDFs**,
    so the page counter can tell the two sides apart.

    ``pages`` maps the PDF bytes to a page count (or ``None`` = unmeasured).
    """

    def converter(src: bytes, engine: str) -> bytes:
        return b"PDF-BEFORE" if src == before else b"PDF-AFTER"

    def rasterizer(pdf: bytes, tool: str, dpi: int):
        return [[(10, 10, 10)] * 3 for _ in range(3)]  # identical for both sides

    def page_counter(pdf: bytes) -> int | None:
        return pages.get(pdf)

    return {
        "render_engine": "fake-engine",
        "render_converter": converter,
        "rasterizer": rasterizer,
        "page_counter": page_counter,
    }


def test_F1_page_count_change_warns_without_failing_L4():
    before = _sample_doc()
    start, end = _target_range(before)
    after = _compliant_edit(before)
    verdict = validate_document(
        before,
        after,
        start=start,
        end=end,
        **_page_render_kwargs(before, after, {b"PDF-BEFORE": 1, b"PDF-AFTER": 2}),
    )
    assert verdict.layer(LAYER_RENDER).status != FAIL
    assert verdict.layer(LAYER_RENDER).detail["page_delta"] == 1
    assert verdict.layer(LAYER_RENDER).detail["pages_before"] == 1
    assert verdict.layer(LAYER_RENDER).detail["pages_after"] == 2
    assert verdict.overall == "warn"
    assert "页数变化" in verdict.notes


def test_F1_page_count_unchanged_is_a_plain_pass():
    before = _sample_doc()
    start, end = _target_range(before)
    after = _compliant_edit(before)
    verdict = validate_document(
        before,
        after,
        start=start,
        end=end,
        **_page_render_kwargs(before, after, {b"PDF-BEFORE": 1, b"PDF-AFTER": 1}),
    )
    assert verdict.layer(LAYER_RENDER).detail["page_delta"] == 0
    assert verdict.overall == PASS
    assert verdict.overall != "warn"


def test_F1_unmeasured_page_count_never_warns():
    before = _sample_doc()
    start, end = _target_range(before)
    after = _compliant_edit(before)
    verdict = validate_document(
        before,
        after,
        start=start,
        end=end,
        **_page_render_kwargs(before, after, {}),  # page_counter ⇒ None both sides
    )
    assert verdict.layer(LAYER_RENDER).detail["page_delta"] is None
    assert verdict.overall != "warn"


def test_counterfactual_F1_removing_the_page_warn_branch_turns_the_warn_case_red(
    monkeypatch,
):
    import forgeflow.documents.validation.stack as stack_mod

    before = _sample_doc()
    start, end = _target_range(before)
    after = _compliant_edit(before)
    kwargs = _page_render_kwargs(before, after, {b"PDF-BEFORE": 1, b"PDF-AFTER": 2})

    def _assert_page_warn(verdict) -> None:
        assert verdict.overall == "warn", f"页数变化未告警（overall={verdict.overall}）"

    _assert_page_warn(
        validate_document(before, after, start=start, end=end, **kwargs)
    )  # green

    # counterfactual: neutralise the page-warn branch (remove its trigger) — the
    # same assertion must go red, proving the branch is load-bearing.
    real_aggregate = stack_mod._aggregate

    def _without_page_branch(layers, config):
        layers[stack_mod.LAYER_RENDER].detail.pop("page_delta", None)
        return real_aggregate(layers, config)

    monkeypatch.setattr(stack_mod, "_aggregate", _without_page_branch)
    with pytest.raises(AssertionError) as excinfo:
        _assert_page_warn(validate_document(before, after, start=start, end=end, **kwargs))
    assert "页数变化未告警" in str(excinfo.value)


# --------------------------------------------------------------------------- #
# 8. F2 / F3 — evidence honesty (capability boundary + no 0/0 misread)          #
# --------------------------------------------------------------------------- #
def test_F2_L1_evidence_never_claims_xsd_schema_validation():
    verdict = structural_layer.structural_verdict(_sample_doc())
    assert verdict.status == PASS
    assert "package-integrity:" in verdict.evidence_ref
    assert "schema" not in verdict.evidence_ref.lower()

    broken = _drop_part(_sample_doc(), "[Content_Types].xml")
    failed = structural_layer.structural_verdict(broken)
    assert failed.status == FAIL
    assert "package-integrity:FAIL[" in failed.evidence_ref
    assert "schema" not in failed.evidence_ref.lower()


def test_F3_l2_l3_empty_denominator_evidence_is_not_zero_over_zero():
    data = _sample_doc()
    # L3: a no-op edit changes nothing in range ⇒ in_range_edited == 0.
    l3 = format_layer.format_verdict(data, data, start=2, end=3)
    assert l3.status == PASS
    assert "style-inherited=n/a (no in-range edit)" in l3.evidence_ref
    assert "0/0" not in l3.evidence_ref
    assert l3.detail["in_range_edited"] == 0  # the numeric field is unchanged

    # L2: a range with no explicit invariant ⇒ compared == 0.
    l2 = invariants_layer.invariants_verdict(data, data, start=0, end=1)
    assert l2.status == PASS
    assert "survived=n/a (no explicit invariant in range)" in l2.evidence_ref
    assert "0/0" not in l2.evidence_ref
    assert l2.detail["compared"] == 0  # the numeric field is unchanged
