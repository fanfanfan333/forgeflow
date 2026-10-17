import io

import pytest
from docx import Document

from forgeflow.documents.content_read import read_content


def test_same_field_in_other_task_cannot_replace_requested_table_goal():
    doc = Document()
    doc.add_heading("工程任务", level=1)
    doc.add_heading("附录 T15", level=2)
    doc.add_paragraph("T15")
    for task, goal in [
        ("T15", "必须保留的较长的发布联锁目标描述" * 20),
        ("T23", "验证目标"),
        ("T150", "另一个编号"),
    ]:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = task + " · 工作项"
        table.cell(1, 0).text = "目标"
        table.cell(1, 1).text = goal
    stream = io.BytesIO()
    doc.save(stream)
    result = read_content(stream.getvalue(), "只读定位 T15 的目标字段", limit=12)
    assert "必须保留的较长的发布联锁目标描述" in result["text"]
    assert result["text"].splitlines()[0].startswith("[table:0/row:1]")
    assert "T15 · 工作项" in result["text"]
    assert "[table:0/row:1]" in result["text"]
    scoped = read_content(stream.getvalue(), task_id="T15", field="目标")
    assert scoped["selection"] == {"task_id": "T15", "field": "目标", "matches": 1}
    assert scoped["next_offset"] is None and not scoped["truncated"]
    assert "T23" not in scoped["text"] and "T150" not in scoped["text"]
    missing = read_content(stream.getvalue(), task_id="T15", field="不存在")
    assert not missing["text"] and missing["selection"]["matches"] == 0
    with pytest.raises(ValueError):
        read_content(stream.getvalue(), task_id="T15")


def test_scoped_field_preserves_long_content_pagination():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "T15 · 工作项"
    table.cell(1, 0).text = "目标"
    table.cell(1, 1).text = "完整内容" * 1000
    stream = io.BytesIO()
    doc.save(stream)
    first = read_content(stream.getvalue(), task_id="T15", field="目标")
    assert first["next_char_offset"] > 0 and first["truncated"]
    second = read_content(
        stream.getvalue(), task_id="T15", field="目标", char_offset=first["next_char_offset"]
    )
    assert second["text"] and second["selection"]["matches"] == 1
