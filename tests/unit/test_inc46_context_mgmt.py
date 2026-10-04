"""INC46 T29 — 上下文管理与压缩（Context Management & Compaction）验收模块。

任务书 T29 的验收命令指向本文件：``pytest tests/unit/test_inc46_context_mgmt.py``。

覆盖（阳性 / 阴性 / 反事实 / 红线 4 + 16）
------------------------------------------
1. **按章节可寻址加载（阳性）** — 只加载**目标区间 ± 邻接 + 文档大纲**；区间外的章节正文
   一律**不**进上下文。
2. **阳性探针（预算）** — 长文档只改第 3 部分：加载 token / 全文 token ≤ 15%，且结果正确。
3. **token 记账** — 未测量 ⇒ ``None``（红线 4），**绝不用 0 冒充**；「已测量的 0」与「未测量」
   被区分开；``run_steps`` 记账加性且保留 ``None``。
4. **会话白名单（阴性）** — 压缩输出缺 ``InvariantSet`` ⇒ **硬断言失败**（必含字段检查）。
5. **跨块边界（阴性）** — 目标区间跨章节边界 ⇒ 仍**完整**加载（不截断）。
6. **回写位置（阴性）** — 分块后**错位回写** ⇒ 被检出（:class:`WriteBackError`）。
7. **反事实（真跑）** — 去掉「必保留字段」白名单 ⇒ 阴性用例**必须转红**。
8. **加性接缝** — ``build_context`` 仅在显式传入 ``document`` 时才走文档加载路径；默认路径
   逐键不变。

所有断言都驱动 ``forgeflow.context`` 的**真实函数**，并且用 ``python-docx`` **合成**真实
DOCX 夹具（无任何未脱敏真实客户数据，红线 16 合规）。
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from forgeflow.context import (
    CONVERSATION_PROTECTED_TAGS,
    TAG_CLARIFICATION,
    TAG_HITL_PENDING,
    TAG_INVARIANTS,
    TAG_USER_INSTRUCTION,
    TAG_VERSION_CHAIN,
    WhitelistConfig,
    assert_outside_unchanged,
    assert_required_preserved,
    compact_plan,
    load_sections,
    run_step_token_payload,
    verify_writeback,
)
from forgeflow.context.budget_accounting import account_chunks
from forgeflow.context.compactor import CompactionStrategy, compact
from forgeflow.context.section_loader import WriteBackError
from forgeflow.context.token_accounting import StepTokenAccounting, account_step
from forgeflow.documents import inspect_docx
from forgeflow.experience.token_budget import estimate_tokens

# pytest asyncio 通过 pyproject 的 ``asyncio_mode = auto`` 自动处理 async 用例，
# 无需模块级 pytestmark（本文件同时含同步与异步用例）。


# --------------------------------------------------------------------------- #
# Synthetic DOCX fixtures (python-docx; deterministic)                          #
# --------------------------------------------------------------------------- #
def _save(document: Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_doc(
    *, section_count: int = 40, paras_per_section: int = 8, sub_on: int | None = None
) -> bytes:
    """A long, sectioned DOCX: ``section_count`` level-1 parts, each with body paragraphs.

    ``sub_on`` (a 1-based part number) additionally gets a level-2 sub-heading plus body,
    so that part's top-level range **crosses an internal heading boundary**.
    """
    document = Document()
    document.add_paragraph("封面说明：本文件为合成夹具。")  # preamble (outside all sections)
    for i in range(1, section_count + 1):
        document.add_heading(f"第{i}部分 小节标题{i}", level=1)
        for _ in range(paras_per_section):
            document.add_paragraph(f"第{i}节正文内容，用于填充篇幅并构成可寻址的章节。" * 2)
        if sub_on is not None and i == sub_on:
            document.add_heading(f"细节{i} 子标题", level=2)
            document.add_paragraph(f"子标题{i}下的额外正文段落。" * 4)
    return _save(document)


def _protected_tokens(plan) -> int:
    """The (estimated) token total of a plan's **whitelist-protected** chunks.

    Used as the compaction budget: exactly enough to keep every protected chunk, so a
    legal compaction must drop the non-protected ones (and must never raise).
    """
    return sum(
        estimate_tokens(c.text)
        for c in plan.all_chunks()
        if plan.whitelist.is_protected(c)
    )


# --------------------------------------------------------------------------- #
# 1. 按章节可寻址加载 — 只加载目标 ± 邻接 + 大纲                                  #
# --------------------------------------------------------------------------- #
def test_loads_only_target_adjacency_and_outline() -> None:
    data = _build_doc(section_count=40)
    plan = load_sections(data, selector="第三部分", adjacency=1)

    assert plan.located, plan.notes
    assert plan.target_chunk is not None
    # 恰好 3 个原文切片块：目标 + 前邻接 + 后邻接。
    assert len(plan.loaded_chunks) == 3
    loaded_text = plan.loaded_text()
    assert "小节标题3" in loaded_text            # 目标
    assert "小节标题2" in loaded_text            # 前邻接
    assert "小节标题4" in loaded_text            # 后邻接
    # 区间外的章节正文**不**进上下文。
    assert "小节标题20" not in loaded_text
    assert "小节标题39" not in loaded_text
    # 大纲完整（全部 40 个一级标题都在大纲里）。
    assert len(plan.outline) == 40
    assert plan.outline_chunk is not None

    # 边界可复算：每个加载块都可由原文切片复原。
    for chunk in plan.loaded_chunks:
        assert plan.source[chunk.start:chunk.end] == chunk.text


def test_adjacency_zero_loads_target_only() -> None:
    data = _build_doc(section_count=40)
    plan = load_sections(data, selector="第三部分", adjacency=0)
    assert len(plan.loaded_chunks) == 1
    assert plan.target_chunk is not None
    assert plan.loaded_chunks[0].chunk_id == plan.target_chunk.chunk_id


def test_selector_ambiguity_does_not_fabricate_target() -> None:
    """选择子未找到 ⇒ 只加载大纲，绝不猜测目标。"""
    data = _build_doc(section_count=5)
    plan = load_sections(data, selector="第九部分", adjacency=1)
    assert not plan.located
    assert plan.target_chunk is None
    assert plan.loaded_chunks == []
    assert "未找到" in plan.notes or "不猜测" in plan.notes


# --------------------------------------------------------------------------- #
# 2. 阳性探针 — 长文档只改第 3 部分：加载 token / 全文 ≤ 15%                      #
# --------------------------------------------------------------------------- #
def test_positive_probe_budget_ratio_le_15_percent() -> None:
    data = _build_doc(section_count=100, paras_per_section=8, sub_on=3)
    plan = load_sections(data, selector="第三部分", adjacency=1)

    full = plan.full_document_tokens
    ratio = plan.budget_ratio()
    assert full is not None and full > 0
    assert ratio is not None
    print(
        "[T29 positive probe] "
        f"full_document_tokens={full} loaded_tokens={plan.loaded_tokens} "
        f"outline_tokens={plan.outline_tokens} total_loaded={plan.total_loaded_tokens} "
        f"ratio={ratio:.4f} (<= 0.15)"
    )
    # 承重断言：加载 token ≤ 全文的 15%（目标区间 + 邻接 + 大纲）。
    assert ratio <= 0.15, f"加载比例 {ratio:.4f} 超过阈值 0.15"
    # 结果正确：目标（含子标题）真的被完整加载。
    assert plan.target_chunk is not None
    assert "小节标题3" in plan.loaded_text()
    assert "细节3 子标题" in plan.loaded_text()


# --------------------------------------------------------------------------- #
# 3. token 记账 — 未测量 ⇒ None（红线 4），禁止写 0                              #
# --------------------------------------------------------------------------- #
def test_account_step_unmeasured_is_none_never_zero() -> None:
    acc = account_step()  # 什么都未记账
    assert acc.prompt_tokens is None
    assert acc.completion_tokens is None
    assert acc.context_tokens is None
    assert acc.measured() is False
    # 明确禁止 0 冒充。
    assert acc.prompt_tokens != 0 and acc.completion_tokens != 0
    # run_steps payload 加性且保留 None。
    payload = run_step_token_payload(acc, base={"step": 1})
    assert payload["step"] == 1
    rec = payload["token_accounting"]
    assert rec["prompt_tokens"] is None and rec["completion_tokens"] is None
    assert rec["measured"] is False


def test_account_step_measured_zero_is_distinct_from_unmeasured() -> None:
    # 已测量的 0（显式传入 0）保持 0，不冒充未测量。
    measured_zero = account_step(prompt=0)
    assert measured_zero.prompt_tokens == 0
    assert measured_zero.measured() is True
    # 文本输入真的被估算为正整数。
    measured = account_step(context="一些上下文文本用于估算 token")
    assert isinstance(measured.context_tokens, int) and measured.context_tokens > 0


def test_account_step_rejects_bool() -> None:
    with pytest.raises(TypeError):
        account_step(prompt=True)  # type: ignore[arg-type]


def test_run_step_payload_none_accounting_is_unmeasured() -> None:
    payload = run_step_token_payload(None)
    rec = payload["token_accounting"]
    assert rec["prompt_tokens"] is None and rec["completion_tokens"] is None
    assert rec["context_tokens"] is None
    assert rec["measured"] is False


def test_loader_ledger_unmeasured_is_none_not_zero() -> None:
    data = _build_doc(section_count=10)
    # counter=None ⇒ 全部块未测量：字段必须是 None，且**明确不是 0**。
    plan = load_sections(data, selector="第三部分", adjacency=1, counter=None)
    assert plan.full_document_tokens is None
    assert plan.budget_ratio() is None
    assert plan.ledger is not None
    assert all(a.estimated_tokens is None for a in plan.ledger.accounts)
    for a in plan.ledger.accounts:
        assert a.estimated_tokens != 0  # None ≈ 未测量，禁止 0 冒充
    assert len(plan.ledger.unmeasured_accounts()) == len(plan.ledger.accounts)


# --------------------------------------------------------------------------- #
# 4. 会话白名单（阴性）— 压缩输出缺 InvariantSet ⇒ 硬失败                        #
# --------------------------------------------------------------------------- #
def test_compaction_preserves_invariantset_and_required_fields() -> None:
    data = _build_doc(section_count=40, sub_on=3)
    plan = load_sections(data, selector="第三部分", adjacency=1)
    assert plan.invariants_chunk is not None
    required = [plan.invariants_chunk]
    if plan.target_chunk is not None:
        required.append(plan.target_chunk)

    # 预算 = 受保护内容本身（刚好放得下白名单，逼出对非白名单块的压缩）。
    budget = _protected_tokens(plan)
    before = {c.chunk_id: c.sha256 for c in required}
    result = compact_plan(plan, budget_tokens=budget, strategy=CompactionStrategy.TRUNCATE)

    kept = {c.chunk_id: c.sha256 for c in result.kept}
    assert_required_preserved(result, required)  # 硬检查
    for chunk_id, digest in before.items():
        assert kept.get(chunk_id) == digest, f"白名单块 {chunk_id} 未被逐字节保留"
    # 确有非白名单块被压缩/丢弃（证明真的压了，不是恒等通过）。
    assert result.dropped or result.compacted
    print(
        "[T29 whitelist before->after] "
        + " | ".join(f"{cid}:{before[cid][:12]}->{kept.get(cid, 'MISSING')[:12]}"
                     for cid in before)
    )


def test_missing_invariantset_fails_hard() -> None:
    """压缩输出缺 InvariantSet ⇒ 断言**硬失败**（必含字段检查）。"""
    data = _build_doc(section_count=40, sub_on=3)
    plan = load_sections(data, selector="第三部分", adjacency=1)
    assert plan.invariants_chunk is not None
    required = [plan.invariants_chunk]

    # 反事实：把白名单清空 + 预算小到连 InvariantSet 都放不下 ⇒ 它被当作普通块丢弃。
    emptied = WhitelistConfig(
        protected_tags=frozenset(),
        protected_kinds=frozenset(),
        protected_chunk_ids=frozenset(),
    )
    unguarded = compact(
        plan.all_chunks(), budget_tokens=1, policy=emptied,
        strategy=CompactionStrategy.DROP_LOW_PRIORITY,
    )
    with pytest.raises(AssertionError) as exc:
        assert_required_preserved(unguarded, required)
    assert "压缩未保留关键状态" in str(exc.value)
    # 可观测证据：无白名单时，InvariantSet 块确实不在保留集里。
    kept_ids = {c.chunk_id for c in unguarded.kept}
    assert plan.invariants_chunk.chunk_id not in kept_ids

    # 对照：同一断言在**有**白名单时通过（证明断言可满足，不是恒真失败）。
    guarded = compact_plan(plan, budget_tokens=_protected_tokens(plan),
                           strategy=CompactionStrategy.TRUNCATE)
    assert_required_preserved(guarded, required)



def test_conversation_whitelist_covers_required_categories_without_fabrication() -> None:
    """白名单覆盖任务书五类必保留；缺席的版本链 ID **只是缺席**，绝不臆造。"""
    for tag in (
        TAG_INVARIANTS,       # InvariantSet
        TAG_HITL_PENDING,     # 待批 pending_id
        TAG_VERSION_CHAIN,    # 版本链 ID（T25）
        TAG_USER_INSTRUCTION, # 用户明确偏好
        TAG_CLARIFICATION,    # 未决澄清
    ):
        assert tag in CONVERSATION_PROTECTED_TAGS, tag
    data = _build_doc(section_count=10)
    plan = load_sections(data, selector="第三部分", adjacency=1)
    # 本夹具没有版本链 / 未决澄清块 ⇒ 白名单里就没有对应块（不臆造一个）。
    assert all(TAG_VERSION_CHAIN not in c.tags for c in plan.all_chunks())


# --------------------------------------------------------------------------- #
# 5. 跨块边界（阴性）— 目标区间跨章节边界 ⇒ 仍完整                              #
# --------------------------------------------------------------------------- #
def test_target_range_crossing_chunk_boundary_is_complete() -> None:
    data = _build_doc(section_count=40, sub_on=3)
    plan = load_sections(data, selector="第三部分", adjacency=0)

    assert plan.target_chunk is not None
    ts, te = plan.target_chunk.start, plan.target_chunk.end
    # 目标块恰好是原文 [ts, te) 的**完整**切片 —— 跨了内部子标题边界也不截断。
    assert plan.source[ts:te] == plan.target_chunk.text
    text = plan.target_chunk.text
    assert "小节标题3" in text          # 章标题
    assert "细节3 子标题" in text        # 内部子标题（跨边界）
    assert f"子标题3下的额外正文段落。" in text  # 子标题下的正文（边界之后）
    # 子标题下的段落**完整**保留（没有被截断在边界处）。
    assert text.count("子标题3下的额外正文段落。") == 4


def test_explicit_range_spanning_two_sections_is_complete() -> None:
    data = _build_doc(section_count=20)
    structure = inspect_docx(data)
    # 取第 3、4 两个一级标题的段落区间（跨章节边界）。
    third = structure.headings[2]
    fourth = structure.headings[3]
    lo = int(third["index"])
    hi = int(fourth["section_end_index"])  # 覆盖第 3 部分 + 第 4 部分
    plan = load_sections(data, target_range=(lo, hi), adjacency=0)

    assert plan.target_chunk is not None
    assert "小节标题3" in plan.target_chunk.text
    assert "小节标题4" in plan.target_chunk.text  # 跨到下一章节仍完整
    assert plan.source[plan.target_chunk.start:plan.target_chunk.end] == plan.target_chunk.text


# --------------------------------------------------------------------------- #
# 6. 回写位置（阴性）— 分块后错位回写 ⇒ 被检出                                    #
# --------------------------------------------------------------------------- #
def test_misplaced_writeback_is_detected() -> None:
    data = _build_doc(section_count=20)
    plan = load_sections(data, selector="第三部分", adjacency=1)
    assert plan.target_chunk is not None
    chunk = plan.target_chunk

    # 正确位置：回写落在原始 start ⇒ 成功，且区间外零变化。
    new_source = verify_writeback(
        plan.source, chunk, "改写后的第三部分正文", at_offset=chunk.start
    )
    assert_outside_unchanged(plan.source, new_source, chunk)
    assert "改写后的第三部分正文" in new_source

    # 错位回写：偏移不是原始 start ⇒ **被检出**（抛错）。
    for wrong in (chunk.start + 1, chunk.start - 1, chunk.end):
        with pytest.raises(WriteBackError) as exc:
            verify_writeback(plan.source, chunk, "错位写入", at_offset=wrong)
        assert "回写位置错误" in str(exc.value)


def test_writeback_on_tampered_source_is_detected() -> None:
    """若原文被篡改致块边界不可复算，回写同样被检出（拒绝写错位置）。"""
    data = _build_doc(section_count=10)
    plan = load_sections(data, selector="第三部分", adjacency=0)
    assert plan.target_chunk is not None
    chunk = plan.target_chunk
    tampered = plan.source[: chunk.start] + "X" + plan.source[chunk.start + 1 :]
    with pytest.raises(WriteBackError):
        verify_writeback(tampered, chunk, "新文本", at_offset=chunk.start)


# --------------------------------------------------------------------------- #
# 7. 反事实（真跑）— 去掉必保留白名单 ⇒ 阴性用例转红                              #
# --------------------------------------------------------------------------- #
def test_counterfactual_removing_whitelist_turns_preservation_red() -> None:
    data = _build_doc(section_count=40, sub_on=3)
    plan = load_sections(data, selector="第三部分", adjacency=1)
    assert plan.invariants_chunk is not None
    required = [plan.invariants_chunk]

    # (a) 默认白名单：保利断言通过。
    guarded = compact_plan(plan, budget_tokens=_protected_tokens(plan))
    assert_required_preserved(guarded, required)  # must NOT raise

    # (b) 反事实（真跑）：清空白名单 + 预算极小 ⇒ 必保留块被丢 ⇒ 断言转红。
    unguarded = compact(
        plan.all_chunks(),
        budget_tokens=1,
        policy=WhitelistConfig(
            protected_tags=frozenset(),
            protected_kinds=frozenset(),
            protected_chunk_ids=frozenset(),
        ),
        strategy=CompactionStrategy.DROP_LOW_PRIORITY,
    )
    with pytest.raises(AssertionError):
        assert_required_preserved(unguarded, required)



# --------------------------------------------------------------------------- #
# 8. 加性接缝 — build_context 仅在显式请求时走文档加载路径                        #
# --------------------------------------------------------------------------- #
async def test_build_context_document_path_is_opt_in_and_additive(force_memory_backend):
    from forgeflow.experience.context_builder import build_context
    from forgeflow.experience.memory_store import clear_memory_entries

    clear_memory_entries()
    data = _build_doc(section_count=40, sub_on=3)

    # 默认路径（不传 document）：逐键不变，无 document_load / token_accounting。
    base = await build_context("t-ctx-eng42", "无关意图", budget_tokens=400)
    d = base.to_dict()
    assert base.document_load is None and base.token_accounting is None
    assert "document_load" not in d and "token_accounting" not in d

    # 显式请求文档加载 ⇒ 追加文档段 + 记账，且 prompt/completion 未测量 ⇒ None。
    bundle = await build_context(
        "t-ctx-eng42",
        "把第三部分改得更正式",
        budget_tokens=400,
        document=data,
        document_selector="第三部分",
        document_adjacency=1,
    )
    sources = {s["source"] for s in bundle.sections}
    assert "document" in sources
    assert bundle.document_load is not None
    assert bundle.document_load["located"] is True
    assert bundle.document_load["budget_ratio"] is not None
    assert bundle.document_load["budget_ratio"] <= 0.15
    assert bundle.token_accounting is not None
    rec = bundle.token_accounting["token_accounting"]
    assert rec["prompt_tokens"] is None and rec["completion_tokens"] is None
    assert isinstance(rec["context_tokens"], int) and rec["context_tokens"] > 0


async def test_build_context_document_bad_bytes_degrades_honestly(force_memory_backend):
    """不可读的 document ⇒ 诚实降级（记录错误、不臆造内容、不破坏 bundle）。"""
    from forgeflow.experience.context_builder import build_context

    bundle = await build_context(
        "t-ctx-eng42", "意图", budget_tokens=200, document=b"not-a-docx",
        document_selector="第三部分",
    )
    assert bundle.document_load is not None
    assert bundle.document_load["located"] is False
    assert "error" in bundle.document_load
    assert all(s["source"] != "document" for s in bundle.sections)


# --------------------------------------------------------------------------- #
# 9. 记账快照（证据）                                                           #
# --------------------------------------------------------------------------- #
def test_section_load_ledger_snapshot_evidence() -> None:
    data = _build_doc(section_count=40, sub_on=3)
    plan = load_sections(data, selector="第三部分", adjacency=1)
    snap = plan.ledger.snapshot()
    print(
        "[T29 ledger snapshot] "
        f"budget={snap['budget_tokens']} measured={snap['measured_tokens']} "
        f"unmeasured={snap['unmeasured_count']} accounts={snap['account_count']} "
        f"protected={[a.chunk_id for a in plan.ledger.protected_chunks()]}"
    )
    # 白名单块（InvariantSet / 编辑目标）在账本里如实标为 protected。
    protected_ids = {a.chunk_id for a in plan.ledger.protected_chunks()}
    assert plan.invariants_chunk is not None
    assert plan.invariants_chunk.chunk_id in protected_ids
    # 已测量口径可复算：总量 = 各已测量块之和。
    assert snap["measured_tokens"] == sum(
        a["estimated_tokens"] for a in snap["accounts"]
        if a["estimated_tokens"] is not None
    )
    # 空输入不产生块（绝不伪造空块）。
    empty_plan = load_sections(_build_doc(section_count=0), selector="第三部分")
    assert empty_plan.located is False
    assert empty_plan.loaded_chunks == []


def test_account_chunks_unmeasured_none_not_zero() -> None:
    """底层账本对未测量同样坚持 None（红线 4 的单元级证据）。"""
    from forgeflow.context.chunker import segment_paragraphs

    chunks = segment_paragraphs("第一段。\n\n第二段。")
    ledger = account_chunks(chunks, budget_tokens=100, counter=None)
    assert all(a.estimated_tokens is None for a in ledger.accounts)
    assert ledger.measured_tokens() == 0
    assert len(ledger.unmeasured_accounts()) == len(chunks)


def test_step_token_accounting_is_json_safe() -> None:
    acc = StepTokenAccounting(prompt_tokens=7, completion_tokens=None, context_tokens=120)
    d = acc.to_dict()
    assert d == {
        "prompt_tokens": 7,
        "completion_tokens": None,
        "context_tokens": 120,
        "measured": True,
    }
    assert estimate_tokens("abcd") == 1  # 复用既有估算器（非精确 tokenizer）
