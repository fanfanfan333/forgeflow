"""INC44 T05 — honesty discipline audit for the PPTX + text/code planes.

Design §7 T05 criterion ② / §8 诚实口径: *假计数扫描 / ``0`` 兜底扫描 / 未测量⇒``None``
/ 缺 extra 诚实降级 / 耗尽诚实报错，全绿*.

Every pin here is a **falsifiable** honesty claim, not a smoke test:

  * **tri-state** — an unmeasured verification dimension is ``None`` (never a fake
    ``True``/``False``) for both the PPTX and the text plane;
  * **no fabricated counts** — the delivered 「修改 N 处」 figures are, in source,
    *either* the recomputed ``diff`` value *or* an honest ``0`` on the failure
    path — never a hardcoded positive;
  * **exhaustion is honest** — an unsatisfiable constraint ends in ``ok=False``
    with ``modified == 0`` and **no** ``artifact_ref`` (no fake deliverable);
  * **missing extra degrades, never 5xx** — a ``.pptx`` whose ``python-pptx`` extra
    is absent (or whose bytes are unparseable) lands ``metadata_only`` with a
    **verbatim** reason, and never reports ``parsed``.

Citation discipline: ``file.py::symbol`` anchors, never ``file:line``.
"""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

from forgeflow.documents import (
    compute_pptx_diff,
    verify_pptx,
    verify_textfile,
)
from forgeflow.documents.pptx_edit import apply_edits as apply_pptx_edits

REPO_ROOT = Path(__file__).resolve().parents[2]
TOOL_HANDLERS = REPO_ROOT / "forgeflow" / "runtime" / "tool_handlers.py"

# NOTE: no module-level ``pytest.mark.asyncio`` here on purpose — this module mixes
# sync honesty pins (source scans / pure-function checks) with ``async`` handler
# tests. ``pyproject.toml`` sets ``asyncio_mode = "auto"``, so the async tests are
# collected without an explicit mark, and the sync ones stay unmarked (a blanket
# mark would raise a PytestWarning on each sync test).


# --------------------------------------------------------------------------- #
# Builders                                                                     #
# --------------------------------------------------------------------------- #
def _pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "季度汇报 2026"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "营收 40 万元"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _ctx() -> SimpleNamespace:
    return SimpleNamespace(
        run_id="run-inc44-honesty",
        step_id="run-inc44-honesty:0:0",
        tenant_id="t-inc44",
        user_id="u-inc44",
        role="manager",
        args={},
    )


@pytest.fixture
def tmp_doc_store(tmp_path, monkeypatch):
    from forgeflow.documents import store as store_mod

    real = store_mod.DocArtifactStore

    def _factory(root=None):
        return real(root=root if root is not None else tmp_path / "docstore")

    import forgeflow.documents as docs_pkg

    monkeypatch.setattr(store_mod, "DocArtifactStore", _factory)
    monkeypatch.setattr(docs_pkg, "DocArtifactStore", _factory)
    return tmp_path


# --------------------------------------------------------------------------- #
# 1. Tri-state: unmeasured ⇒ None (never a fake True/False)                    #
# --------------------------------------------------------------------------- #
async def test_verify_pptx_unmeasured_dimensions_are_none():
    report = verify_pptx(_pptx_bytes(), {})
    assert report.openable is True
    assert report.structure_ok is None
    assert report.data_ok is None
    assert report.requirement_ok is None


async def test_verify_textfile_unmeasured_dimensions_are_none():
    report = verify_textfile(b"line1\nline2\n", {})
    assert report.openable is True
    assert report.structure_ok is None
    assert report.data_ok is None
    assert report.requirement_ok is None


async def test_verify_pptx_unopenable_bytes_are_honest():
    report = verify_pptx(b"not a presentation", {"max_chars": 10})
    assert report.openable is False
    assert report.structure_ok is None
    assert report.data_ok is None
    assert report.requirement_ok is None


# --------------------------------------------------------------------------- #
# 2. No fabricated counts — the figures are ``diff``-derived or an honest 0     #
# --------------------------------------------------------------------------- #
def test_delivered_counts_are_never_hardcoded_positives():
    """Every ``modified`` / ``added`` / ``removed`` / ``numeric_changes`` in the
    document + text handlers is **either** the recomputed ``diff.<field>`` **or**
    an honest ``0`` (the failure path) — never a fabricated positive literal.
    """
    source = TOOL_HANDLERS.read_text(encoding="utf-8")
    for field in ("modified", "added", "removed", "numeric_changes"):
        values = re.findall(rf'"{field}":\s*([^,\n]+)', source)
        assert values, f"{field} not found — scan would be vacuous"
        for raw in values:
            token = raw.strip()
            assert token in (f"diff.{field}", "0"), (field, token)


async def test_pptx_success_counts_equal_the_recomputed_diff(tmp_path, tmp_doc_store):
    """The reported figures are exactly the pure ``compute_diff`` result."""
    from forgeflow.runtime.tool_handlers import document_edit

    old = _pptx_bytes()
    # Recompute the diff the same way the handler does.
    new, _changes = apply_pptx_edits(
        old, [{"op": "replace_text", "match": "40", "replace": "42"}]
    )
    expected = compute_pptx_diff(old, new)

    target = tmp_path / "deck.pptx"
    target.write_bytes(old)
    result = await document_edit(
        {
            "paths": [str(target)],
            "edits": [{"op": "replace_text", "match": "40", "replace": "42"}],
        },
        _ctx(),
    )
    assert result["ok"] is True
    assert result["modified"] == expected.modified
    assert result["added"] == expected.added
    assert result["removed"] == expected.removed
    assert result["numeric_changes"] == expected.numeric_changes


# --------------------------------------------------------------------------- #
# 3. Exhaustion is honest — ok=False, count 0, NO artifact_ref                 #
# --------------------------------------------------------------------------- #
async def test_pptx_edit_exhaustion_is_honest(tmp_path, tmp_doc_store):
    from forgeflow.runtime.tool_handlers import document_edit

    target = tmp_path / "deck.pptx"
    target.write_bytes(_pptx_bytes())
    result = await document_edit(
        {
            "paths": [str(target)],
            "edits": [{"op": "replace_text", "match": "40", "replace": "42"}],
            "max_chars": 1,  # unsatisfiable ⇒ bounded retry then honest failure
        },
        _ctx(),
    )
    assert result["ok"] is False
    assert "校验未通过" in result["reason"]
    assert result["modified"] == 0
    assert result["validation"]["requirement_ok"] is False
    assert "artifact_ref" not in result


async def test_textfile_edit_exhaustion_is_honest(tmp_path, tmp_doc_store):
    from forgeflow.runtime.tool_handlers import textfile_edit

    target = tmp_path / "app.py"
    target.write_bytes(b"def add(a, b):\n    return a + b\n")
    result = await textfile_edit(
        {
            "paths": [str(target)],
            "edits": [{"op": "replace_text", "match": "add", "replace": "plus"}],
            "max_chars": 1,
        },
        _ctx(),
    )
    assert result["ok"] is False
    assert "校验未通过" in result["reason"]
    assert result["modified"] == 0
    assert "artifact_ref" not in result


# --------------------------------------------------------------------------- #
# 4. Missing extra / unparseable ⇒ metadata_only + verbatim reason (HTTP<500)   #
# --------------------------------------------------------------------------- #
def test_pptx_missing_extra_degrades_to_metadata_only_with_verbatim_reason(monkeypatch):
    from forgeflow.resources import summaries

    data = _pptx_bytes()  # build BEFORE hiding the extra (the builder itself imports pptx)
    # Simulate the optional extra being absent (``import pptx`` ⇒ ImportError).
    monkeypatch.setitem(sys.modules, "pptx", None)
    status, summary, detail = summaries.summarize_presentation(data)
    assert status == "metadata_only"
    assert "pptx support unavailable" in detail
    # Unmeasured ⇒ None, never a fabricated 0 (the honesty rule).
    assert summary.rows is None and summary.chars is None


def test_pptx_unparseable_bytes_degrade_honestly():
    from forgeflow.resources import summaries

    status, _summary, detail = summaries.summarize_presentation(b"not a pptx")
    assert status == "metadata_only"
    assert "unparseable pptx" in detail

    # …and it never claims ``parsed`` through the dispatcher either.
    status2, _s2, detail2 = summaries.summarize_bytes(b"not a pptx", filename="x.pptx")
    assert status2 == "metadata_only"
    assert "unparseable pptx" in detail2


async def test_pptx_upload_of_bad_bytes_is_never_parsed(force_memory_backend):
    from forgeflow.resources.service import ResourceService, reset_resource_index

    reset_resource_index()
    try:
        record = await ResourceService().register_file(
            "t-inc44", name="broken.pptx", data=b"not a pptx", created_by="u"
        )
        assert record.status == "metadata_only"
        assert record.detail  # verbatim reason present
    finally:
        reset_resource_index()
