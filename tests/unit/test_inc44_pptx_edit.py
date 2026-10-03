"""INC44 T02 — PPTX document-editing layered pins (offline; no model, no PG).

Mirrors ``tests/unit/test_inc43_docx_edit.py`` for the presentation format:

1. **Layering** — ``apply_edits`` (the Tool layer) really rewrites the bytes via
   ``python-pptx``, while ``resolve_intent`` (the LLM layer) never touches them.
2. **Recomputable diff** — ``compute_diff``'s four figures are exact for a known
   pair (the shared ``textdiff`` engine, same rule as DOCX).
3. **Honest verification** — ``verify_pptx`` is tri-state: ``None`` for an
   unmeasured dimension, ``False`` only for a *measured* violation.
4. **Honest degradation** — bytes that are not a readable PPTX raise
   ``PptxInspectionError`` (never a fabricated structure); a missing-target op
   changes nothing and reports ``0``.
"""

from __future__ import annotations

import hashlib
import io
from types import SimpleNamespace

import pytest

from forgeflow.documents import (
    PptxEditOp,
    PptxInspectionError,
    UnknownEditOpError,
    apply_pptx_edits,
    compute_pptx_diff,
    inspect_pptx,
    open_pptx,
    resolve_pptx_intent,
    verify_pptx,
)

pytestmark = pytest.mark.asyncio


# --------------------------------------------------------------------------- #
# Builders / helpers                                                           #
# --------------------------------------------------------------------------- #
def _pptx_bytes(*, slides: int = 2) -> bytes:
    """A known PPTX: slide 0 has a title with a number + a body text box."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    layout = presentation.slide_layouts[5]  # "Title Only"
    first = presentation.slides.add_slide(layout)
    first.shapes.title.text = "总金额 40 元"
    box = first.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "备注 请核对"
    for index in range(1, slides):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = f"第 {index + 1} 章 明细"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _slide0_title(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    return str(presentation.slides[0].shapes.title.text or "")


def _ctx() -> SimpleNamespace:
    return SimpleNamespace(run_id="run-inc44", tenant_id="t-pptx", role="manager", args={})


class _FakeModel:
    def __init__(self, content: str) -> None:
        self._content = content
        self.calls = 0

    async def ainvoke(self, _messages):
        self.calls += 1
        return SimpleNamespace(content=self._content)


@pytest.fixture
def no_model(monkeypatch):
    import forgeflow.models.provider as provider

    monkeypatch.setattr(provider, "get_model", lambda *a, **k: None)


@pytest.fixture
def tmp_doc_store(tmp_path, monkeypatch):
    from forgeflow.documents import store as store_mod

    real = store_mod.DocArtifactStore

    def _factory(root=None):
        return real(root=root if root is not None else tmp_path / "docstore")

    import forgeflow.documents as docs_pkg

    monkeypatch.setattr(store_mod, "DocArtifactStore", _factory)
    monkeypatch.setattr(docs_pkg, "DocArtifactStore", _factory)
    return tmp_path


# --------------------------------------------------------------------------- #
# 1. Layering                                                                  #
# --------------------------------------------------------------------------- #
async def test_apply_edits_really_changes_bytes_and_preserves_slides():
    old = _pptx_bytes()
    structure = inspect_pptx(old)

    new, changes = apply_pptx_edits(old, [PptxEditOp(op="replace_text", match="40", replace="38")])

    assert hashlib.sha256(new).hexdigest() != hashlib.sha256(old).hexdigest()
    assert changes == 1
    assert "38" in _slide0_title(new)
    # The slide count is preserved.
    assert inspect_pptx(new).slides == structure.slides == 2


async def test_set_slide_title_and_set_shape_text_write_the_frame():
    old = _pptx_bytes()

    new, changes = apply_pptx_edits(
        old, [PptxEditOp(op="set_slide_title", index=0, text="新标题")]
    )
    assert changes == 1
    assert _slide0_title(new) == "新标题"

    # Address a shape by its 0-based position on the slide (numeric selector).
    new2, changes2 = apply_pptx_edits(
        old, [PptxEditOp(op="set_shape_text", index=0, shape="0", text="替换后的标题")]
    )
    assert changes2 == 1
    assert _slide0_title(new2) == "替换后的标题"


async def test_missing_target_ops_change_nothing_and_report_zero():
    old = _pptx_bytes()
    # Slide index out of range ⇒ no change, no fabricated count.
    # (``apply_edits`` always re-saves, so byte identity is not meaningful for a
    # no-op — the *semantic* diff is what must be empty.)
    new, changes = apply_pptx_edits(
        old, [PptxEditOp(op="set_slide_title", index=99, text="x")]
    )
    assert changes == 0
    assert _slide0_title(new) == _slide0_title(old)
    assert compute_pptx_diff(old, new).to_dict() == {
        "modified": 0,
        "added": 0,
        "removed": 0,
        "numeric_changes": 0,
    }

    # An unknown shape selector ⇒ no change.
    _, changes2 = apply_pptx_edits(
        old, [PptxEditOp(op="set_shape_text", index=0, shape="does-not-exist", text="x")]
    )
    assert changes2 == 0


async def test_unknown_edit_op_is_refused():
    with pytest.raises(UnknownEditOpError):
        PptxEditOp(op="delete_slide")
    with pytest.raises(UnknownEditOpError):
        PptxEditOp.from_dict({"op": "nope"})
    with pytest.raises(UnknownEditOpError):
        PptxEditOp.from_dict({"op": "replace_text", "replace": "x"})
    with pytest.raises(UnknownEditOpError):
        PptxEditOp.from_dict({"op": "set_slide_title", "text": "x"})  # no index


async def test_resolve_intent_never_writes_bytes(monkeypatch):
    data = _pptx_bytes()
    structure = inspect_pptx(data)
    fingerprint = hashlib.sha256(data).hexdigest()

    fake = _FakeModel('{"edits":[{"op":"replace_text","match":"40","replace":"38"}]}')
    monkeypatch.setattr("forgeflow.models.provider.get_model", lambda *a, **k: fake)

    ops = await resolve_pptx_intent(data, "把金额 40 改成 38", structure)

    assert ops == [PptxEditOp(op="replace_text", match="40", replace="38")]
    assert fake.calls == 1
    assert hashlib.sha256(data).hexdigest() == fingerprint
    assert _slide0_title(data) == "总金额 40 元"


async def test_resolve_intent_returns_none_when_model_unavailable(no_model):
    data = _pptx_bytes()
    assert await resolve_pptx_intent(data, "改一下", inspect_pptx(data)) is None


async def test_resolve_intent_returns_none_on_unusable_reply(monkeypatch):
    data = _pptx_bytes()
    structure = inspect_pptx(data)
    monkeypatch.setattr(
        "forgeflow.models.provider.get_model", lambda *a, **k: _FakeModel("不是 JSON")
    )
    assert await resolve_pptx_intent(data, "改一下", structure) is None
    monkeypatch.setattr(
        "forgeflow.models.provider.get_model",
        lambda *a, **k: _FakeModel('{"edits":[{"op":"delete_slide"}]}'),
    )
    assert await resolve_pptx_intent(data, "改一下", structure) is None


# --------------------------------------------------------------------------- #
# 2. Diff is recomputable                                                      #
# --------------------------------------------------------------------------- #
async def test_compute_diff_exact_values_for_known_pair():
    old = _pptx_bytes()
    new, _ = apply_pptx_edits(old, [PptxEditOp(op="replace_text", match="40", replace="38")])

    report = compute_pptx_diff(old, new)
    assert report.to_dict() == {
        "modified": 1,
        "added": 0,
        "removed": 0,
        "numeric_changes": 1,
    }


async def test_compute_diff_is_empty_for_identical_presentations():
    data = _pptx_bytes()
    assert compute_pptx_diff(data, data).to_dict() == {
        "modified": 0,
        "added": 0,
        "removed": 0,
        "numeric_changes": 0,
    }


# --------------------------------------------------------------------------- #
# 3. Inspection + verification are honest                                      #
# --------------------------------------------------------------------------- #
async def test_inspect_pptx_reports_measured_structure():
    structure = inspect_pptx(_pptx_bytes())
    assert structure.slides == 2
    assert structure.text_frames >= 2
    assert structure.char_count > 0
    payload = structure.to_dict(limit=1)
    assert payload["slides"] == 2 and payload["title_count"] == 2
    assert payload["truncated"] is True  # the limit bit ⇒ explicit, not silent


async def test_inspect_pptx_raises_honestly_on_bad_bytes():
    for bad in (b"", b"not a pptx at all"):
        with pytest.raises(PptxInspectionError):
            inspect_pptx(bad)
        with pytest.raises(PptxInspectionError):
            open_pptx(bad)


async def test_verify_pptx_is_tri_state():
    data = _pptx_bytes()

    empty = verify_pptx(data, {})
    assert empty.openable is True
    assert empty.structure_ok is None
    assert empty.data_ok is None
    assert empty.requirement_ok is None

    ok = verify_pptx(data, {"expected_slides": 2})
    assert ok.structure_ok is True
    bad = verify_pptx(data, {"expected_slides": 99})
    assert bad.structure_ok is False

    over = verify_pptx(data, {"max_chars": 1})
    assert over.requirement_ok is False
    assert over.structure_ok is None

    unopenable = verify_pptx(b"garbage", {"expected_slides": 2})
    assert unopenable.openable is False
    assert unopenable.structure_ok is None


async def test_verify_pptx_data_ok_flags_silent_number_change():
    old = _pptx_bytes()
    new, _ = apply_pptx_edits(old, [PptxEditOp(op="replace_text", match="40", replace="38")])
    silent = verify_pptx(new, {"original_numbers": ["40"], "allowed_missing_numbers": []})
    assert silent.data_ok is False
    declared = verify_pptx(new, {"original_numbers": ["40"], "allowed_missing_numbers": ["40"]})
    assert declared.data_ok is True
