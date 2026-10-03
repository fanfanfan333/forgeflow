"""INC46 T23 — five-layer validation stack over the real T26 fidelity corpus.

Why this file exists
--------------------
``tests/unit/test_inc46_validation_stack.py`` pins each layer on synthetic docs.
This file runs the **whole stack over the 41-document T26 fidelity corpus**
(``tests/fixtures/docx_corpus`` — comments / revisions / fields / footnotes /
merged cells / nested tables / textboxes / floating images / headers-footers /
numbering / content controls / hyperlinks / equations / styles / large document),
so the layers are exercised against *real* OOXML structure, not one toy doc.

What it pins
------------
* **Positive** — for every corpus sample, a **compliant** edit makes **L1–L3 all
  ``pass`` with a non-empty ``evidence_ref``** (no empty coverage).
* **Negative (5 per-layer probes, all真跑)** — a real defect per layer is caught
  by that layer (L1 broken package / L2 changed amount / L3 out-of-range reformat
  / L4 injected render diff / L5 low judge).
* **None semantics (red line 15)** — with no engine / no judge, L4 & L5 are
  ``None`` (**never** ``pass``). This box has no LibreOffice, so the probe is
  asserted for real.
* **Counterfactual (5 layers + L5-not-alone)** — disabling a layer through
  ``ValidationConfig.disabled_layers`` turns that layer ``None`` and makes the
  "this layer caught the defect" assertion **red** (真跑, named red point).

All corpus documents are **synthesised** by ``tests/corpus/build_corpus.py`` — no
real customer data (红线 16). This file only adds cases; it never edits existing
tests (红线 1).
"""

from __future__ import annotations

import io
import zipfile

import pytest

from forgeflow.documents import verify_docx
from forgeflow.documents.docx_edit import EditOp, apply_edits
from forgeflow.documents.docx_inspect import open_docx
from forgeflow.documents.validation import (
    ValidationConfig,
    validate_document,
)
from forgeflow.documents.validation import render as render_layer
from forgeflow.documents.validation import semantic as semantic_layer
from forgeflow.documents.validation.verdict import (
    FAIL,
    LAYER_FORMAT,
    LAYER_INVARIANTS,
    LAYER_RENDER,
    LAYER_SEMANTIC,
    LAYER_STRUCTURAL,
    NEEDS_REVIEW,
    PASS,
)
from tests.corpus.build_corpus import CORPUS_DIR, load_manifest

# --------------------------------------------------------------------------- #
# Corpus loading (same manifest/marker discipline as the T26 fidelity suite)   #
# --------------------------------------------------------------------------- #
MANIFEST = load_manifest()
FLAT_SAMPLES: list[tuple[str, dict]] = [
    (feature["id"], sample)
    for feature in MANIFEST["features"]
    for sample in feature["samples"]
]
SAMPLE_IDS = [sample["file"] for _, sample in FLAT_SAMPLES]

_BASE_FILE = "merged_cells_01.docx"


def _sample(filename: str) -> dict:
    return next(sample for _, sample in FLAT_SAMPLES if sample["file"] == filename)


def _read(sample: dict) -> bytes:
    return (CORPUS_DIR / sample["file"]).read_bytes()


def _target_range(data: bytes, marker: str) -> tuple[int, int]:
    texts = [paragraph.text for paragraph in open_docx(data).paragraphs]
    matches = [index for index, text in enumerate(texts) if marker in text]
    assert len(matches) == 1, f"语料 {marker} 命中 {matches}"
    return matches[0], matches[0] + 1


def _compliant_edit(data: bytes, sample: dict) -> bytes:
    edited, changes = apply_edits(
        data,
        [EditOp(op="replace_text", match=sample["marker"], replace=sample["replacement"])],
    )
    assert changes == 1, changes
    return edited


# --------------------------------------------------------------------------- #
# corruption helpers (real byte surgery on the corpus)                          #
# --------------------------------------------------------------------------- #
def _read_parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _repack(parts: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in sorted(parts):
            archive.writestr(name, parts[name])
    return buffer.getvalue()


def _drop_part(data: bytes, name: str) -> bytes:
    parts = _read_parts(data)
    assert name in parts, name
    del parts[name]
    return _repack(parts)


def _mutate_text(data: bytes, old: bytes, new: bytes) -> bytes:
    parts = _read_parts(data)
    blob = parts["word/document.xml"]
    assert old in blob, old
    parts["word/document.xml"] = blob.replace(old, new)
    return _repack(parts)


def _bolden_paragraph(data: bytes, index: int) -> bytes:
    from docx import Document

    document = Document(io.BytesIO(data))
    document.paragraphs[index].runs[0].bold = True
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _grid(color: tuple[int, int, int]) -> list[list[tuple[int, int, int]]]:
    return [[color] * 4 for _ in range(4)]


def _low_judge(_request):
    return semantic_layer.SemanticScore(
        score=0.05,
        dimensions={name: 0.05 for name in semantic_layer.DIMENSIONS},
        judge="corpus-low",
    )


def _render_kwargs(before: bytes, after: bytes) -> dict:
    """A render fixture that never needs LibreOffice.

    The converter keys off the two DOCX byte strings and the rasteriser returns a
    deterministic grid, so the **real** pixel-diff path runs — the layer is
    genuinely measured, never stubbed to a verdict.
    """

    def converter(src: bytes, engine: str) -> bytes:
        return b"PDF-BEFORE" if src == before else b"PDF-AFTER"

    def rasterizer(pdf: bytes, tool: str, dpi: int):
        return _grid((0, 0, 0) if pdf == b"PDF-BEFORE" else (255, 255, 255))

    return {
        "render_engine": "fake-engine",
        "render_converter": converter,
        "rasterizer": rasterizer,
    }


# --------------------------------------------------------------------------- #
# assertion helpers (named so the counterfactual can make them go red)         #
# --------------------------------------------------------------------------- #
def _assert_layer_failed(verdict, layer: str) -> None:
    """The named red point: ``layer`` must have caught the injected defect."""
    actual = verdict.layer(layer)
    assert actual.status == FAIL, f"{layer} 未捕获缺陷（status={actual.status!r}，{actual.evidence_ref}）"


def _assert_positive_mechanical_layers(verdict, filename: str) -> None:
    for layer in (LAYER_STRUCTURAL, LAYER_INVARIANTS, LAYER_FORMAT):
        actual = verdict.layer(layer)
        assert actual.status == PASS, f"{filename}: {layer} 未通过（{actual.evidence_ref}）"
        assert actual.evidence_ref, f"{filename}: {layer} 的 evidence_ref 为空（红线 4）"


# --------------------------------------------------------------------------- #
# A. Positive — the whole corpus passes L1–L3 with non-empty evidence          #
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("feature_id,sample", FLAT_SAMPLES, ids=SAMPLE_IDS)
def test_corpus_compliant_edit_passes_l1_l2_l3(feature_id, sample):
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _compliant_edit(before, sample)

    verdict = validate_document(before, after, start=start, end=end)
    _assert_positive_mechanical_layers(verdict, sample["file"])
    assert verdict.overall in (PASS, NEEDS_REVIEW)
    assert verdict.layer(LAYER_RENDER).status is None
    assert verdict.layer(LAYER_SEMANTIC).status is None


def test_corpus_legacy_verify_plane_still_reads_every_corpus_document():
    """The module→package migration did not break the ``verify_docx`` seam."""
    for _, sample in FLAT_SAMPLES:
        report = verify_docx(_read(sample), {"expected_tables": len(open_docx(_read(sample)).tables)})
        assert report.openable is True, sample["file"]
    assert verify_docx(b"not a docx").openable is False


# --------------------------------------------------------------------------- #
# B. None semantics (red line 15) — unmeasured ⇒ None, never pass              #
# --------------------------------------------------------------------------- #
def test_none_semantics_render_layer_is_unmeasured_on_a_box_without_an_engine():
    engine = render_layer.probe_render_engine()
    if engine is not None:
        pytest.skip("this host has a render engine; the monkeypatched unit nail covers the missing case")
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    verdict = validate_document(before, _compliant_edit(before, sample), start=start, end=end)
    render_verdict = verdict.layer(LAYER_RENDER)
    assert render_verdict.status is None
    assert render_verdict.status != PASS
    assert LAYER_RENDER in verdict.unmeasured


def test_none_semantics_semantic_layer_is_unmeasured_without_a_judge(monkeypatch):
    monkeypatch.setattr(semantic_layer, "_DEFAULT_JUDGE", None)
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    verdict = validate_document(before, _compliant_edit(before, sample), start=start, end=end)
    semantic = verdict.layer(LAYER_SEMANTIC)
    assert semantic.status is None
    assert semantic.status != PASS
    assert LAYER_SEMANTIC in verdict.unmeasured


# --------------------------------------------------------------------------- #
# C. Negative — a real defect per layer is caught by that layer                #
# --------------------------------------------------------------------------- #
def test_negative_L1_broken_package_is_failed():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _drop_part(before, "[Content_Types].xml")
    verdict = validate_document(before, after, start=start, end=end)
    _assert_layer_failed(verdict, LAYER_STRUCTURAL)
    assert verdict.overall == FAIL


def test_negative_L2_changed_amount_is_failed():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _mutate_text(before, b"1,000.00", b"9,999.99")
    verdict = validate_document(before, after, start=start, end=end)
    _assert_layer_failed(verdict, LAYER_INVARIANTS)


def test_negative_L3_out_of_range_reformatting_is_failed():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _bolden_paragraph(before, 0)  # paragraph 0 is outside [start, end)
    verdict = validate_document(before, after, start=start, end=end)
    _assert_layer_failed(verdict, LAYER_FORMAT)


def test_negative_L4_render_difference_is_failed():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _compliant_edit(before, sample)
    verdict = validate_document(
        before,
        after,
        start=start,
        end=end,
        **_render_kwargs(before, after),
    )
    _assert_layer_failed(verdict, LAYER_RENDER)
    assert verdict.layer(LAYER_RENDER).detail["diff_ratio"] == 1.0


def test_negative_L5_low_score_is_needs_review():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    verdict = validate_document(
        before, _compliant_edit(before, sample), start=start, end=end, judge=_low_judge
    )
    assert verdict.layer(LAYER_SEMANTIC).status == NEEDS_REVIEW


# --------------------------------------------------------------------------- #
# D. Counterfactual — disabling a layer makes its probe go RED (真跑)           #
# --------------------------------------------------------------------------- #
def test_counterfactual_L1_disabling_structural_lets_a_broken_package_through():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _drop_part(before, "[Content_Types].xml")

    enabled = validate_document(before, after, start=start, end=end)
    _assert_layer_failed(enabled, LAYER_STRUCTURAL)  # green: L1 catches it

    disabled = validate_document(
        before, after, start=start, end=end,
        config=ValidationConfig(disabled_layers=frozenset({LAYER_STRUCTURAL})),
    )
    with pytest.raises(AssertionError) as excinfo:
        _assert_layer_failed(disabled, LAYER_STRUCTURAL)
    assert LAYER_STRUCTURAL in str(excinfo.value)


def test_counterfactual_L2_disabling_invariants_lets_a_changed_amount_through():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _mutate_text(before, b"1,000.00", b"9,999.99")

    enabled = validate_document(before, after, start=start, end=end)
    _assert_layer_failed(enabled, LAYER_INVARIANTS)

    disabled = validate_document(
        before, after, start=start, end=end,
        config=ValidationConfig(disabled_layers=frozenset({LAYER_INVARIANTS})),
    )
    with pytest.raises(AssertionError) as excinfo:
        _assert_layer_failed(disabled, LAYER_INVARIANTS)
    assert LAYER_INVARIANTS in str(excinfo.value)


def test_counterfactual_L3_disabling_format_lets_a_reformatted_paragraph_through():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _bolden_paragraph(before, 0)

    enabled = validate_document(before, after, start=start, end=end)
    _assert_layer_failed(enabled, LAYER_FORMAT)

    disabled = validate_document(
        before, after, start=start, end=end,
        config=ValidationConfig(disabled_layers=frozenset({LAYER_FORMAT})),
    )
    with pytest.raises(AssertionError) as excinfo:
        _assert_layer_failed(disabled, LAYER_FORMAT)
    assert LAYER_FORMAT in str(excinfo.value)


def test_counterfactual_L4_disabling_render_lets_a_visual_diff_through():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _compliant_edit(before, sample)
    render_kwargs = _render_kwargs(before, after)

    enabled = validate_document(before, after, start=start, end=end, **render_kwargs)
    _assert_layer_failed(enabled, LAYER_RENDER)

    disabled = validate_document(
        before,
        after,
        start=start,
        end=end,
        config=ValidationConfig(disabled_layers=frozenset({LAYER_RENDER})),
        **render_kwargs,
    )
    with pytest.raises(AssertionError) as excinfo:
        _assert_layer_failed(disabled, LAYER_RENDER)
    assert LAYER_RENDER in str(excinfo.value)


def test_counterfactual_L5_disabling_semantic_lets_a_low_score_through():
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _compliant_edit(before, sample)

    enabled = validate_document(before, after, start=start, end=end, judge=_low_judge)
    assert enabled.layer(LAYER_SEMANTIC).status == NEEDS_REVIEW

    def _assert_review(verdict) -> None:
        assert verdict.layer(LAYER_SEMANTIC).status == NEEDS_REVIEW, (
            f"{LAYER_SEMANTIC} 未报 needs_review（status="
            f"{verdict.layer(LAYER_SEMANTIC).status!r}）"
        )

    disabled = validate_document(
        before, after, start=start, end=end, judge=_low_judge,
        config=ValidationConfig(disabled_layers=frozenset({LAYER_SEMANTIC})),
    )
    with pytest.raises(AssertionError) as excinfo:
        _assert_review(disabled)
    assert LAYER_SEMANTIC in str(excinfo.value)


def test_counterfactual_L5_low_score_never_alone_releases():
    """A low semantic score must raise ``needs_review`` — it can never be a release."""
    sample = _sample(_BASE_FILE)
    before = _read(sample)
    start, end = _target_range(before, sample["marker"])
    after = _compliant_edit(before, sample)

    def _assert_not_released(verdict) -> None:
        assert verdict.overall != PASS, f"低分语义仍被放行（overall={verdict.overall}）"

    enabled = validate_document(before, after, start=start, end=end, judge=_low_judge)
    # every mechanical layer passed …
    _assert_positive_mechanical_layers(enabled, sample["file"])
    # … yet the overall verdict is NOT a release
    assert enabled.overall == NEEDS_REVIEW
    _assert_not_released(enabled)  # green

    # counterfactual: remove the semantic layer and the very same documents become
    # a plain "pass" — proof that here L5 is the only thing holding the release
    # back, which is why the low score can never be allowed to be "pass" itself.
    disabled = validate_document(
        before, after, start=start, end=end, judge=_low_judge,
        config=ValidationConfig(disabled_layers=frozenset({LAYER_SEMANTIC})),
    )
    assert disabled.overall == PASS
    with pytest.raises(AssertionError) as excinfo:
        _assert_not_released(disabled)
    assert "放行" in str(excinfo.value)
