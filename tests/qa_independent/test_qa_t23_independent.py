"""T23 第二轮独立验证 —— QA(Edward) 独立探针（D–H 项）。

本文件 **不** 复跑工程师用例，而是独立构造 before/after 文档与反事实注入点，
逐项验证五层验证栈的 DoD。落点 tests/qa_independent/（本次不提交）。
所有断言自跑、证据自出。
"""
from __future__ import annotations

import io
import shutil
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from forgeflow.documents.validation import (
    format as format_layer,
    render as render_layer,
    semantic as semantic_layer,
    structural as structural_layer,
)
from forgeflow.documents.validation import stack as stack_mod
from forgeflow.documents.validation.verdict import (
    FAIL,
    LAYER_FORMAT,
    LAYER_INVARIANTS,
    LAYER_RENDER,
    LAYER_SEMANTIC,
    LAYER_STRUCTURAL,
    NEEDS_REVIEW,
    PASS,
    LayerVerdict,
    ValidationConfig,
)

MARK = "[[TGT::qa23]]"
REPL = "已编辑::qa23"
AMOUNT = "1,000.00"


# --------------------------------------------------------------------------- #
# builders                                                                     #
# --------------------------------------------------------------------------- #
def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _rpr_target_doc() -> bytes:
    """Heading(0) / 目标段(1, 带显式 rPr) / 区间外体段(2, 金额)."""
    d = Document()
    d.add_heading("报告标题", level=1)
    p = d.add_paragraph()
    r = p.add_run(f"占位 {MARK} 金额 {AMOUNT} 元，日期 2026-10-03。")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
    r.font.name = "Arial"
    d.add_paragraph(f"区间外正文，金额 2,000.00 元。")
    return _save(d)


def _plain_doc() -> bytes:
    d = Document()
    d.add_heading("报告标题", level=1)
    d.add_paragraph(f"占位 {MARK}，金额 {AMOUNT} 元，日期 2026-10-03。")
    d.add_paragraph("区间外正文，金额 2,000.00 元。")
    return _save(d)


def _target_range(data: bytes) -> tuple[int, int]:
    from forgeflow.documents.docx_inspect import open_docx

    texts = [p.text for p in open_docx(data).paragraphs]
    idx = next(i for i, t in enumerate(texts) if MARK in t)
    return idx, idx + 1


def _compliant(data: bytes) -> bytes:
    from forgeflow.documents.docx_edit import EditOp, apply_edits

    edited, n = apply_edits(data, [EditOp(op="replace_text", match=MARK, replace=REPL)])
    assert n == 1, n
    return edited


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def _repack(parts: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n in sorted(parts):
            z.writestr(n, parts[n])
    return buf.getvalue()


def _drop(data: bytes, name: str) -> bytes:
    p = _parts(data)
    assert name in p, name
    del p[name]
    return _repack(p)


def _edit_xml(data: bytes, old: bytes, new: bytes) -> bytes:
    p = _parts(data)
    blob = p["word/document.xml"]
    assert old in blob, old
    p["word/document.xml"] = blob.replace(old, new)
    return _repack(p)


def _bolden_para(data: bytes, index: int) -> bytes:
    d = Document(io.BytesIO(data))
    d.paragraphs[index].runs[0].bold = True
    return _save(d)


# =========================================================================== #
# D. L3 DoD 子项「区间内 run 级属性保留 / 被改文本继承原 run 样式」定向构造      #
# =========================================================================== #
def _assert_no_style_lost(v):
    assert v.status == PASS, f"L3 未通过（{v.evidence_ref}）"
    assert v.detail["style_inheritance_failures"] == []


def _assert_style_lost(v):
    assert v.status == FAIL, f"L3 未捕获样式丢失（status={v.status!r} {v.evidence_ref}）"
    assert "style-lost:" in v.evidence_ref, v.evidence_ref


def test_D1_in_range_edit_keeps_run_style_passes():
    before = _rpr_target_doc()
    s, e = _target_range(before)
    v = format_layer.format_verdict(before, _compliant(before), start=s, end=e)
    _assert_no_style_lost(v)
    assert v.detail["in_range_edited"] == 1
    assert v.detail["out_of_range_changes"] == []


def test_D2_in_range_edit_drops_run_style_fails_with_style_lost():
    before = _rpr_target_doc()
    s, e = _target_range(before)
    d = Document(io.BytesIO(before))
    run = d.paragraphs[s].runs[0]
    rpr = run._r.find(qn("w:rPr"))
    assert rpr is not None
    run._r.remove(rpr)  # 丢掉 run 级样式（b/color/sz/rFonts 全没）
    run.text = run.text.replace(MARK, REPL)  # 文本也变了
    after = _save(d)
    v = format_layer.format_verdict(before, after, start=s, end=e)
    _assert_style_lost(v)
    assert v.detail["style_inheritance_failures"] == [f"paragraph[{s}]"]


def test_D3_counterfactual_neutralising_style_inheritance_check_turns_D2_red(monkeypatch):
    before = _rpr_target_doc()
    s, e = _target_range(before)
    d = Document(io.BytesIO(before))
    run = d.paragraphs[s].runs[0]
    rpr = run._r.find(qn("w:rPr"))
    run._r.remove(rpr)
    run.text = run.text.replace(MARK, REPL)
    after = _save(d)

    # 真跑绿：样式继承检查承重
    _assert_style_lost(format_layer.format_verdict(before, after, start=s, end=e))

    # 反事实：摘掉「样式继承检查」（把取 run 签名的辅助函数置为恒 None）⇒ 断言转红
    monkeypatch.setattr(format_layer, "_first_run_signature", lambda p: None)
    with pytest.raises(AssertionError) as ei:
        _assert_style_lost(format_layer.format_verdict(before, after, start=s, end=e))
    assert "L3 未捕获样式丢失" in str(ei.value)


# =========================================================================== #
# E. None 语义真实性（红线 15/4）—— 真实探测，非硬编码                     #
# =========================================================================== #
def test_E1_render_engine_probe_is_real_not_hardcoded(monkeypatch):
    real = shutil.which("soffice") or shutil.which("libreoffice")
    assert render_layer.probe_render_engine() == real  # 与真实 which 一致

    monkeypatch.setattr(
        render_layer.shutil, "which",
        lambda name: "/fake/bin/soffice" if name == "soffice" else None,
    )
    assert render_layer.probe_render_engine() == "/fake/bin/soffice"


def test_E1b_L4_measures_once_engine_resolves(monkeypatch):
    before, after = _rpr_target_doc(), _rpr_target_doc()
    monkeypatch.setattr(
        render_layer.shutil, "which",
        lambda name: "/fake/bin/soffice" if name == "soffice" else None,
    )
    grid = [[(7, 7, 7)] * 3 for _ in range(3)]
    v = render_layer.render_verdict(
        before, after,
        converter=lambda src, eng: b"%PDF-fake",
        rasterizer=lambda pdf, tool, dpi: grid,
    )
    assert v.status in (PASS, FAIL), f"引擎存在时 L4 仍未测量：{v.status!r} {v.evidence_ref}"


def test_E1c_second_engine_probe_is_real():
    real = shutil.which("soffice") or shutil.which("libreoffice")
    assert structural_layer.probe_second_engine() == real


def test_E1d_page_count_reader_never_fabricates_zero():
    assert render_layer.probe_pdf_page_count(b"not a pdf at all") is None


def test_E2_L5_none_is_genuine_and_judge_injection_measures(monkeypatch):
    before = _plain_doc()
    s, e = _target_range(before)
    after = _compliant(before)

    monkeypatch.setattr(semantic_layer, "_DEFAULT_JUDGE", None)
    none_v = semantic_layer.semantic_verdict(before, after, start=s, end=e)
    assert none_v.status is None and none_v.status != PASS

    semantic_layer.set_default_judge(
        lambda req: semantic_layer.SemanticScore(
            score=0.9, dimensions={n: 0.9 for n in semantic_layer.DIMENSIONS}, judge="qa"
        )
    )
    try:
        measured = semantic_layer.semantic_verdict(before, after, start=s, end=e)
        assert measured.status in (PASS, NEEDS_REVIEW)
    finally:
        semantic_layer.set_default_judge(None)


# =========================================================================== #
# F. 新 F1 攻击面 —— 页数变化 warn 的分支/护栏/优先级                          #
# =========================================================================== #
def _L(name, status, detail=None):
    return LayerVerdict(name, status, f"ev:{name}", detail or {})


def _stack(l1=PASS, l2=PASS, l3=PASS, l4=PASS, l4d=None, l5=PASS):
    return {
        LAYER_STRUCTURAL: _L(LAYER_STRUCTURAL, l1),
        LAYER_INVARIANTS: _L(LAYER_INVARIANTS, l2),
        LAYER_FORMAT: _L(LAYER_FORMAT, l3),
        LAYER_RENDER: _L(LAYER_RENDER, l4, l4d or {}),
        LAYER_SEMANTIC: _L(LAYER_SEMANTIC, l5),
    }


def _agg(layers, cfg=None):
    return stack_mod._aggregate(layers, cfg or ValidationConfig())


def test_F1a_none_page_delta_never_warns():
    overall, note = _agg(_stack(l4d={"page_delta": None, "pages_before": None, "pages_after": None}))
    assert overall == PASS, note


def test_F1b_bool_guard_true_false_do_not_warn():
    # bool 绝不能当页数变化（isinstance(int) and not isinstance(bool) 护栏）
    assert _agg(_stack(l4d={"page_delta": True, "pages_before": 0, "pages_after": 1}))[0] == PASS
    assert _agg(_stack(l4d={"page_delta": False, "pages_before": 1, "pages_after": 1}))[0] == PASS


def test_F1c_nonzero_int_warns_zero_does_not():
    assert _agg(_stack(l4d={"page_delta": 1, "pages_before": 1, "pages_after": 2}))[0] == "warn"
    assert _agg(_stack(l4d={"page_delta": 0, "pages_before": 1, "pages_after": 1}))[0] == PASS
    assert _agg(_stack(l4d={}))[0] == PASS  # 键缺失


def test_F1d_page_decrease_warns_too():
    overall, note = _agg(_stack(l4d={"page_delta": -1, "pages_before": 2, "pages_after": 1}))
    assert overall == "warn", note


def test_F1e_priority_matrix():
    # L1–L3 fail > L4 fail
    o, n = _agg(_stack(l2=FAIL, l4=FAIL, l4d={"page_delta": 1}))
    assert o == FAIL and "硬失败层：L2" in n, (o, n)
    # L4 fail（非致命）> 页数 warn：走降级分支
    o, n = _agg(_stack(l4=FAIL, l4d={"page_delta": 1}))
    assert o == "warn" and "降级为 warn" in n, (o, n)
    # L4 fail 致命 ⇒ fail
    o, n = _agg(_stack(l4=FAIL), ValidationConfig(l4_fail_is_fatal=True))
    assert o == FAIL, (o, n)
    # 页数 warn > L5 needs_review
    o, n = _agg(_stack(l4d={"page_delta": 1, "pages_before": 1, "pages_after": 2}, l5=NEEDS_REVIEW))
    assert o == "warn", (o, n)
    # L5 needs_review > pass
    o, n = _agg(_stack(l5=NEEDS_REVIEW))
    assert o == NEEDS_REVIEW, (o, n)
    # 全 pass
    o, n = _agg(_stack())
    assert o == PASS and "全部" in n, (o, n)
    # 存在未测量层（None）但机械层通过 ⇒ pass（并在 notes 里被点出，见 validate_document）
    o, n = _agg(_stack(l5=None))
    assert o == PASS and "未测量" in n, (o, n)


# 端到端：page_counter 控制页数
def _page_kwargs(before, after, pages):
    def conv(src, eng):
        return b"PDF-B" if src == before else b"PDF-A"

    def raster(pdf, tool, dpi):
        return [[(9, 9, 9)] * 3 for _ in range(3)]  # 两侧同一栅格 ⇒ diff 0

    return {
        "render_engine": "fake",
        "render_converter": conv,
        "rasterizer": raster,
        "page_counter": lambda pdf: pages.get(pdf),
    }


def test_F1f_e2e_page_change_warns_without_failing_L4():
    before = _plain_doc()
    s, e = _target_range(before)
    after = _compliant(before)
    from forgeflow.documents.validation import validate_document

    up = validate_document(before, after, start=s, end=e,
                           **_page_kwargs(before, after, {b"PDF-B": 1, b"PDF-A": 2}))
    assert up.layer(LAYER_RENDER).status != FAIL
    assert up.layer(LAYER_RENDER).detail["page_delta"] == 1
    assert up.overall == "warn" and "页数变化" in up.notes

    down = validate_document(before, after, start=s, end=e,
                             **_page_kwargs(before, after, {b"PDF-B": 2, b"PDF-A": 1}))
    assert down.overall == "warn" and "页数变化" in down.notes

    none = validate_document(before, after, start=s, end=e,
                             **_page_kwargs(before, after, {}))
    assert none.layer(LAYER_RENDER).detail["page_delta"] is None
    assert none.overall != "warn"


def test_F5_counterfactual_third_injection_point(monkeypatch):
    """自设的第三种注入点：把渲染层的 _page_delta（页数测量）置为恒「未测量」⇒ 断言转红。"""
    from forgeflow.documents.validation import validate_document

    before = _plain_doc()
    s, e = _target_range(before)
    after = _compliant(before)
    kwargs = _page_kwargs(before, after, {b"PDF-B": 1, b"PDF-A": 2})

    def _assert_page_warn(v):
        assert v.overall == "warn", f"页数变化未告警（overall={v.overall}）"

    _assert_page_warn(validate_document(before, after, start=s, end=e, **kwargs))  # 绿

    monkeypatch.setattr(
        render_layer, "_page_delta",
        lambda b, a, c: (None, None, None, ";pages=not measured"),
    )
    with pytest.raises(AssertionError) as ei:
        _assert_page_warn(validate_document(before, after, start=s, end=e, **kwargs))
    assert "页数变化未告警" in str(ei.value)


# =========================================================================== #
# G. DoD 阴性探针（5 条）                                                      #
# =========================================================================== #
def test_G1_changed_amount_L2_fail():
    from forgeflow.documents.validation import invariants as inv_layer

    before = _plain_doc()
    s, e = _target_range(before)
    after = _edit_xml(before, AMOUNT.encode(), b"9,999.99")
    v = inv_layer.invariants_verdict(before, after, start=s, end=e)
    assert v.status == FAIL and any("amount" in x for x in v.detail["lost"]), v.evidence_ref


def test_G2_bolden_out_of_range_L3_fail():
    before = _plain_doc()
    s, e = _target_range(before)
    after = _bolden_para(before, 2)  # 区间外体段变粗
    v = format_layer.format_verdict(before, after, start=s, end=e)
    assert v.status == FAIL and v.detail["out_of_range_changes"], v.evidence_ref


def test_G3_out_of_range_text_change_L3_fail():
    before = _plain_doc()
    s, e = _target_range(before)
    after = _edit_xml(before, "区间外正文".encode(), "区间外正文（改）".encode())
    v = format_layer.format_verdict(before, after, start=s, end=e)
    # 末段落按「镜像锚定」标记为 tail-1（format.py 明确设计）
    assert v.status == FAIL and v.detail["out_of_range_changes"] == ["paragraph[tail-1]"], v.evidence_ref


def test_G4_broken_rels_L1_fail():
    before = _plain_doc()
    p = _parts(before)
    rels = p["word/_rels/document.xml.rels"]
    p["word/_rels/document.xml.rels"] = rels.replace(b"Target=", b"Target=missing_", 1)
    v = structural_layer.structural_verdict(_repack(p))
    assert v.status == FAIL and "package-integrity:FAIL[" in v.evidence_ref, v.evidence_ref


def test_G5_inserted_new_number_L2_fail():
    from forgeflow.documents.validation import invariants as inv_layer

    before = _plain_doc()
    s, e = _target_range(before)
    # 区间内目标段尾部追加一个新数字（"新增 7 项"）⇒ 「不新增数字」基线应 fail
    after = _edit_xml(before, "2026-10-03。".encode(), "2026-10-03。新增 7 项".encode())
    v = inv_layer.invariants_verdict(before, after, start=s, end=e)
    assert v.status == FAIL and v.detail["new_numbers"], v.evidence_ref
    assert any(n.startswith("7x") for n in v.detail["new_numbers"]), v.detail["new_numbers"]


# =========================================================================== #
# H. F2 / F3 验收                                                              #
# =========================================================================== #
def test_H1_L1_evidence_is_package_integrity_never_schema():
    ok = structural_layer.structural_verdict(_plain_doc())
    assert ok.status == PASS
    assert "package-integrity:" in ok.evidence_ref
    assert "schema" not in ok.evidence_ref.lower()

    bad = structural_layer.structural_verdict(_drop(_plain_doc(), "[Content_Types].xml"))
    assert bad.status == FAIL
    assert "package-integrity:FAIL[" in bad.evidence_ref
    assert "schema" not in bad.evidence_ref.lower()


def test_H2_L1_required_subchecks_are_load_bearing():
    base = _plain_doc()
    # 1) 删 [Content_Types].xml
    assert structural_layer.structural_verdict(_drop(base, "[Content_Types].xml")).status == FAIL
    # 2) 删 _rels/.rels
    assert structural_layer.structural_verdict(_drop(base, "_rels/.rels")).status == FAIL
    # 3) 注入未解析 rel target
    p = _parts(base)
    rels = p["word/_rels/document.xml.rels"]
    p["word/_rels/document.xml.rels"] = rels.replace(b"Target=", b"Target=broken_", 1)
    assert structural_layer.structural_verdict(_repack(p)).status == FAIL


def test_H3_empty_denominator_evidence_is_na_not_zero_over_zero():
    from forgeflow.documents.validation import invariants as inv_layer

    data = _plain_doc()
    # L3：无区间内编辑 ⇒ in_range_edited == 0
    l3 = format_layer.format_verdict(data, data, start=1, end=2)
    assert l3.status == PASS
    assert "n/a" in l3.evidence_ref and "0/0" not in l3.evidence_ref
    assert l3.detail["in_range_edited"] == 0

    # L2：区间内无显式不变量 ⇒ compared == 0
    l2 = inv_layer.invariants_verdict(data, data, start=0, end=1)
    assert l2.status == PASS
    assert "n/a" in l2.evidence_ref and "0/0" not in l2.evidence_ref
    assert l2.detail["compared"] == 0
