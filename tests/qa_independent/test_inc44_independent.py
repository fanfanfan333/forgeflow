"""QA (严过关 / Edward) INC44-C — **independent** verification harness.

This file is written from scratch by QA and deliberately does **NOT** import or
rely on the engineer's ``tests/unit/test_inc44_*`` / ``tests/integration/test_inc44_*``
suites. Wherever possible it drives the platform with QA's *own* inputs and, for
the load-bearing claims, builds an explicit **counterfactual** (mutate the
mechanism under test and require the corresponding pin to go RED) so a green
result cannot be vacuous.

Run (memory profile, offline):
    cd ForgeFlow-main
    <venv>/python -m pytest tests/qa_independent/test_inc44_independent.py -q

Points map 1:1 to the INC44-C verification brief (1..9).
"""

from __future__ import annotations

import hashlib
import io
import os
import sys
import uuid
from types import SimpleNamespace
from urllib.parse import quote

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

# --------------------------------------------------------------------------- #
# Shared helpers                                                               #
# --------------------------------------------------------------------------- #
TENANT = "t-qa-inc44"


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _pptx_bytes(title: str = "总金额 40 元", body: str = "备注 请核对") -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _docx_bytes(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


class _FakeModel:
    def __init__(self, content: str) -> None:
        self._content = content
        self.calls = 0

    async def ainvoke(self, _messages):
        self.calls += 1
        return SimpleNamespace(content=self._content)


def _ctx() -> SimpleNamespace:
    return SimpleNamespace(
        run_id="run-qa", step_id="run-qa:0:0", tenant_id=TENANT,
        user_id="qa", role="manager", args={},
    )


# =========================================================================== #
# 1. T01 — closed loop really persists + publishes (with counterfactual)       #
# =========================================================================== #
@pytest.fixture
def clean_skills():
    from forgeflow.repositories.memory.skill_repo import clear_skill_store

    clear_skill_store()
    yield
    clear_skill_store()


async def _seed_experiences(repo, tenant: str, n: int = 3):
    from forgeflow.experience.embedding import deterministic_embedding
    from forgeflow.experience.models import ExperienceRecord

    text = "分析销售数据 数据分析"
    emb = deterministic_embedding(text)
    for i in range(n):
        await repo.save(
            ExperienceRecord(
                tenant_id=tenant, run_id=f"r-{i}",
                summary=f"{text} — 共执行 2 个步骤，结果：success",
                outcome="success",
                reusable_steps=[{"tool": "data.query"}, {"tool": "report.render"}],
                tags=["数据分析"], embedding=emb,
            )
        )


async def test_P1_synthesize_loop_publish_closes_and_is_load_bearing(clean_skills, monkeypatch):
    from forgeflow.repositories.memory.experience_repo import MemoryExperienceRepository
    from forgeflow.repositories.memory.policy_repo import MemoryPolicyRepository
    from forgeflow.repositories.memory.skill_repo import (
        MemorySkillCandidateRepository,
        MemorySkillRepository,
    )
    from forgeflow.skills import engineering as eng
    from forgeflow.skills.tester import structural_score

    tenant = f"t-qa-{uuid.uuid4().hex[:6]}"
    exp_repo = MemoryExperienceRepository()
    cand_repo = MemorySkillCandidateRepository()
    skill_repo = MemorySkillRepository()
    pol_repo = MemoryPolicyRepository()
    await _seed_experiences(exp_repo, tenant)

    contract = await eng.synthesize(
        tenant, mode="auto", candidate_repo=cand_repo, experience_repo=exp_repo
    )
    assert contract.is_complete() is True
    cands = await cand_repo.list_candidates(tenant)
    assert cands, "synthesize must leave a candidate"
    cid = cands[0].id

    res = await eng.run_engineering_loop(
        tenant, cid, "manager-1", "manager",
        candidate_repo=cand_repo, skill_repo=skill_repo, policy_repo=pol_repo,
    )

    # (1) the loop's own evaluation is really stored
    ev = await cand_repo.get_evaluation_for(tenant, cid)
    assert ev is not None, "loop must persist an evaluation (P1 fix)"
    assert ev.verdict == "pass"
    assert ev.metrics["score"] == structural_score(res.contract)
    # (2) candidate advanced + derived lifecycle agrees with the returned state
    stored = await cand_repo.get_candidate(tenant, cid)
    assert stored.status == "approved"
    assert eng.derive_lifecycle(stored) == "REVIEW" == res.lifecycle
    # (3) the *verified* contract is the one that will be published
    assert stored.draft_spec == res.contract.to_draft_spec()
    # (4) publish now really succeeds (no 403)
    version = await eng.version_and_publish(
        tenant, cid, "manager-1", "manager",
        candidate_repo=cand_repo, skill_repo=skill_repo, policy_repo=pol_repo,
    )
    parts = version.semver.split(".")
    assert len(parts) == 3 and all(p.isdigit() for p in parts)
    assert version.eval_score is not None

    # ---- COUNTERFACTUAL: remove the persistence side effect -------------- #
    # Re-run the whole loop with ``_persist_loop_evaluation`` neutered; the
    # positive pin above MUST go red (evaluation absent + publish 403).
    tenant2 = f"t-qa-{uuid.uuid4().hex[:6]}"
    await _seed_experiences(exp_repo, tenant2)

    async def _noop(*_a, **_k):
        return 0.0

    monkeypatch.setattr(eng, "_persist_loop_evaluation", _noop)
    await eng.synthesize(
        tenant2, mode="auto", candidate_repo=cand_repo, experience_repo=exp_repo
    )
    cid2 = (await cand_repo.list_candidates(tenant2))[0].id
    res2 = await eng.run_engineering_loop(
        tenant2, cid2, "manager-1", "manager",
        candidate_repo=cand_repo, skill_repo=skill_repo, policy_repo=pol_repo,
    )
    assert res2.passed is True  # the loop still *computed* success
    assert await cand_repo.get_evaluation_for(tenant2, cid2) is None, (
        "counterfactual failed: persistence is NOT load-bearing"
    )
    from forgeflow.skills.errors import GovernanceError

    with pytest.raises(GovernanceError):
        await eng.version_and_publish(
            tenant2, cid2, "manager-1", "manager",
            candidate_repo=cand_repo, skill_repo=skill_repo, policy_repo=pol_repo,
        )


# =========================================================================== #
# 2 + 3. Exhaustion is honest; critique-block persists nothing                 #
# =========================================================================== #
async def test_P2_exhaustion_is_honest_and_still_403(clean_skills):
    from forgeflow.repositories.memory.policy_repo import MemoryPolicyRepository
    from forgeflow.repositories.memory.skill_repo import (
        MemorySkillCandidateRepository,
        MemorySkillRepository,
    )
    from forgeflow.skills import engineering as eng
    from forgeflow.skills.errors import GovernanceError
    from forgeflow.skills.models import SkillCandidateRecord

    tenant = f"t-qa-{uuid.uuid4().hex[:6]}"
    cand_repo = MemorySkillCandidateRepository()
    cand = SkillCandidateRecord(
        tenant_id=tenant, name="weak", domain="数据分析",
        experience_ids=["e1"], status="draft",
        draft_spec={
            "prompt": "一句话目标",
            "steps": ["唯一步骤"],
            "tools": ["report.render"],
            "applicable_when": {},
        },
    )
    await cand_repo.save_candidate(cand)

    res = await eng.run_engineering_loop(
        tenant, cand.id, "manager-1", "manager", max_repair=0,
        candidate_repo=cand_repo, skill_repo=MemorySkillRepository(),
        policy_repo=MemoryPolicyRepository(),
    )
    assert res.passed is False
    ev = await cand_repo.get_evaluation_for(tenant, cand.id)
    assert ev is not None and ev.verdict == "fail"
    assert (await cand_repo.get_candidate(tenant, cand.id)).status == "draft"

    with pytest.raises(GovernanceError) as ei:
        await eng.version_and_publish(
            tenant, cand.id, "manager-1", "manager",
            candidate_repo=cand_repo, skill_repo=MemorySkillRepository(),
            policy_repo=MemoryPolicyRepository(),
        )
    assert ei.value.status_code == 403, "exhausted candidate must still 403"


async def test_P3_critique_block_persists_nothing(clean_skills):
    from forgeflow.repositories.memory.policy_repo import MemoryPolicyRepository
    from forgeflow.repositories.memory.skill_repo import (
        MemorySkillCandidateRepository,
        MemorySkillRepository,
    )
    from forgeflow.skills import engineering as eng
    from forgeflow.skills.errors import GovernanceError
    from forgeflow.skills.models import SkillCandidateRecord

    tenant = f"t-qa-{uuid.uuid4().hex[:6]}"
    cand_repo = MemorySkillCandidateRepository()
    cand = SkillCandidateRecord(
        tenant_id=tenant, name="empty", domain="数据分析",
        experience_ids=["e1"], status="draft", draft_spec={},
    )
    await cand_repo.save_candidate(cand)

    res = await eng.run_engineering_loop(
        tenant, cand.id, "manager-1", "manager", max_repair=0,
        candidate_repo=cand_repo, skill_repo=MemorySkillRepository(),
        policy_repo=MemoryPolicyRepository(),
    )
    assert res.passed is False and res.critique.must_fix
    assert await cand_repo.get_evaluation_for(tenant, cand.id) is None
    with pytest.raises(GovernanceError):
        await eng.version_and_publish(
            tenant, cand.id, "manager-1", "manager",
            candidate_repo=cand_repo, skill_repo=MemorySkillRepository(),
            policy_repo=MemoryPolicyRepository(),
        )


# =========================================================================== #
# 4. T02 layering: LLM never writes, Tool really writes, diff recomputable      #
# =========================================================================== #
async def test_P4_layer_split_pptx_and_docx(monkeypatch):
    from forgeflow.documents import apply_edits as apply_docx_edits
    from forgeflow.documents import (
        apply_pptx_edits,
        inspect_docx,
        inspect_pptx,
        resolve_pptx_intent,
    )
    from forgeflow.documents import resolve_intent as resolve_docx_intent

    # --- PPTX ------------------------------------------------------------- #
    data = _pptx_bytes()
    struct = inspect_pptx(data)
    fp = _sha(data)
    fake = _FakeModel('{"edits":[{"op":"replace_text","match":"40","replace":"42"}]}')
    monkeypatch.setattr("forgeflow.models.provider.get_model", lambda *a, **k: fake)

    ops = await resolve_pptx_intent(data, "把 40 改成 42", struct)
    assert ops and ops[0].replace == "42"
    assert _sha(data) == fp, "LLM layer must not write bytes"

    new, n = apply_pptx_edits(data, ops)
    assert n == 1
    assert _sha(new) != fp, "Tool layer must really write bytes"
    from forgeflow.documents.pptx_inspect import open_pptx, shape_paragraph_texts

    texts = "\n".join(shape_paragraph_texts(open_pptx(new)))
    assert "42" in texts and "40" not in texts

    # --- DOCX ------------------------------------------------------------- #
    ddata = _docx_bytes(["总金额 40 元", "备注 A"])
    dstruct = inspect_docx(ddata)
    dfp = _sha(ddata)
    fake2 = _FakeModel('{"edits":[{"op":"replace_text","match":"40","replace":"42"}]}')
    monkeypatch.setattr("forgeflow.models.provider.get_model", lambda *a, **k: fake2)

    dops = await resolve_docx_intent(ddata, "把 40 改成 42", dstruct)
    assert dops and dops[0].replace == "42"
    assert _sha(ddata) == dfp, "DOCX LLM layer must not write bytes"

    dnew, dn = apply_docx_edits(ddata, dops)
    assert dn == 1 and _sha(dnew) != dfp
    import docx

    reopened = "\n".join(p.text for p in docx.Document(io.BytesIO(dnew)).paragraphs)
    assert "42" in reopened and "40" not in reopened


def test_P4_diff_is_recomputable_from_independent_cases():
    """Derive the four figures by hand from crafted pairs; compare to compute_*."""
    from forgeflow.documents import compute_pptx_diff, compute_textfile_diff

    # textfile: one modified line
    old = b"alpha\nbeta\ngamma\n"
    new = b"alpha\nBETA\ngamma\n"
    r = compute_textfile_diff(old, new)
    assert (r.modified, r.added, r.removed, r.numeric_changes) == (1, 0, 0, 0)
    # one removed line
    assert compute_textfile_diff(old, b"alpha\ngamma\n").removed == 1
    # one added line
    assert compute_textfile_diff(old, b"alpha\nbeta\ngamma\ndelta\n").added == 1
    # numeric-only change (same token set count but numbers differ)
    rn = compute_textfile_diff(b"value = 40\n300\n", b"value = 42\n300\n")
    assert (rn.modified, rn.numeric_changes) == (1, 1)
    # identical ⇒ all zero
    assert compute_textfile_diff(old, old).to_dict() == {
        "modified": 0, "added": 0, "removed": 0, "numeric_changes": 0,
    }

    # pptx: replace a single number in the title
    from forgeflow.documents.pptx_edit import apply_edits as apply_pptx_edits

    p_old = _pptx_bytes(title="总金额 40 元")
    p_new, _ = apply_pptx_edits(
        p_old, [{"op": "replace_text", "match": "40", "replace": "42"}]
    )
    pr = compute_pptx_diff(p_old, p_new)
    assert (pr.modified, pr.added, pr.removed, pr.numeric_changes) == (1, 0, 0, 1)


def test_P4_align_pairs_in_place_rewrites():
    from forgeflow.documents.textdiff import align_lines

    aligned = align_lines(["A", "B", "C"], ["A", "X", "C"])
    kinds = {k for k, _, _ in aligned}
    assert "pair" in kinds and "add" not in kinds and "remove" not in kinds


# =========================================================================== #
# 5. T02 adversarial EOL / encoding / BOM preservation                          #
# =========================================================================== #
def test_P5_adversarial_eol_encoding_bom_are_preserved():
    from forgeflow.documents import apply_textfile_edits, inspect_textfile

    # CRLF stays CRLF (no bare LF smuggled in)
    crlf = "第一行\r\n第二行\r\n".encode()
    out, _ = apply_textfile_edits(crlf, [{"op": "replace_text", "match": "第二行", "replace": "改动行"}])
    assert b"\r\n" in out and b"\n" not in out.replace(b"\r\n", b"")
    assert inspect_textfile(out).eol == "crlf"

    # UTF-8 BOM stays BOM (byte-identical BOM prefix)
    bom = b"\xef\xbb\xbf" + "行一\n行二\n".encode()
    out2, _ = apply_textfile_edits(bom, [{"op": "replace_text", "match": "行二", "replace": "行三"}])
    assert out2.startswith(b"\xef\xbb\xbf")
    assert inspect_textfile(out2).has_bom is True
    assert not out2[len(b"\xef\xbb\xbf"):].startswith(b"\xef\xbb\xbf")  # exactly one BOM

    # GB18030 stays GB18030 (still decodable as gb18030; no UTF-8 BOM)
    gb = "中文内容\n".encode("gb18030")
    assert inspect_textfile(gb).encoding == "gb18030"
    out3, _ = apply_textfile_edits(gb, [{"op": "replace_text", "match": "中文", "replace": "英文"}])
    assert out3.decode("gb18030") == "英文内容\n"
    assert not out3.startswith(b"\xef\xbb\xbf")

    # lossy decode is explicitly flagged (never hidden)
    lossy = inspect_textfile(b"\xff\xfe\xff\xfe")
    assert lossy.lossy_decode is True
    assert lossy.encoding == "latin-1(replace)"


# =========================================================================== #
# 6. T03 resource seam — pptx + text, additive & honest                         #
# =========================================================================== #
@pytest.fixture
def memory_res(tmp_path, monkeypatch):
    from forgeflow.repositories.memory.resource_repo import clear_resource_store
    from forgeflow.resources.service import reset_resource_index

    monkeypatch.setattr(
        __import__("forgeflow.config", fromlist=["get_settings"]).get_settings(),
        "resource_store_root", str(tmp_path / "blobs"),
    )
    reset_resource_index()
    clear_resource_store()
    yield
    reset_resource_index()
    clear_resource_store()


async def test_P6_pptx_kind_and_extension(memory_res):
    from forgeflow.resources import summaries

    assert ".pptx" in summaries.SUPPORTED_FILE_EXTENSIONS
    assert summaries.content_kind("slides.pptx") == "document"
    assert summaries.is_supported_file("slides.pptx") is True
    assert summaries.mime_for("slides.pptx").endswith("presentationml.presentation")


async def test_P6_pptx_document_paths_and_py_text_paths_paths_unchanged(memory_res,
                                                                        force_memory_backend):
    from forgeflow.resources.service import ResourceService

    svc = ResourceService()
    pptx = await svc.register_file(TENANT, name="deck.pptx", data=_pptx_bytes())
    assert pptx.status == "parsed"
    r1 = svc.resolve_task_inputs({"resources": [pptx.id]})
    assert r1.get("document_paths") and os.path.isfile(r1["document_paths"][0])
    assert r1["document_names"] == ["deck.pptx"]
    assert r1["paths"] == r1["document_paths"]  # item-for-item unchanged

    py = await svc.register_file(TENANT, name="app.py", data=b"x = 1\n")
    r2 = svc.resolve_task_inputs({"resources": [py.id]})
    assert r2.get("text_paths") and os.path.isfile(r2["text_paths"][0])
    assert r2["text_names"] == ["app.py"]
    assert r2["paths"] == r2["text_paths"]  # item-for-item unchanged
    assert "document_paths" not in r2


async def test_P6_missing_pptx_extra_degrades_without_fake_counts(monkeypatch):
    from forgeflow.resources import summaries

    data = _pptx_bytes()  # build before hiding the import
    monkeypatch.setitem(sys.modules, "pptx", None)
    status, summary, detail = summaries.summarize_presentation(data)
    assert status == "metadata_only"
    assert "pptx support unavailable" in detail
    assert summary.chars is None, "unmeasured char count must be None, never a fake 0"
    assert summary.rows is None

    status2, s2, d2 = summaries.summarize_bytes(data, filename="deck.pptx")
    assert status2 == "metadata_only" and "pptx support unavailable" in d2


# =========================================================================== #
# 7. T04 routing mutual exclusion (both counterfactuals)                        #
# =========================================================================== #
_CTX = None


def _rctx():
    from forgeflow.runtime.orchestrator import RequestContext

    return RequestContext(tenant_id=TENANT, user_id="qa", role="manager")


async def test_P7_py_is_textfile_not_analysis(memory_res, force_memory_backend):
    from forgeflow.resources.service import ResourceService
    from forgeflow.runtime import orchestrator as orch
    from forgeflow.runtime.orchestrator import TaskCreate

    py = await ResourceService().register_file(TENANT, name="app.py", data=b"import os\n")
    task = TaskCreate(intent="改代码", context={"resources": [py.id]})
    assert orch._is_textfile_task(task, _rctx()) is True
    assert orch._is_analysis_task(task, _rctx()) is False
    assert orch._is_document_task(task, _rctx()) is False


async def test_P7_csv_is_still_analysis(memory_res, force_memory_backend):
    from forgeflow.resources.service import ResourceService
    from forgeflow.runtime import orchestrator as orch
    from forgeflow.runtime.orchestrator import TaskCreate

    csv = await ResourceService().register_file(
        TENANT, name="leads.csv", data=b"id,amount\n1,10\n"
    )
    task = TaskCreate(intent="分析数据", context={"resources": [csv.id]})
    assert orch._is_analysis_task(task, _rctx()) is True
    assert orch._is_textfile_task(task, _rctx()) is False


async def test_P7_md_json_and_pptx_predicates():
    from forgeflow.runtime import orchestrator as orch
    from forgeflow.runtime.orchestrator import TaskCreate

    for name in ("notes.md", "config.json"):
        t = TaskCreate(intent="改文件", context={"paths": [f"/srv/{name}"]})
        assert orch._is_textfile_task(t, _rctx()) is True, name
        assert orch._is_analysis_task(t, _rctx()) is False, name
    t2 = TaskCreate(intent="改幻灯片", context={"paths": ["/srv/deck.pptx"]})
    assert orch._is_document_task(t2, _rctx()) is True
    assert orch._is_textfile_task(t2, _rctx()) is False
    assert orch._is_analysis_task(t2, _rctx()) is False


async def test_P7_counterfactual_textfile_exclusion_is_load_bearing(
    memory_res, force_memory_backend, monkeypatch
):
    """Remove the analysis-exclusion ⟹ the ``.py`` pin above would go RED."""
    from forgeflow.resources.service import ResourceService
    from forgeflow.runtime import orchestrator as orch
    from forgeflow.runtime.orchestrator import TaskCreate

    py = await ResourceService().register_file(TENANT, name="app.py", data=b"import os\n")
    task = TaskCreate(intent="改代码", context={"resources": [py.id]})

    monkeypatch.setattr(orch, "_is_textfile_task", lambda t, c: False)
    # Without the exclusion the .py FILE now looks like a data file ⇒ analysis.
    assert orch._is_analysis_task(task, _rctx()) is True, (
        "counterfactual failed: the textfile exclusion is not what keeps .py off analysis"
    )


async def test_P7_counterfactual_blanket_textfile_would_break_csv(
    memory_res, force_memory_backend, monkeypatch
):
    """Making ``_is_textfile_task`` always True ⟹ the CSV pin would go RED."""
    from forgeflow.resources.service import ResourceService
    from forgeflow.runtime import orchestrator as orch
    from forgeflow.runtime.orchestrator import TaskCreate

    csv = await ResourceService().register_file(
        TENANT, name="leads.csv", data=b"id,amount\n1,10\n"
    )
    task = TaskCreate(intent="分析数据", context={"resources": [csv.id]})

    monkeypatch.setattr(orch, "_is_textfile_task", lambda t, c: True)
    assert orch._is_analysis_task(task, _rctx()) is False, (
        "counterfactual failed: a blanket textfile predicate would mis-route a CSV"
    )


# =========================================================================== #
# 8. T04 catalogue ⇔ registry ⇔ grants; document names unchanged                #
# =========================================================================== #
def test_P8_whitelist_and_document_names_unchanged():
    from forgeflow.runtime import tool_registry
    from forgeflow.runtime.gate import (
        PLATFORM_PLAN_TOOLS,
        PLATFORM_TOOL_CATALOGUE,
        TOOL_PERMISSION_MAP,
    )

    tool_registry.load_default_bindings()  # explicit: known_ids() reads a global registry
    assert set(PLATFORM_PLAN_TOOLS) == set(tool_registry.known_ids()), "orphan/dead binding"
    for tool in ("textfile.inspect", "textfile.edit"):
        assert tool in PLATFORM_TOOL_CATALOGUE
        assert tool in PLATFORM_PLAN_TOOLS
        assert tool not in TOOL_PERMISSION_MAP
        assert tool_registry.resolve(tool) is not None
    # document.* names are unchanged (no presentation.* introduced)
    for name in ("document.inspect", "document.edit", "artifact.save"):
        assert name in PLATFORM_TOOL_CATALOGUE
    assert not any(t.startswith("presentation.") for t in PLATFORM_TOOL_CATALOGUE)


# =========================================================================== #
# 9. runs.py artifact download — :path fix, no traversal, regressions intact    #
# =========================================================================== #
def _run_record(run_id: str, tenant: str):
    from forgeflow.runtime.orchestrator import RunRecord

    return RunRecord(
        run_id=run_id, thread_id=f"{run_id}-th", tenant_id=tenant, agent_id=None,
        intent="编辑文件", status="completed", outcome="success", steps=[], errors=[],
        created_at="2026-10-04T00:00:00+00:00", completed_at="2026-10-04T00:00:01+00:00",
        session_id=run_id,
    )


@pytest.fixture
def run_env(tmp_path, monkeypatch):
    import forgeflow.documents as docs_pkg
    from forgeflow.documents import store as store_mod
    from forgeflow.runtime.orchestrator import reset_run_store

    real = store_mod.DocArtifactStore

    def _factory(root=None):
        return real(root=root if root is not None else tmp_path / "docstore")

    monkeypatch.setattr(store_mod, "DocArtifactStore", _factory)
    monkeypatch.setattr(docs_pkg, "DocArtifactStore", _factory)
    reset_run_store()
    yield
    reset_run_store()


def _client(tenant: str) -> TestClient:
    from forgeflow.api.hub_deps import resolve_tenant
    from forgeflow.api.routers import runs as runs_router

    app = FastAPI()
    app.dependency_overrides[resolve_tenant] = lambda: tenant
    app.include_router(runs_router.router, prefix="/runs")
    return TestClient(app)


def _seed(run_id: str, artifacts, tenant: str = TENANT):
    from forgeflow.runtime.orchestrator import get_run_store

    rec = _run_record(run_id, tenant)
    rec.artifacts = list(artifacts)
    get_run_store().save(rec)


def test_P9_route_uses_path_converter_and_slashed_id_downloads(run_env):
    from forgeflow.documents import DocArtifactStore

    blob = _pptx_bytes(title="下载标题 40")
    ref = DocArtifactStore().put(blob, ".pptx")
    assert "/" in ref, "content_ref must contain a slash for this test to bite"

    run_id = "run-qa-dl"
    art = {
        "id": f"{run_id}:artifact:{ref}:pptx",
        "kind": "document_pptx", "title": "d.edited.pptx", "format": "pptx",
        "content": "", "content_ref": ref, "source": "document.edit",
    }
    _seed(run_id, [art])
    client = _client(TENANT)
    resp = client.get(f"/runs/{run_id}/artifacts/{quote(art['id'], safe='')}")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].endswith("presentationml.presentation")
    assert resp.content == blob

    # the route really carries the :path converter
    from forgeflow.api.routers import runs as runs_router

    paths = [getattr(r, "path", "") for r in runs_router.router.routes]
    assert any(p.endswith("/artifacts/{artifact_id:path}") for p in paths)


def test_P9_no_traversal_no_cross_run_leak(run_env):
    from forgeflow.documents import DocArtifactStore

    ref_a = DocArtifactStore().put(b"AAAA", ".txt")
    ref_b = DocArtifactStore().put(b"BBBB", ".txt")
    run_a, run_b = "run-qa-a", "run-qa-b"
    _seed(run_a, [{"id": f"{run_a}:artifact:{ref_a}:text", "format": "text",
                   "content": "", "content_ref": ref_a}])
    _seed(run_b, [{"id": f"{run_b}:artifact:{ref_b}:text", "format": "text",
                   "content": "", "content_ref": ref_b}])
    client = _client(TENANT)

    # a traversal-shaped id can only ever resolve against the run's own artifacts
    for evil in ("../../etc/passwd", "..%2f..%2fetc%2fpasswd",
                 f"{run_b}:artifact:{ref_b}:text"):
        r = client.get(f"/runs/{run_a}/artifacts/{quote(evil, safe='')}")
        assert r.status_code == 404, (evil, r.status_code)
    # and the cross-run artifact is only reachable via its OWN run
    r_ok = client.get(f"/runs/{run_b}/artifacts/{quote(f'{run_b}:artifact:{ref_b}:text', safe='')}")
    assert r_ok.status_code == 200 and r_ok.content == b"BBBB"
    # a different tenant sees 404 (fail-closed)
    r_other = _client("t-other").get(
        f"/runs/{run_b}/artifacts/{quote(f'{run_b}:artifact:{ref_b}:text', safe='')}"
    )
    assert r_other.status_code == 404


def test_P9_legacy_markdown_download_unchanged(run_env):
    run_id = "run-qa-md"
    art = {
        "id": f"{run_id}:report:1", "kind": "report", "title": "report.md",
        "format": "markdown", "content": "# hello\n\nbody\n", "content_ref": "",
    }
    _seed(run_id, [art])
    resp = _client(TENANT).get(f"/runs/{run_id}/artifacts/{quote(art['id'], safe='')}")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/markdown")
    assert resp.content.decode("utf-8") == "# hello\n\nbody\n"


# REMOVED (2026-10-04, INC48 follow-up round). Formerly
# ``test_A_inc43_pin_diff_is_scoped`` (+ its ``_INC44_PIN_FLIP_COMMIT`` constant).
#
# Verdict: REMOVE — it asserted the *diff shape* of the already-landed commit
# ``6a83da0``. Once that commit was in history the patch it read could never
# change, so the case was a permanently-true snapshot: non-vacuous only during
# the uncommitted window it was written in, and never again a live regression
# guard (it could not fail for any change to current code).
#
# The invariant it was aiming at is behavioural and is still pinned, by
# ``tests/unit/test_inc43_docx_resource.py::test_pptx_is_a_document_and_xlsx_stays_a_table``
# (``.pptx`` is a document, ``.xlsx`` stays a table) and
# ``::test_limits_lists_docx_as_supported`` (``.pptx`` is advertised). Those
# cases fail on a real behaviour regression; this one could not.
