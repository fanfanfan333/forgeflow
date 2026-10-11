"""Verify a saved instruction skill against real TXT, Word and Excel files.

Uses configured development credentials; creates new verification resources and
runs, never overwrites user files. Creates the 文档清洗 example only if absent.
"""

from __future__ import annotations

import io
import json
import time
from pathlib import Path
from urllib.parse import quote

import httpx
from docx import Document
from openpyxl import Workbook, load_workbook

from forgeflow.config import get_settings

INSTRUCTIONS = """# 文档清洗

## 什么时候使用
用户上传 Word、文本或表格，希望清理文字或规范内容时使用。

## 处理方法
1. 先读取上传文件，按真实文件类型选择工具，说明发现的问题。
2. 清理行首行尾空白、连续的多余空格；只在有明确依据时修复重复或错误文字。
3. 不修改数字、金额、日期、姓名或公式，不删除有效内容。不确定的修改先询问用户。
4. 使用实际编辑工具生成文件，保存修改后的文件并返回下载结果。
5. 核对修改后的内容，简要列出改动。文件没有问题时说明检查结果，不声称作了不存在的修改。"""


def wait(client, run_id, timeout=240):
    deadline = time.monotonic() + timeout
    next_notice = time.monotonic() + 20
    while time.monotonic() < deadline:
        response = client.get(f"/runs/{run_id}")
        response.raise_for_status()
        record = response.json()
        if record["status"] not in {"running", "queued"}:
            return record
        if time.monotonic() >= next_notice:
            print(json.dumps({"run_id": run_id, "status": record["status"]}), flush=True)
            next_notice = time.monotonic() + 20
        time.sleep(1)
    raise TimeoutError(run_id)


def main():
    output = Path("_audit")
    output.mkdir(exist_ok=True)
    with httpx.Client(base_url="http://127.0.0.1:8010", timeout=30, trust_env=False) as client:
        login = client.post(
            "/auth/login",
            json={
                "user_id": "manager-1",
                "password": get_settings().dev_login_password.get_secret_value(),
            },
        )
        login.raise_for_status()
        client.headers["Authorization"] = "Bearer " + login.json()["access_token"]
        listed = client.get("/skills", params={"q": "文档清洗", "limit": 100})
        listed.raise_for_status()
        skill = next(
            (
                s
                for s in listed.json()["items"]
                if s["name"] == "文档清洗" and s["status"] == "published"
            ),
            None,
        )
        if skill is None:
            created = client.post(
                "/skills/instructions",
                json={
                    "name": "文档清洗",
                    "description": "整理文档中的多余空白、重复文字与格式问题，保留原始数据。",
                    "instructions": INSTRUCTIONS,
                    "example": "清洗这份文档，整理多余空格，保留数字和原意，返回修改后的文件。",
                },
            )
            created.raise_for_status()
            skill = created.json()
        print(
            json.dumps(
                {"skill_id": skill["id"], "version": skill["current_version"]}, ensure_ascii=False
            ),
            flush=True,
        )
        fixtures = [("txt", "客户：  张三\n金额：  100万元\n".encode())]
        doc = Document()
        doc.add_paragraph("客户：  张三")
        doc.add_paragraph("金额：  100万元")
        buffer = io.BytesIO()
        doc.save(buffer)
        fixtures.append(("docx", buffer.getvalue()))
        book = Workbook()
        book.active["A1"] = "客户"
        book.active["B1"] = "金额"
        book.active["A2"] = "  张三  "
        book.active["B2"] = 100
        book.active["B3"] = "=SUM(B2)"
        buffer = io.BytesIO()
        book.save(buffer)
        fixtures.append(("xlsx", buffer.getvalue()))
        for fmt, original in fixtures:
            uploaded = client.post(
                "/resources/files", files={"file": (f"skill-cleaning-verification.{fmt}", original)}
            )
            uploaded.raise_for_status()
            admitted = client.post(
                "/workspace/tasks",
                json={
                    "intent": "清洗这份文档，去掉多余空格，保留所有姓名、数字、金额和公式，返回处理后的文件。",
                    "context": {"resources": [uploaded.json()["id"]], "skill_ids": [skill["id"]]},
                },
            )
            admitted.raise_for_status()
            record = wait(client, admitted.json()["run_id"])
            artifact = next(
                (
                    a
                    for a in record.get("artifacts", [])
                    if a.get("format") == ("text" if fmt == "txt" else fmt)
                ),
                None,
            )
            if record["status"] != "completed" or artifact is None:
                print(
                    json.dumps(
                        {
                            "format": fmt,
                            "status": record["status"],
                            "errors": record.get("errors"),
                            "answer": record.get("llm", {}).get("final_answer"),
                        },
                        ensure_ascii=False,
                    ),
                    flush=True,
                )
                raise AssertionError(f"No completed {fmt} delivery")
            response = client.get(
                f"/runs/{record['run_id']}/artifacts/{quote(artifact['id'], safe='')}"
            )
            response.raise_for_status()
            data = response.content
            if fmt == "docx":
                text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
            elif fmt == "xlsx":
                result = load_workbook(io.BytesIO(data))
                assert result.active["A2"].value == "张三"
                assert result.active["B2"].value == 100
                assert result.active["B3"].value == "=SUM(B2)"
                result.close()
                text = "张三 100"
            else:
                text = data.decode("utf-8-sig")
            assert "张三" in text and "100" in text and "  " not in text
            assert skill["id"] in record["skills_used"]
            assert record["declared_inputs"]["skill_ids"] == [skill["id"]]
            (output / f"skill-cleaning-result.{fmt}").write_bytes(data)
            print(
                json.dumps(
                    {
                        "format": fmt,
                        "run_id": record["run_id"],
                        "cleaning": "passed",
                        "data_preserved": True,
                        "skill_used": True,
                    },
                    ensure_ascii=False,
                ),
                flush=True,
            )


if __name__ == "__main__":
    main()
