"""Exercise real compound tasks in a disposable QA tenant, checking downloaded bytes."""

from __future__ import annotations

import asyncio
import io
import json
import secrets
import sys
import time
import uuid
from pathlib import Path
from urllib.parse import quote

import httpx
from docx import Document
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def fixtures():
    code = b'#include <stdio.h>\nint helper(void) { return 7; }\nint main(void)\n{\n    printf("hello word");\n}\n'
    doc = Document()
    doc.add_heading("  项目验收  ", 0)
    doc.add_paragraph("客户：  张三   日期：2026-10-15")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "  项目  "
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "  北辰  项目  "
    table.cell(1, 1).text = "12345.67"
    buffer = io.BytesIO()
    doc.save(buffer)
    word = buffer.getvalue()
    book = Workbook()
    sheet = book.active
    sheet.title = "销售"
    sheet.append(["客户", "金额", "税率"])
    sheet.append(["  张三  ", 12345.67, 0.06])
    sheet.append(["  李四  ", 890, 0.13])
    sheet["B4"] = "=SUM(B2:B3)"
    sheet["D2"] = "=B2*C2"
    sheet["B2"].number_format = "0.00"
    book.create_sheet("说明")["A1"] = "  保留该工作表  "
    buffer = io.BytesIO()
    book.save(buffer)
    return [
        (
            "compound.c.txt",
            code,
            "将 hello word 修正为 hello world，在 main 函数结束前添加 return 0;，保留 helper 函数的 return 7;，返回修改后的完整文件。",
            "code",
        ),
        (
            "compound.docx",
            word,
            "清洗这份 Word 文档中的多余空格，保留全部姓名、日期、金额以及表格结构，返回修改后的 Word 文件。",
            "word",
        ),
        (
            "compound.xlsx",
            buffer.getvalue(),
            "清洗这份 Excel 文件的文字首尾空格，保留所有数字、公式、工作表和数字格式，返回修改后的 Excel 文件。",
            "excel",
        ),
    ]


def verify(kind, data):
    if kind == "code":
        text = data.decode("utf-8-sig")
        main = text.split("int main(void)", 1)[1]
        return {
            "literal_corrected": 'printf("hello world")' in text,
            "helper_preserved": "int helper(void) { return 7; }" in text,
            "return_inside_main": "return 0;" in main
            and main.index("return 0;") < main.rindex("}"),
            "balanced_braces": text.count("{") == text.count("}"),
        }
    if kind == "word":
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        return {
            "opens": True,
            "identity_date_preserved": "张三" in text and "2026-10-15" in text,
            "table_preserved": len(doc.tables) == 1
            and len(doc.tables[0].rows) == 2
            and len(doc.tables[0].columns) == 2,
            "amount_preserved": doc.tables[0].cell(1, 1).text == "12345.67",
            "whitespace_cleaned": doc.paragraphs[0].text == "项目验收"
            and doc.tables[0].cell(0, 0).text == "项目",
        }
    book = load_workbook(io.BytesIO(data), data_only=False)
    sheet = book["销售"]
    return {
        "opens": True,
        "sheets_preserved": book.sheetnames == ["销售", "说明"],
        "values_preserved": sheet["B2"].value == 12345.67
        and sheet["B3"].value == 890
        and sheet["C2"].value == 0.06,
        "formulas_preserved": sheet["B4"].value == "=SUM(B2:B3)" and sheet["D2"].value == "=B2*C2",
        "number_format_preserved": sheet["B2"].number_format == "0.00",
        "other_sheet_cleaned": book["说明"]["A1"].value == "保留该工作表",
        "whitespace_cleaned": sheet["A2"].value == "张三" and sheet["A3"].value == "李四",
    }


async def main():
    from forgeflow.auth.passwords import hash_password
    from forgeflow.auth.users import get_by_username, upsert_local_user
    from forgeflow.database import close_pool, init_pool
    from forgeflow.tenant_scope import tenant_scope

    output = Path("docs/quality/evidence/2026-10-15/complex-tasks")
    output.mkdir(parents=True, exist_ok=True)
    identity = "qa-complex-" + uuid.uuid4().hex
    password = secrets.token_urlsafe(32)
    pool = await init_pool()
    results = []
    with tenant_scope(identity):
        assert await get_by_username(pool, identity) is None
        user = await upsert_local_user(
            pool, identity, hash_password(password), "manager", tenant_id=identity
        )
    try:
        async with httpx.AsyncClient(
            base_url="http://127.0.0.1:8010", trust_env=False, timeout=30
        ) as client:
            login = await client.post(
                "/auth/login",
                json={"user_id": identity, "password": password, "workspace_id": identity},
            )
            login.raise_for_status()
            client.headers["Authorization"] = "Bearer " + login.json()["access_token"]
            for name, original, intent, kind in fixtures():
                result = {"case": kind, "intent": intent}
                started = time.monotonic()
                try:
                    upload = await client.post("/resources/files", files={"file": (name, original)})
                    upload.raise_for_status()
                    admitted = await client.post(
                        "/workspace/tasks",
                        json={"intent": intent, "context": {"resources": [upload.json()["id"]]}},
                    )
                    admitted.raise_for_status()
                    run_id = admitted.json()["run_id"]
                    result["run_id"] = run_id
                    deadline = time.monotonic() + 300
                    while True:
                        response = await client.get("/runs/" + run_id)
                        response.raise_for_status()
                        run = response.json()
                        if run["status"] not in {"running", "queued", "pending"}:
                            break
                        if time.monotonic() >= deadline:
                            await client.post("/runs/" + run_id + "/abort")
                            raise TimeoutError("QA task exceeded 300 seconds; abort requested")
                        await asyncio.sleep(2)
                    result.update(
                        status=run["status"],
                        outcome=run.get("outcome"),
                        acceptance=run.get("llm", {}).get("acceptance"),
                        errors=run.get("errors"),
                    )
                    artifacts = [
                        a
                        for a in run.get("artifacts", [])
                        if a.get("content_ref") and a.get("kind") != "report_markdown"
                    ]
                    checks = []
                    for artifact in artifacts:
                        download = await client.get(
                            "/runs/" + run_id + "/artifacts/" + quote(artifact["id"], safe="")
                        )
                        download.raise_for_status()
                        try:
                            checks.append(verify(kind, download.content))
                            (output / (kind + "-download" + Path(name).suffix)).write_bytes(
                                download.content
                            )
                        except Exception as exc:
                            checks.append({"opens": False, "error": str(exc)[:200]})
                    result["artifact_checks"] = checks
                    result["passed"] = (
                        run["status"] == "completed"
                        and run.get("outcome") == "success"
                        and (result.get("acceptance") or {}).get("complete") is True
                        and bool(checks)
                        and all(all(v is True for v in check.values()) for check in checks)
                    )
                except Exception as exc:
                    result.update(passed=False, error=str(exc)[:300])
                result["seconds"] = round(time.monotonic() - started, 2)
                results.append(result)
                (output / "results.json").write_text(
                    json.dumps(
                        {
                            "tenant": identity,
                            "model": "configured Ollama default",
                            "cases": results,
                        },
                        ensure_ascii=False,
                        indent=2,
                    ),
                    encoding="utf-8",
                )
                print(json.dumps(result, ensure_ascii=False), flush=True)
    finally:
        with tenant_scope(identity):
            async with pool.acquire() as conn:
                await conn.execute("UPDATE auth_users SET disabled=TRUE WHERE id=$1", user["id"])
        await close_pool()
    return 0 if len(results) == 3 and all(r["passed"] for r in results) else 1


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
