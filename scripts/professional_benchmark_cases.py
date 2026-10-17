"""Independent synthetic regression inputs; never register these as human golden data."""

import io


def cases():
    from docx import Document
    from openpyxl import Workbook

    items = []
    code_cases = [
        (
            "compound",
            'int helper(void) { return 7; }\nint main(void) { printf("hello word"); }',
            "将 hello word 修正为 hello world，在 main 函数内增加 return 0;，保留 helper 的 return 7;",
            ["hello world", "return 7;"],
        ),
        (
            "return_only",
            'int main(void) { printf("ready"); }',
            "在 main 函数结束前增加 return 0;，保留输出文字",
            ["ready"],
        ),
        (
            "legacy_main",
            'main() { printf("hello word"); }',
            "优化代码，将 hello word 修改为 hello world，补上 main 函数内的 return 0;",
            ["int main(void)", "hello world"],
        ),
        (
            "helper_after",
            'int main(void) { printf("hello word"); }\nint helper(void) { return 3; }',
            "将 hello word 修正为 hello world，在 main 内增加 return 0;，保留 helper 的 return 3;",
            ["hello world", "return 3;"],
        ),
        (
            "brace_string",
            'int main(void) { printf("hello word }"); }',
            "将 hello word 修改为 hello world，并在 main 函数内增加 return 0;，保留字符串中的大括号",
            ["hello world }"],
        ),
        (
            "comment_return",
            'int main(void) { /* return 0; is missing */ printf("hello word"); }',
            "将 hello word 修正为 hello world，并在 main 函数结束前增加真正的 return 0;，保留注释",
            ["hello world", "/* return 0; is missing */"],
        ),
        (
            "existing_return",
            'int main(void) { printf("hello word"); return 0; }',
            "将 hello word 修正为 hello world，保留 main 内已有的 return 0;，不要重复添加",
            ["hello world"],
        ),
        (
            "two_literals",
            'int main(void) { printf("hello word"); printf("sucess"); }',
            "将 hello word 修正为 hello world，将 sucess 修改为 success，并在 main 函数内增加 return 0;",
            ["hello world", "success"],
        ),
        (
            "nested_block",
            'int main(void) { if (1) { printf("hello word"); } }',
            "将 hello word 修正为 hello world，并在 main 的最外层结束大括号前增加 return 0;，保留 if 分支",
            ["hello world", "if (1)"],
        ),
        (
            "crlf",
            'int main(void) {\n    printf("hello word");\n}',
            "将 hello word 修正为 hello world，在 main 函数内增加 return 0;，保留其余代码",
            ["hello world"],
        ),
    ]
    for name, body, intent, literals in code_cases:
        text = "#include <stdio.h>\n" + body + "\n"
        if name == "crlf":
            text = text.replace("\n", "\r\n")
        items.append(
            {
                "id": "code-" + name,
                "domain": "code",
                "files": [(name + ".txt", text.encode())],
                "intent": intent + "，返回修改后的完整文件。",
                "assertions": [{"kind": "main_return"}]
                + [{"kind": "contains", "value": literal} for literal in literals],
            }
        )
    for index in range(4):
        doc = Document()
        doc.add_heading("  验收项目  ", 0)
        doc.add_paragraph("姓名：  张三   日期：2026-10-15")
        table = doc.add_table(rows=2 + index, cols=2)
        table.cell(0, 0).text = "  项目  "
        table.cell(0, 1).text = "金额"
        for row in range(1, 2 + index):
            table.cell(row, 0).text = "  客户  条目  "
            table.cell(row, 1).text = str(12345 + row) + ".67"
        out = io.BytesIO()
        doc.save(out)
        items.append(
            {
                "id": f"document-word-{index}",
                "domain": "document",
                "files": [(f"word-{index}.docx", out.getvalue())],
                "intent": "清洗这份 Word 文档的多余空格，保留全部姓名、日期、金额和表格，返回修改后的 Word 文件。",
                "assertions": [
                    {"kind": "docx_contains", "format": "docx", "value": value}
                    for value in ["验收项目", "张三", "2026-10-15", "客户 条目", "12346.67"]
                ],
            }
        )
    for index in range(3):
        book = Workbook()
        sheet = book.active
        sheet.title = "销售"
        sheet.append(["客户", "金额"])
        sheet.append(["  张三  ", 12345.67 + index])
        sheet["B3"] = "=SUM(B2:B2)"
        sheet["B2"].number_format = "0.00"
        book.create_sheet("说明")["A1"] = "  保留说明  "
        out = io.BytesIO()
        book.save(out)
        items.append(
            {
                "id": f"document-excel-{index}",
                "domain": "document",
                "files": [(f"sheet-{index}.xlsx", out.getvalue())],
                "intent": "清洗这份 Excel 文件的文字首尾空格，保留所有数字、公式、工作表和数字格式，返回修改后的 Excel 文件。",
                "assertions": [
                    {
                        "kind": "xlsx_cell",
                        "format": "xlsx",
                        "sheet": "销售",
                        "cell": "A2",
                        "value": "张三",
                    },
                    {
                        "kind": "xlsx_cell",
                        "format": "xlsx",
                        "sheet": "销售",
                        "cell": "B2",
                        "value": 12345.67 + index,
                    },
                    {
                        "kind": "xlsx_cell",
                        "format": "xlsx",
                        "sheet": "销售",
                        "cell": "B3",
                        "value": "=SUM(B2:B2)",
                    },
                    {
                        "kind": "xlsx_cell",
                        "format": "xlsx",
                        "sheet": "说明",
                        "cell": "A1",
                        "value": "保留说明",
                    },
                    {
                        "kind": "xlsx_format",
                        "format": "xlsx",
                        "sheet": "销售",
                        "cell": "B2",
                        "value": "0.00",
                    },
                ],
            }
        )
    for index in range(3):
        text = f"  客户：  张三  \n  金额：  {12345 + index}.67  \n  编号：  A-000{index}  \n"
        items.append(
            {
                "id": f"document-text-{index}",
                "domain": "document",
                "files": [(f"text-{index}.txt", text.encode())],
                "intent": "清洗文本的多余空格，保留全部姓名、金额和编号，返回修改后的文本文件。",
                "assertions": [
                    {"kind": "contains", "value": value}
                    for value in ["客户： 张三", f"{12345 + index}.67", f"A-000{index}"]
                ],
            }
        )
    sources = [
        (
            "json",
            "https://docs.python.org/3/library/json.html",
            "解释 loads 和 dumps 的用途",
            "loads",
        ),
        ("pathlib", "https://docs.python.org/3/library/pathlib.html", "解释 Path 的作用", "Path"),
        (
            "asyncio",
            "https://docs.python.org/3/library/asyncio-task.html",
            "解释任务取消与 CancelledError",
            "CancelledError",
        ),
        (
            "sqlite",
            "https://docs.python.org/3/library/sqlite3.html",
            "解释参数绑定的用途",
            "sqlite3",
        ),
        (
            "csv",
            "https://docs.python.org/3/library/csv.html",
            "解释 DictReader 的用途",
            "DictReader",
        ),
        (
            "unittest",
            "https://docs.python.org/3/library/unittest.html",
            "解释 TestCase 的用途",
            "TestCase",
        ),
        (
            "hashlib",
            "https://docs.python.org/3/library/hashlib.html",
            "解释 sha256 的用途",
            "sha256",
        ),
        (
            "http",
            "https://www.rfc-editor.org/rfc/rfc9110.html",
            "解释 HTTP 的请求与响应模型",
            "HTTP",
        ),
        ("worktree", "https://git-scm.com/docs/git-worktree", "解释 worktree 的用途", "worktree"),
        (
            "rls",
            "https://www.postgresql.org/docs/current/ddl-rowsecurity.html",
            "解释行级安全策略的用途",
            "POLICY",
        ),
    ]
    for name, url, question, term in sources:
        items.append(
            {
                "id": "research-" + name,
                "domain": "research",
                "files": [],
                "intent": f"研究原始来源 {url}，{question}。实际读取原文，交付带精确原文引用和来源链接的简短报告，并校验引用与结论；资料不支持时明确说明。",
                "assertions": [
                    {
                        "kind": "research_verified",
                        "format": "markdown",
                        "value": url.split("https://", 1)[1].split(".html")[0],
                    },
                    {"kind": "contains", "format": "markdown", "value": term},
                ],
            }
        )
    assert len(items) == 30 and all(
        sum(c["domain"] == domain for c in items) == 10
        for domain in ("code", "document", "research")
    )
    return items
