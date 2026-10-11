"""Verify the live local API, model execution and Word upload/edit/download.

Uses configured development credentials without printing passwords or tokens.
Writes only new probe runs/resources; never changes an existing user document.
"""

from __future__ import annotations

import argparse
import io
import json
import time
from urllib.parse import quote

import httpx
from docx import Document

from forgeflow.config import get_settings


def wait_for_run(client: httpx.Client, run_id: str, timeout: float) -> dict:
    deadline = time.monotonic() + timeout
    next_progress = time.monotonic() + 20
    while time.monotonic() < deadline:
        response = client.get(f"/runs/{run_id}")
        response.raise_for_status()
        record = response.json()
        if record.get("status") in {"completed", "failed", "aborted", "interrupted", "rejected"}:
            return record
        if time.monotonic() >= next_progress:
            print(json.dumps({"run_id": run_id, "status": record.get("status")}), flush=True)
            next_progress = time.monotonic() + 20
        time.sleep(1)
    raise TimeoutError(f"Run did not settle within {timeout}s: {run_id}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8010")
    parser.add_argument("--user", default="admin")
    parser.add_argument("--timeout", type=float, default=180)
    parser.add_argument(
        "--natural-language",
        action="store_true",
        help="Resolve document edits through the real model",
    )
    args = parser.parse_args()
    settings = get_settings()
    with httpx.Client(base_url=args.base_url, timeout=30, trust_env=False) as client:
        health = client.get("/health")
        health.raise_for_status()
        print(json.dumps({"health": health.json()}, ensure_ascii=False), flush=True)
        login = client.post(
            "/auth/login",
            json={
                "user_id": args.user,
                "password": settings.dev_login_password.get_secret_value(),
            },
        )
        if login.status_code != 200:
            print(json.dumps({"login": "failed", "status_code": login.status_code}), flush=True)
            return 1
        client.headers["Authorization"] = "Bearer " + login.json()["access_token"]
        model = client.get("/runs/model-status")
        model.raise_for_status()
        print(json.dumps({"configured_models": model.json()}, ensure_ascii=False), flush=True)
        response = client.post(
            "/workspace/tasks",
            json={
                "intent": "请用一句中文说明你能协助完成什么工作。这是连通性验证，不需要搜索，也不要调用外部工具。",
            },
        )
        response.raise_for_status()
        first = wait_for_run(client, response.json()["run_id"], args.timeout)
        final = (first.get("llm") or {}).get("final_answer") or {}
        if first.get("status") != "completed" or not final.get("text"):
            print(
                json.dumps(
                    {
                        "model_execution": "failed",
                        "status": first.get("status"),
                        "errors": first.get("errors", []),
                    },
                    ensure_ascii=False,
                ),
                flush=True,
            )
            return 1
        print(
            json.dumps(
                {
                    "model_execution": "passed",
                    "runtime_mode": first.get("runtime_mode"),
                    "run_id": first["run_id"],
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
        document = Document()
        document.add_paragraph("文档验收：旧名称需要更新。")
        document.add_paragraph("销售额：100万元。")
        buffer = io.BytesIO()
        document.save(buffer)
        uploaded = client.post(
            "/resources/files",
            files={
                "file": (
                    "workspace-verification.docx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            },
        )
        uploaded.raise_for_status()
        resource_id = uploaded.json()["id"]
        edit_context = {"resources": [resource_id]}
        if not args.natural_language:
            edit_context["edits"] = [
                {"op": "replace_text", "match": "旧名称", "replace": "新名称"},
                {"op": "replace_text", "match": "100万元", "replace": "120万元"},
            ]
        response = client.post(
            "/workspace/tasks",
            json={
                "intent": "请读取已附加的 Word 文档，将旧名称替换为新名称，将销售额从100万元修改为120万元（这是用户明确要求的数据修正），其他内容保留，保存修改后的文件并返回下载产物。",
                "context": edit_context,
            },
        )
        response.raise_for_status()
        record = wait_for_run(client, response.json()["run_id"], args.timeout)
        artifact = next(
            (item for item in record.get("artifacts", []) if item.get("format") == "docx"), None
        )
        if artifact is None:
            print(
                json.dumps(
                    {
                        "document_delivery": "failed",
                        "run_id": record["run_id"],
                        "status": record.get("status"),
                        "errors": record.get("errors", []),
                    },
                    ensure_ascii=False,
                ),
                flush=True,
            )
            return 1
        download = client.get(
            f"/runs/{record['run_id']}/artifacts/{quote(artifact['id'], safe='')}"
        )
        download.raise_for_status()
        texts = "\n".join(
            paragraph.text for paragraph in Document(io.BytesIO(download.content)).paragraphs
        )
        if (
            "新名称" not in texts
            or "旧名称" in texts
            or "120万元" not in texts
            or "100万元" in texts
        ):
            print(
                json.dumps(
                    {"document_delivery": "failed", "reason": "downloaded content unchanged"}
                ),
                flush=True,
            )
            return 1
        print(
            json.dumps(
                {
                    "document_delivery": "passed",
                    "run_id": record["run_id"],
                    "download_bytes": len(download.content),
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
