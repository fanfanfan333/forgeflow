"""Non-executable acceptance assertions over actual downloaded artifacts."""

from __future__ import annotations

import io
import re
from typing import Any


def validate_assertions(assertions: list[dict]) -> None:
    if not isinstance(assertions, list) or not 1 <= len(assertions) <= 30:
        raise ValueError("An evaluation needs 1–30 explicit artifact assertions")
    for assertion in assertions:
        if not isinstance(assertion, dict):
            raise ValueError("Acceptance assertions must be objects")
        kind, value = assertion.get("kind"), assertion.get("value")
        if kind not in {
            "contains",
            "absent",
            "main_return",
            "docx_contains",
            "xlsx_cell",
            "xlsx_format",
            "xlsx_sheets",
            "research_verified",
            "code_definition",
            "report_answer_contains",
        }:
            raise ValueError(f"Unsupported acceptance assertion: {kind}")
        if not isinstance(assertion.get("format", "text"), str):
            raise ValueError("Artifact format must be a string")
        if kind in {
            "contains",
            "absent",
            "docx_contains",
            "research_verified",
            "report_answer_contains",
        } and (not isinstance(value, str) or not 1 <= len(value) <= 12000):
            raise ValueError("Text assertions require a bounded nonempty literal")
        if kind == "code_definition" and (not isinstance(value, str) or not value.isidentifier()):
            raise ValueError("Code definition assertion requires an identifier")
        if kind in {"xlsx_cell", "xlsx_format"} and (
            not isinstance(assertion.get("sheet"), str)
            or not assertion["sheet"]
            or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", str(assertion.get("cell", "")))
        ):
            raise ValueError("Spreadsheet assertions require a sheet and cell coordinate")
        if kind == "xlsx_sheets" and (
            not isinstance(value, list)
            or not value
            or any(not isinstance(name, str) or not name for name in value)
        ):
            raise ValueError("Spreadsheet sheets assertion requires sheet names")


def check_artifacts(
    artifacts: list[tuple[str, bytes]],
    assertions: list[dict],
    invocations: list[dict] | None = None,
) -> dict[str, bool]:
    validate_assertions(assertions)
    outcomes = {}
    for index, assertion in enumerate(assertions):
        kind = assertion.get("kind")
        expected = assertion.get("value")
        fmt = assertion.get("format", "text")
        formats = {fmt}
        if kind == "research_verified" and "format" not in assertion:
            formats = {"text", "markdown", "md"}
        if kind == "code_definition":
            if not isinstance(expected, str) or not expected.isidentifier():
                raise ValueError("Code definition assertion requires an identifier")
            outcomes[f"{index}:{kind}"] = any(
                record.get("tool") == "code.lint"
                and record.get("status") == "ok"
                and record.get("executed") is True
                and isinstance(record.get("payload"), dict)
                and record["payload"].get("symbol") == expected
                and any(
                    match.get("kind") == "definition"
                    for match in record["payload"].get("matches", [])
                )
                for record in (invocations or [])
            )
            continue
        candidates = [data for actual, data in artifacts if actual in formats]
        passed = False
        for data in candidates:
            if kind == "report_answer_contains":
                answer = re.search(
                    r"(?ms)^## 最终答案\s*\n(.*?)(?=^## |\Z)", data.decode("utf-8-sig")
                )
                passed = answer is not None and expected in answer.group(1)
            elif kind in {"contains", "absent"}:
                if not isinstance(expected, str) or not expected:
                    raise ValueError("Text assertions require a nonempty literal")
                matched = expected in data.decode("utf-8-sig")
                passed = matched if kind == "contains" else not matched
            elif kind == "main_return":
                from forgeflow.documents.textfile_edit import main_return_satisfied

                passed = main_return_satisfied(data)
            elif kind == "docx_contains":
                from docx import Document

                doc = Document(io.BytesIO(data))
                text = "\n".join(
                    [p.text for p in doc.paragraphs]
                    + [
                        cell.text
                        for table in doc.tables
                        for row in table.rows
                        for cell in row.cells
                    ]
                )
                passed = isinstance(expected, str) and bool(expected) and expected in text
            elif kind in {"xlsx_cell", "xlsx_format", "xlsx_sheets"}:
                from openpyxl import load_workbook

                book = load_workbook(io.BytesIO(data), data_only=False)
                try:
                    if kind == "xlsx_sheets":
                        passed = book.sheetnames == expected
                    else:
                        cell = book[assertion["sheet"]][assertion["cell"]]
                        passed = (
                            cell.value if kind == "xlsx_cell" else cell.number_format
                        ) == expected
                finally:
                    book.close()
            elif kind == "research_verified":
                verified = any(
                    r.get("tool") == "research.verify"
                    and r.get("status") == "ok"
                    and (r.get("payload") or {}).get("validation", {}).get("claims_supported")
                    is True
                    and (r.get("payload") or {}).get("validation", {}).get("citations_ok") is True
                    for r in (invocations or [])
                )
                passed = (
                    verified and isinstance(expected, str) and expected in data.decode("utf-8-sig")
                )
            else:
                raise ValueError(f"Unsupported acceptance assertion: {kind}")
            if passed:
                break
        outcomes[f"{index}:{kind}"] = bool(passed)
    return outcomes


def run_verdict(record: Any, checks: dict[str, bool]) -> bool:
    status = record.get("status") if isinstance(record, dict) else record.status
    outcome = record.get("outcome") if isinstance(record, dict) else record.outcome
    return status == "completed" and outcome == "success" and bool(checks) and all(checks.values())
