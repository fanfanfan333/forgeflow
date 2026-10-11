"""Durable single-worker synchronization to the independently deployed OV server."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import logging
import os
import re
import tempfile
import threading
import time
from collections.abc import Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

import httpx

from forgeflow.config import get_settings

logger = logging.getLogger(__name__)
_lock = asyncio.Lock()
_journal_replace_lock = threading.Lock()


def _save(path: Path, job: dict[str, Any]) -> None:
    temp: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as file:
            temp = Path(file.name)
            json.dump(job, file, ensure_ascii=False)
            file.flush()
            os.fsync(file.fileno())
        # Windows rename operations targeting the same file must not overlap.
        with _journal_replace_lock:
            os.replace(temp, path)
    finally:
        if temp is not None:
            temp.unlink(missing_ok=True)


async def enqueue(
    tenant: str | None, user: str, kind: str, object_id: str, payload: dict[str, Any]
) -> bool:
    settings = get_settings()
    endpoint = urlsplit(settings.openviking_url)
    if (
        endpoint.scheme not in {"http", "https"}
        or not endpoint.hostname
        or endpoint.username
        or endpoint.password
        or endpoint.query
        or endpoint.fragment
    ):
        return False
    if endpoint.scheme == "http" and endpoint.hostname not in {"127.0.0.1", "localhost", "::1"}:
        return False
    if not user or any(c in user for c in "/\\%?#") or user in {".", ".."}:
        return False
    if (
        not settings.openviking_sync_enabled
        or not settings.openviking_url
        or f"{tenant}:{user}" not in settings.openviking_user_keys
    ):
        return False
    identity = hashlib.sha256(f"{tenant}:{user}:{kind}:{object_id}".encode()).hexdigest()
    root = Path(settings.openviking_outbox_root)
    root.mkdir(parents=True, exist_ok=True)
    job = {
        "tenant": tenant,
        "user": user,
        "kind": kind,
        "id": object_id,
        "payload": payload,
        "status": "pending",
        "attempts": 0,
        "updated_at": datetime.now(UTC).isoformat(),
    }
    async with _lock:
        job_path = root / (identity + ".json")
        if job_path.exists():
            previous = json.loads(job_path.read_text(encoding="utf-8"))
            if previous["payload"] == payload and previous["status"] in {
                "pending",
                "waiting",
                "done",
                "ready",
            }:
                return True
            job["history"] = (
                previous.get("history", [])
                + [
                    {
                        key: previous.get(key)
                        for key in ("status", "attempts", "updated_at", "error", "result")
                    }
                ]
            )[-10:]
        if not kind.startswith("delete_"):
            delete_identity = hashlib.sha256(
                f"{tenant}:{user}:delete_{kind}:{object_id}".encode()
            ).hexdigest()
            delete_path = root / (delete_identity + ".json")
            if delete_path.exists():
                tombstone = json.loads(delete_path.read_text(encoding="utf-8"))
                if tombstone["status"] != "done":
                    tombstone["status"] = "cancelled"
                    _save(delete_path, tombstone)
        if kind.startswith("delete_"):
            previous = hashlib.sha256(
                f"{tenant}:{user}:{kind.removeprefix('delete_')}:{object_id}".encode()
            ).hexdigest()
            old_path = root / (previous + ".json")
            if old_path.exists():
                old = json.loads(old_path.read_text(encoding="utf-8"))
                if old.get("result", {}).get("task_id"):
                    payload["task_ids"] = [old["result"]["task_id"]]
                old["status"] = "cancelled"
                _save(old_path, old)
        _save(job_path, job)
    return True


def _text(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        return "\n".join(
            [p.text for p in doc.paragraphs]
            + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        )
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            return "\n".join(
                f"{sheet.title}: "
                + " | ".join(str(c.value) if c.value is not None else "" for c in row)
                for sheet in workbook
                for row in sheet.iter_rows(max_row=1000)
            )
        finally:
            workbook.close()
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages[:100])
    if suffix == ".pptx":
        from pptx import Presentation

        return "\n".join(
            shape.text
            for slide in Presentation(io.BytesIO(data)).slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
    from forgeflow.resources.summaries import _decode

    return _decode(data)[0]


def _registered_code(record: Any) -> str:
    from forgeflow.resources.code_sources import _SKIP_DIRS, EXTENSION_LANGUAGE

    root = Path(record.locator.identifier).resolve(strict=True)
    chunks = []
    size = 0
    for directory, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = sorted(
            d
            for d in dirs
            if d not in _SKIP_DIRS
            and not d.startswith(".")
            and not (Path(directory) / d).is_symlink()
        )
        for name in sorted(files):
            path = Path(directory) / name
            if (
                path.suffix.lower() not in EXTENSION_LANGUAGE
                or name.startswith(".")
                or path.is_symlink()
            ):
                continue
            if root not in path.resolve().parents or path.stat().st_size > 65536:
                continue
            chunk = (
                f"\n## {path.relative_to(root).as_posix()}\n"
                + _text(path.read_bytes(), name)[:6000]
            )
            chunks.append(chunk)
            size += len(chunk)
            if size >= 180000 or len(chunks) >= 100:
                return "".join(chunks) + "\n[代码索引达到上限，完整源代码保持不变]"
    return "".join(chunks)


async def queue_resource(service: Any, record: Any) -> None:
    if not get_settings().openviking_sync_enabled:
        return
    try:
        if record.kind == "git_repo" and record.locator.source_type == "local_path":
            content = await asyncio.to_thread(_registered_code, record)
        elif record.kind != "file":
            content = json.dumps(
                {"name": record.name, "kind": record.kind, "id": record.id}, ensure_ascii=False
            )
        else:
            ref = record.locator.storage_ref
            content = await asyncio.to_thread(_text, service.blobs.read(ref), record.name)
        truncated = len(content) > 200000
        await enqueue(
            record.tenant_id,
            record.created_by,
            "resource",
            record.id,
            {
                "name": record.name,
                "content": content[:200000]
                + ("\n[索引内容达到上限，原文件保持完整]" if truncated else ""),
                "truncated": truncated,
            },
        )
    except Exception as exc:
        logger.warning("OpenViking resource enqueue failed: %s", type(exc).__name__)
        if record.kind == "file":
            await enqueue(
                record.tenant_id,
                record.created_by,
                "resource",
                record.id,
                {
                    "name": record.name,
                    "source_ref": record.locator.storage_ref,
                    "conversion_error": type(exc).__name__,
                },
            )


async def queue_skill(tenant: str, user: str, skill: Any, spec: dict[str, Any]) -> None:
    if not get_settings().openviking_sync_enabled:
        return
    content = str(
        spec.get("instructions") or spec.get("prompt") or json.dumps(spec, ensure_ascii=False)
    )
    for identity in get_settings().openviking_user_keys:
        mapped_tenant, _, mapped_user = identity.partition(":")
        if mapped_tenant != tenant:
            continue
        await enqueue(
            tenant,
            mapped_user,
            "skill",
            skill.id,
            {
                "name": skill.name,
                "description": skill.description,
                "version": skill.current_version,
                "content": content,
            },
        )


async def sync_skill_state(tenant: str | None, skill: Any, repo: Any) -> None:
    if not tenant or not get_settings().openviking_sync_enabled:
        return
    try:
        if skill.status == "published" and skill.current_version:
            version = await repo.get_version(tenant, skill.id, skill.current_version)
            if version and version.approved_by:
                await queue_skill(tenant, skill.owner or "", skill, version.spec)
        else:
            for identity in get_settings().openviking_user_keys:
                mapped_tenant, _, mapped_user = identity.partition(":")
                if mapped_tenant == tenant:
                    await enqueue(tenant, mapped_user, "delete_skill", skill.id, {})
    except Exception as exc:
        logger.warning("OpenViking skill mirror unavailable: %s", type(exc).__name__)


async def queue_run(ctx: Any, record: Any) -> None:
    """Mirror actual deliverables and conversation; never break the native result."""
    if not get_settings().openviking_sync_enabled:
        return
    try:
        from forgeflow.documents.store import DocArtifactStore

        for artifact in record.artifacts:
            ref = artifact.get("content_ref")
            if ref:
                path = DocArtifactStore().resolve(ref)
                content = await asyncio.to_thread(_text, path.read_bytes(), path.name)
            else:
                content = artifact.get("content")
                if not isinstance(content, str) or not content:
                    continue
            await enqueue(
                ctx.tenant_id,
                ctx.user_id or "",
                "resource",
                artifact["id"],
                {
                    "name": artifact.get("title") or "任务产物",
                    "content": content[:200000],
                    "run_id": record.run_id,
                    "version_ref": ref,
                },
            )
        answer = record.llm.get("final_answer", {})
        conversation = {
            "intent": record.intent,
            "answer": str(answer.get("text") or "") if isinstance(answer, dict) else "",
            "outcome": record.outcome,
            "experience_id": getattr(record, "experience_id", None),
            "learning_eligible": record.outcome == "success"
            and (
                any(artifact.get("content_ref") for artifact in record.artifacts)
                or sum(bool(step.get("tool")) for step in getattr(record, "steps", [])) >= 2
                or bool(re.search(r"记住|remember", record.intent, re.IGNORECASE))
            ),
        }
        await enqueue(ctx.tenant_id, ctx.user_id or "", "episode", record.run_id, conversation)
        await enqueue(
            ctx.tenant_id,
            ctx.user_id or "",
            "run",
            record.run_id,
            conversation,
        )
    except Exception as exc:
        logger.warning("OpenViking result enqueue unavailable: %s", type(exc).__name__)


async def compile_learned_pattern(job: dict[str, Any]) -> dict[str, Any]:
    """Only own successful, scrubbed, same-tool-pattern experiences form a draft."""
    from forgeflow.repositories import get_experience_repository
    from forgeflow.skills.candidate_compiler import compile_candidate

    if job["payload"].get("outcome") != "success":
        return {"status": "skipped", "reason": "unsuccessful_run", "published": False}

    root = Path(get_settings().openviking_outbox_root)
    ids = [job["payload"]["experience_id"]]
    for path in root.glob("*.json"):
        other = json.loads(path.read_text(encoding="utf-8"))
        if (other["tenant"], other["user"], other["kind"]) != (job["tenant"], job["user"], "run"):
            continue
        if (
            other["status"] == "done"
            and other["payload"].get("outcome") == "success"
            and other["payload"].get("experience_id")
        ):
            ids.append(other["payload"]["experience_id"])
    repo = get_experience_repository()
    seed = await repo.get(job["tenant"], job["payload"].get("experience_id", ""))

    def signature(e: Any) -> tuple:
        return tuple(s.get("tool") for s in e.reusable_steps if s.get("tool"))

    if not seed or not signature(seed):
        return {"status": "insufficient", "reason": "no_reusable_tool_pattern"}
    selected = []
    for experience_id in set(ids):
        experience = await repo.get(job["tenant"], experience_id)
        if (
            experience
            and signature(experience) == signature(seed)
            and experience.scrub_status == "scrubbed"
        ):
            selected.append(experience_id)

    class ScopedExperiences:
        async def list(self, tenant: str, **kwargs: Any) -> Sequence[Any]:
            return [record for record in await repo.list(tenant, **kwargs) if record.id in selected]

        async def find_similar(
            self, tenant: str, embedding: Any, **kwargs: Any
        ) -> Sequence[tuple[Any, float]]:
            return [
                (record, score)
                for record, score in await repo.find_similar(tenant, embedding, **kwargs)
                if record.id in selected
            ]

        async def get(self, tenant: str, experience_id: str) -> Any:
            return await repo.get(tenant, experience_id) if experience_id in selected else None

    candidate = (
        await compile_candidate(job["tenant"], mode="auto", experience_repo=ScopedExperiences())
        if selected
        else None
    )
    return {
        "status": candidate.status if candidate else "insufficient",
        "candidate_id": candidate.id if candidate and candidate.status == "draft" else None,
        "experience_count": len(selected),
        "published": False,
    }


async def _process(client: httpx.AsyncClient, job: dict[str, Any]) -> dict[str, Any]:
    user, payload = job["user"], job["payload"]
    stable = hashlib.sha256(job["id"].encode()).hexdigest()[:32]
    root = f"viking://user/{user}/"

    async def request(method: str, path: str, **kwargs: Any) -> Any:
        response = await client.request(method, path, **kwargs)
        if response.status_code == 404 and method == "DELETE":
            return {"deleted": True}
        response.raise_for_status()
        body = response.json()
        if body.get("status") != "ok":
            raise ValueError("remote_operation_failed")
        return body["result"]

    kind = job["kind"]
    if (
        kind == "run"
        and payload.get("intent") is not None
        and not payload.get(
            "learning_eligible", bool(re.search(r"记住|remember", payload["intent"], re.IGNORECASE))
        )
    ):
        previous = job.get("result", {})
        if previous.get("task_id"):
            response = await client.get("/api/v1/tasks/" + previous["task_id"])
            if response.status_code == 200 and response.json()["result"]["status"] not in {
                "completed",
                "failed",
                "cancelled",
            }:
                await client.post("/api/v1/tasks/" + previous["task_id"] + "/cancel", json={})
                return {**previous, "state": "waiting", "learning": "not_needed"}
        sid = "ff-" + stable
        if not previous.get("session_id"):
            created = await client.post("/api/v1/sessions", json={"session_id": sid})
            if created.status_code not in {200, 201, 409}:
                created.raise_for_status()
            info = await request("GET", "/api/v1/sessions/" + sid)
            count = info.get("total_message_count", info.get("message_count", 0))
            if count == 0:
                await request(
                    "POST",
                    f"/api/v1/sessions/{sid}/messages/batch",
                    json={
                        "messages": [
                            {"role": "user", "content": payload["intent"]},
                            {"role": "assistant", "content": payload["answer"]},
                        ]
                    },
                )
        return {
            "session_id": previous.get("session_id") or sid,
            "state": "complete",
            "learning": "not_needed",
            "reason": "conversation_stored_extraction_not_required",
        }
    if kind.startswith("delete_") and kind != "delete_skill":
        for task_id in payload.get("task_ids", []):
            task = await client.get("/api/v1/tasks/" + task_id)
            if task.status_code == 404:
                continue
            task.raise_for_status()
            state = task.json()["result"]["status"]
            if state not in {"completed", "cancelled", "failed"}:
                cancel = await client.post("/api/v1/tasks/" + task_id + "/cancel", json={})
                if cancel.status_code not in {200, 202, 404, 409}:
                    cancel.raise_for_status()
                return {"state": "waiting", "cancelling_task": task_id}
    if kind == "delete_episode":
        uri = root + "memories/forgeflow/episodes/" + stable
        await request("DELETE", "/api/v1/fs", params={"uri": uri, "recursive": "true"})
        return {"state": "deleted", "memory_uri": uri}
    if kind == "delete_skill":
        name = "ff-" + stable
        existing = await client.get(
            "/api/v1/skills/" + name, params={"target_uri": root + "skills"}
        )
        if existing.status_code == 404:
            return {"state": "deleted", "uri": root + "skills/" + name}
        existing.raise_for_status()
        if existing.json()["result"].get("root_uri") != root + "skills/" + name:
            raise ValueError("skill_scope_mismatch")
        deleted = await request(
            "DELETE", "/api/v1/skills/" + name, params={"target_uri": root + "skills"}
        )
        return {"state": "deleted", "delete": deleted, "uri": root + "skills/" + name}
    if kind == "episode":
        uri = root + "memories/forgeflow/episodes/" + stable
        if not job.get("result", {}).get("memory_uri"):
            content = f"# 来源会话\n\n来源运行：{job['id']}\n状态：{payload['outcome']}\n\n## 用户原始要求\n{payload['intent']}\n\n## 实际回答\n{payload['answer']}\n\n这是来源记录，未经事实提炼；失败回答不能作为已验证事实。"
            await request(
                "POST",
                "/api/v1/content/write",
                json={
                    "uri": uri + "/source.md",
                    "content": content,
                    "mode": "replace",
                    "wait": False,
                },
            )
            return {"memory_uri": uri, "state": "waiting", "memory_strategy": "source_episode"}
        found = await request(
            "POST",
            "/api/v1/search/find",
            json={"query": payload["intent"][:4000], "target_uri": uri, "limit": 3},
        )
        hits = [
            item
            for category in ("memories", "resources", "skills")
            for item in found.get(category, [])
        ]
        indexed = any(item.get("uri", "").startswith(uri + "/") for item in hits)
        return {
            "memory_uri": uri,
            "state": "complete" if indexed else "waiting",
            "memory_strategy": "source_episode",
            "indexed": indexed,
        }
    if kind == "delete_run":
        results = []
        for sid in payload.get("sessions", []):
            if not sid.startswith("ff-") or any(c in sid for c in "/\\%?#"):
                raise ValueError("invalid_session_reference")
            results.append(
                await request(
                    "DELETE",
                    "/api/v1/fs",
                    params={"uri": root + "sessions/" + sid, "recursive": "true"},
                )
            )
        return {"state": "deleted", "sessions_deleted": len(results)}
    if job.get("result", {}).get("task_id"):
        try:
            result = await request("GET", "/api/v1/tasks/" + job["result"]["task_id"])
        except httpx.HTTPStatusError as exc:
            if exc.response.status_code != 404:
                raise
            job.setdefault("failed_tasks", []).append(
                {"task_id": job["result"]["task_id"], "status": "unavailable"}
            )
            job["result"] = {}
            raise ValueError("remote_task_missing") from exc
        state = result.get("status", result.get("state"))
        if state in {"failed", "cancelled"}:
            job.setdefault("failed_tasks", []).append(result)
            job["result"] = {"previous_task_id": result.get("task_id")}
            raise ValueError("remote_extraction_failed")
        if state != "completed" and kind in {"resource", "skill"}:
            own_indexed = job["result"].get("own_indexed", False)
            if not own_indexed:
                uri = job["result"]["uri"]
                found = await request(
                    "POST",
                    "/api/v1/search/find",
                    json={
                        "query": payload["name"],
                        "target_uri": uri,
                        "limit": 3,
                        "level": [0, 1, 2],
                    },
                )
                indexed = any(
                    item.get("uri", "").startswith(uri + "/")
                    for category in ("memories", "resources", "skills")
                    for item in found.get(category, [])
                )
                if indexed:
                    content = await request(
                        "GET",
                        "/api/v1/content/read",
                        params={"uri": uri + ("/SKILL.md" if kind == "skill" else "/source.md")},
                    )
                    own_indexed = f"Source ID: {job['id']}" in str(content) and (
                        kind != "skill" or f"Version: {payload.get('version')}" in str(content)
                    )
            return {
                **job["result"],
                "state": "available" if own_indexed else "waiting",
                "own_indexed": own_indexed,
                "task": result,
                "ancestor_refresh_pending": True,
            }
        return {
            **job["result"],
            "state": "complete" if state == "completed" else "waiting",
            "task": result,
        }
    if kind == "skill":
        name = "ff-" + stable
        existing = await client.get(
            "/api/v1/skills/" + name, params={"target_uri": root + "skills"}
        )
        if existing.status_code == 404:
            method, endpoint = "POST", "/api/v1/skills"
        else:
            existing.raise_for_status()
            remote_uri = existing.json().get("result", {}).get("root_uri", "")
            if remote_uri != root + "skills/" + name:
                raise ValueError("skill_scope_mismatch")
            method, endpoint = "PUT", "/api/v1/skills/" + name
        result = await request(
            method,
            endpoint,
            json={
                "data": {
                    "name": name,
                    "description": payload.get("description") or payload["name"],
                    "content": f"# {payload['name']}\n\nSource ID: {job['id']}\nVersion: {payload.get('version')}\n\n{payload['content']}",
                },
                "target_uri": root + "skills",
                "wait": False,
            },
        )
        return {
            "uri": root + "skills/" + name,
            "write": result,
            "state": "waiting" if result.get("task_id") else "complete",
            "task_id": result.get("task_id"),
        }
    if kind == "resource":
        uri = root + "resources/forgeflow/" + stable
        body = payload.get("content")
        if body is None and payload.get("source_ref"):
            from forgeflow.resources.storage import FileBlobStore

            body = await asyncio.to_thread(
                _text, FileBlobStore().read(payload["source_ref"]), payload["name"]
            )
        if body is None:
            raise ValueError("resource_content_unavailable")
        content = f"# {payload['name']}\n\nSource ID: {job['id']}\n\n{body[:200000]}"
        uploaded = await request(
            "POST",
            "/api/v1/resources/temp_upload",
            files={"file": ("source.md", content.encode("utf-8"), "text/markdown")},
        )
        result = await request(
            "POST",
            "/api/v1/resources",
            json={
                "temp_file_id": uploaded["temp_file_id"],
                "to": uri,
                "create_parent": True,
                "wait": False,
            },
        )
        if not result.get("task_id"):
            raise ValueError("resource_import_missing_task")
        return {"uri": uri, "write": result, "task_id": result["task_id"], "state": "waiting"}
    if kind in {"delete_resource", "delete_skill"}:
        uri = root + (
            "resources/forgeflow/" + stable if kind == "delete_resource" else "skills/ff-" + stable
        )
        result = await request("DELETE", "/api/v1/fs", params={"uri": uri, "recursive": "true"})
        return {"uri": uri, "delete": result, "state": "deleted"}
    if kind == "run":
        sid = (
            "ff-"
            + stable
            + (f"-retry-{len(job['failed_tasks'])}" if job.get("failed_tasks") else "")
        )
        response = await client.post("/api/v1/sessions", json={"session_id": sid})
        if response.status_code not in (200, 201, 409):
            response.raise_for_status()
        info = await request("GET", "/api/v1/sessions/" + sid)
        count = info.get("total_message_count", info.get("message_count", 0))
        if count < 1:
            await request(
                "POST",
                f"/api/v1/sessions/{sid}/messages",
                json={"role": "user", "content": payload["intent"]},
            )
        if count < 2:
            await request(
                "POST",
                f"/api/v1/sessions/{sid}/messages",
                json={
                    "role": "assistant",
                    "content": payload["answer"]
                    + "\n\n真实任务状态："
                    + payload["outcome"]
                    + "\n来源运行："
                    + job["id"],
                },
            )
        commit = (
            await request("POST", f"/api/v1/sessions/{sid}/commit", json={})
            if not info.get("commit_count")
            else {"status": "already_committed", "task_id": info.get("task_id")}
        )
        if commit.get("status") == "already_committed" and not commit.get("task_id"):
            tasks = await request("GET", "/api/v1/tasks", params={"limit": 100})
            items = tasks.get("items", tasks.get("tasks", [])) if isinstance(tasks, dict) else tasks
            if not isinstance(items, list) or any(not isinstance(item, dict) for item in items):
                raise ValueError("invalid_extraction_task_list")
            matching = [
                t
                for t in items
                if t.get("resource_id") == sid and t.get("task_type") == "session_commit"
            ]
            if not matching:
                raise ValueError("committed_session_missing_extraction_task")
            commit["task_id"] = matching[0]["task_id"]
        return {
            "session_id": sid,
            "commit": commit,
            "task_id": commit.get("task_id"),
            "state": "waiting" if commit.get("task_id") else "complete",
        }
    raise ValueError("unsupported_sync_kind")


async def drain_once() -> int:
    settings = get_settings()
    root = Path(settings.openviking_outbox_root)
    if not settings.openviking_sync_enabled or not root.exists():
        return 0
    count = 0
    processed = 0
    paths = list(root.glob("*.json"))
    queued = {path: json.loads(path.read_text(encoding="utf-8")) for path in paths}
    active = sum(
        job["kind"] in {"resource", "skill", "episode"} and job["status"] == "waiting"
        for job in queued.values()
    )
    active_learning = sum(
        job["kind"] == "run" and job["status"] == "waiting" for job in queued.values()
    )

    def priority(path: Path) -> tuple:
        return (
            0
            if queued[path]["kind"].startswith("delete_")
            else 1
            if queued[path]["status"] == "waiting"
            else 2,
            path.stat().st_mtime,
        )

    for path in sorted(paths, key=priority):
        async with _lock:
            job = json.loads(path.read_text(encoding="utf-8"))
            if job["status"] in {"done", "cancelled"} or job.get("attempts", 0) >= 10:
                continue
            if job.get("next_attempt_at", 0) > time.time():
                continue
            if (
                job["kind"] == "run"
                and not job.get("result", {}).get("task_id")
                and active_learning >= 1
            ):
                continue
            if (
                job["kind"] in {"resource", "skill", "episode"}
                and not (
                    job.get("result", {}).get("task_id") or job.get("result", {}).get("memory_uri")
                )
                and active >= 1
            ):
                continue
            key = settings.openviking_user_keys.get(f"{job['tenant']}:{job['user']}")
            if key is None:
                continue
        processed += 1
        if processed > 50:
            break
        try:
            async with httpx.AsyncClient(
                base_url=settings.openviking_url,
                headers={"X-API-Key": key.get_secret_value()},
                trust_env=False,
                follow_redirects=False,
                timeout=120,
            ) as client:
                job["result"] = await _process(client, job)
            if job["kind"] in {"resource", "skill", "episode"}:
                if job["status"] == "waiting" and job["result"].get("state") != "waiting":
                    active -= 1
                elif job["status"] != "waiting" and job["result"].get("state") == "waiting":
                    active += 1
            if job["kind"] == "run":
                if job["status"] == "waiting" and job["result"].get("state") != "waiting":
                    active_learning -= 1
                elif job["status"] != "waiting" and job["result"].get("state") == "waiting":
                    active_learning += 1
            job["status"] = {"waiting": "waiting", "available": "ready"}.get(
                job["result"].get("state"), "done"
            )
            if (
                job["status"] == "done"
                and job["kind"] == "run"
                and job["payload"].get("experience_id")
            ):
                try:
                    job["result"]["skill_learning"] = await compile_learned_pattern(job)
                except Exception as exc:
                    job["result"]["skill_learning"] = {
                        "status": "unavailable",
                        "error": type(exc).__name__,
                    }
            job.pop("error", None)
            job.pop("next_attempt_at", None)
            count += 1
        except (httpx.HTTPError, ValueError, KeyError) as exc:
            job["attempts"] += 1
            job["status"] = "failed" if job["attempts"] >= 10 else "retry"
            job["error"] = type(exc).__name__
            job["next_attempt_at"] = time.time() + min(900, 15 * 2 ** min(job["attempts"], 6))
        job["updated_at"] = datetime.now(UTC).isoformat()
        async with _lock:
            current = json.loads(path.read_text(encoding="utf-8"))
            if current["payload"] != job["payload"] or current["status"] == "cancelled":
                continue
            _save(path, job)
    return count


async def worker() -> None:
    while True:
        try:
            await drain_once()
        except Exception as exc:
            logger.warning("OpenViking sync worker retry: %s", type(exc).__name__)
        await asyncio.sleep(5)
