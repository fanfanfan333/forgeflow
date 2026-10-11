"""Bounded, read-only previews of uploaded binary documents."""

from io import BytesIO
from itertools import islice
from pathlib import Path
from typing import Any


def preview_binary(data: bytes, name: str, limit: int) -> dict[str, Any]:
    base: dict[str, Any] = {
        "available": True,
        "format": "text",
        "columns": [],
        "rows": [],
        "content": "",
        "truncated": False,
        "note": "",
    }
    lines: list[str] = []
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".docx":
            from docx import Document

            doc = Document(BytesIO(data))
            for block in doc.iter_inner_content():
                if hasattr(block, "text"):
                    lines.extend(block.text.splitlines() or [""])
                else:
                    for row in block.rows:
                        lines.append(" | ".join(cell.text for cell in row.cells))
                        if len(lines) > limit:
                            break
                if len(lines) > limit:
                    break
        elif suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            book = load_workbook(BytesIO(data), read_only=True, data_only=False)
            try:
                sheet = book.active
                if sheet is None:
                    raise ValueError("工作簿没有可读取的工作表")
                rows = list(
                    islice(
                        sheet.iter_rows(max_col=min(sheet.max_column or 50, 50), values_only=True),
                        limit + 2,
                    )
                )

                def convert(row: tuple[Any, ...]) -> list[str]:
                    return ["" if cell is None else str(cell)[:4000] for cell in row]

                base.update(
                    format="table",
                    columns=convert(rows[0]) if rows else [],
                    rows=[convert(row) for row in rows[1 : limit + 1]],
                    truncated=len(rows) > limit + 1,
                    note=f"工作表：{sheet.title}；最多显示 50 列，公式按原文显示",
                )
                return base
            finally:
                book.close()
        elif suffix == ".pptx":
            from pptx import Presentation

            deck = Presentation(BytesIO(data))
            for index, slide in enumerate(deck.slides):
                lines.append(f"第 {index + 1} 页")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        lines.extend(shape.text.splitlines())
                    if len(lines) > limit:
                        break
                if len(lines) > limit:
                    break
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            for page in islice(reader.pages, 5):
                lines.extend((page.extract_text() or "").splitlines())
                if len(lines) > limit:
                    break
            base["note"] = "文本预览最多读取前 5 页；扫描件需使用文字识别"
            base["truncated"] = len(reader.pages) > 5
        else:
            raise ValueError("该文档类型暂不支持内容预览")
        base["truncated"] = base["truncated"] or len(lines) > limit or sum(map(len, lines)) > 40000
        base["content"] = "\n".join(lines[:limit])[:40000]
        if not base["content"].strip():
            base.update(available=False, note="未提取到可显示的文字；文件可能为空或为扫描件")
        return base
    except ImportError:
        base.update(available=False, note="预览依赖未安装，请检查该文档类型的处理依赖")
    except Exception as exc:
        base.update(
            available=False, note=f"无法解析文件（{type(exc).__name__}），请检查文件是否完整或加密"
        )
    return base
