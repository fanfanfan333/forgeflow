"""INC45 — real XLSX structural inspection via ``openpyxl``.

Layer contract (design §1.1 / §8): this module **only reads** bytes. It never
writes a file and never calls a model — it is the honest "inspect" layer that
reports a workbook's *measured* structure (worksheet names, the selected
sheet's row/column counts, its header row, non-empty cell count and character
count). Nothing is guessed: bytes ``openpyxl`` cannot open raise
:class:`SheetInspectionError`, which the handler converts into an honest failure.

Why a dedicated module (and reader)
----------------------------------
The resource-summary read path (:func:`forgeflow.resources.summaries.summarize_excel`)
is a deliberately pure-stdlib ``zipfile`` + XML reader and is pinned by the
resource tests — it is **not** touched here. This module is the *write-plane's*
reader: it uses ``openpyxl`` (the OOXML write fact standard) so the structure the
write plane edits and the structure it reports can never disagree.

Dependency posture: ``openpyxl`` is an optional extra (``xlsx``). :func:`open_workbook`
imports it lazily; when it is absent the call raises :class:`SheetInspectionError`
with a verbatim reason (the caller degrades honestly — ``not_executed``, HTTP < 500),
mirroring the ``.docx`` / ``python-docx`` posture exactly.

Honesty rules:
  * a numeric fact is ``None`` when it could not be measured — never a fake ``0``;
  * formulas are read as formulas (``data_only=False``) so an inspect never
    silently reports a cached value the write plane would then clobber.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from typing import Any

__all__ = [
    "SheetStructure",
    "SheetInspectionError",
    "open_workbook",
    "inspect_sheet",
    "sheet_numbers",
]


class SheetInspectionError(ValueError):
    """Raised when bytes are not a readable XLSX (never silently tolerated)."""


def _has_vba(data: bytes) -> bool:
    """Whether an OOXML package really carries a VBA project (``.xlsm`` macro)."""
    try:
        with zipfile.ZipFile(io.BytesIO(bytes(data))) as archive:
            return any(name.endswith("vbaProject.bin") for name in archive.namelist())
    except Exception:  # noqa: BLE001 — a non-zip blob simply has no VBA part
        return False


def open_workbook(data: bytes) -> Any:
    """Open ``data`` as an ``openpyxl`` ``Workbook`` (or raise honestly).

    This is the single open path used by inspect / edit / validation, so an
    unreadable byte string (or a missing ``openpyxl`` extra) is reported
    identically everywhere. ``data_only=False`` keeps **formulas as formulas**
    (design §1.1 P2-2): an inspect never reports a cached value, and the write
    plane can preserve the formula on every untouched cell.
    """
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise SheetInspectionError("XLSX 字节为空")
    try:  # imported lazily so the module stays import-safe without the extra
        import openpyxl

        payload = io.BytesIO(bytes(data))
        # ``keep_vba`` only matters for a macro-enabled workbook; sniffing the
        # package avoids asking openpyxl to keep a VBA part that is not there.
        return openpyxl.load_workbook(payload, data_only=False, keep_vba=_has_vba(bytes(data)))
    except ImportError as exc:
        raise SheetInspectionError(f"xlsx 支持不可用：{exc}") from exc
    except Exception as exc:  # noqa: BLE001 — any parse failure is the point
        raise SheetInspectionError(f"无法解析 XLSX：{exc}") from exc


def _count_cells_and_chars(worksheet: Any) -> tuple[int, int]:
    """Return ``(non_empty_cells, char_count)`` over the worksheet's real cells.

    A cell counts only when it really carries a value (a ``None`` / empty string
    is not a cell); ``char_count`` is the summed length of every counted cell's
    string form, so a ``≤N 字`` requirement is measured, never estimated.
    """
    cells = 0
    chars = 0
    for row in worksheet.iter_rows(values_only=True):
        for value in row:
            if value is None or (isinstance(value, str) and value == ""):
                continue
            cells += 1
            chars += len(str(value))
    return cells, chars


@dataclass
class SheetStructure:
    """The measured structure of an XLSX workbook."""

    sheets: list[str] = field(default_factory=list)
    active: str = ""
    #: Row / column counts of the inspected sheet — ``None`` when openpyxl did
    #: not report a dimension (never a fabricated ``0``).
    rows: int | None = None
    columns: int | None = None
    header: list[str] = field(default_factory=list)
    cells: int | None = None
    char_count: int | None = None

    def to_dict(self, limit: int | None = None) -> dict[str, Any]:
        """JSON-safe structure summary.

        ``limit`` bounds ``header`` so a very wide sheet cannot blow past the
        executor's payload ceiling; when it bites, ``header_count`` is reported
        verbatim alongside ``"truncated": True`` so the cap is explicit rather
        than a silent loss.
        """
        header = list(self.header)
        truncated = False
        if limit is not None and limit >= 0 and len(header) > limit:
            header = header[:limit]
            truncated = True
        out: dict[str, Any] = {
            "sheets": list(self.sheets),
            "sheet_count": len(self.sheets),
            "active": self.active,
            "rows": self.rows,
            "columns": self.columns,
            "header_count": len(self.header),
            "header": header,
            "cells": self.cells,
            "char_count": self.char_count,
        }
        if truncated:
            out["truncated"] = True
        return out


def inspect_sheet(data: bytes, *, sheet: str | None = None) -> SheetStructure:
    """Read the **real** structure of an XLSX byte string.

    ``sheet`` selects a worksheet by name; when it is ``None`` the workbook's
    active worksheet is inspected. A named worksheet that does not exist raises
    :class:`SheetInspectionError` (the target is never guessed).

    Raises:
        SheetInspectionError: ``data`` is empty / unreadable, the workbook has no
            worksheet, or ``sheet`` names a worksheet that does not exist.
    """
    workbook = open_workbook(data)
    names = [str(name) for name in workbook.sheetnames]
    if not names:
        raise SheetInspectionError("工作簿没有任何工作表")
    if sheet:
        wanted = str(sheet)
        if wanted not in names:
            raise SheetInspectionError(f"工作表不存在：{wanted}")
        worksheet = workbook[wanted]
    else:
        worksheet = workbook.active
    if worksheet is None:  # pragma: no cover — defensive; a named sheet always exists
        raise SheetInspectionError("工作簿没有任何工作表")

    max_row = getattr(worksheet, "max_row", None)
    max_column = getattr(worksheet, "max_column", None)
    rows = int(max_row) if isinstance(max_row, int) else None
    columns = int(max_column) if isinstance(max_column, int) else None

    first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    header = ["" if value is None else str(value) for value in first_row]
    cells, char_count = _count_cells_and_chars(worksheet)

    return SheetStructure(
        sheets=names,
        active=str(worksheet.title),
        rows=rows,
        columns=columns,
        header=header,
        cells=cells,
        char_count=char_count,
    )


def sheet_numbers(data: bytes) -> list[str]:
    """Every numeric token across the workbook's cell text, in sheet/cell order.

    Used by the edit handler to build the "numbers that must survive" guard: the
    original numbers minus the ones an explicit edit legitimately liberates. The
    numeric rule is re-used from :mod:`forgeflow.documents.docx_inspect` so the
    whole document domain shares exactly one definition of "what is a number".
    """
    from forgeflow.documents.docx_inspect import numbers_in_text

    workbook = open_workbook(data)
    out: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            for value in row:
                if value is None:
                    continue
                out.extend(numbers_in_text(str(value)))
    return out
