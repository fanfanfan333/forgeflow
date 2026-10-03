"""INC46 T23 — **L2** invariant validation (编辑前后不变量比对).

What L2 checks
--------------
L2 answers the one question that matters for a *bounded* edit: **did everything
that had to survive actually survive?** It compares the :class:`InvariantSet` of
the located range — as measured by the frozen T20
:func:`forgeflow.documents.invariants.extract_invariants` — *before* and *after*
the edit, and it does so for **both** halves of the T20 contract:

* **explicit invariants** — every amount / date / auto-number / reference / image
  / table that really occurs inside the range must still be there afterwards;
* **implicit baseline** (always in force, independent of the instruction) —
  1. **目标区间外内容零变化**  2. **章节数不变**  3. **不新增数字**.

On top of that L2 compares **per-cell table text + merge structure**
(``w:gridSpan`` / ``w:vMerge``) so a silently merged/unmerged or re-typed cell is
caught even when the document still opens.

Interval semantics (consumes T20's locator)
-------------------------------------------
The range is ``[start, end)`` in **paragraph-index space** — exactly the interval
a :class:`~forgeflow.documents.locator.LocatorResult` yields
(``chosen.index`` … ``chosen.section_end_index``)."Outside the range" is measured
**anchored from both ends** (the head ``[0, start)`` and the tail
``[end, len)``), so a compliant edit that grows the range in place never
mis-aligns the untouched head/tail.

Honesty (red line 4 / 15)
-------------------------
L2 never invents a value: it compares *measured* structures only. A layer that
cannot read the bytes raises honestly (the caller turns that into an honest
``fail``); it never reports an unmeasured dimension as ``pass``. Recognition is
**not** re-implemented here — L2 calls T20's ``extract_invariants`` and the
frozen ``docx_inspect`` helpers; it must never fork the recognisers or touch
``docx_inspect._NUMERIC_RE``.
"""

from __future__ import annotations

from collections import Counter
from typing import Any

from forgeflow.documents.docx_inspect import (
    document_numbers,
    heading_level,
    open_docx,
)
from forgeflow.documents.invariants import (
    AMOUNT_KIND,
    DATE_KIND,
    IMAGE_KIND,
    NUMBERING_KIND,
    REFERENCE_KIND,
    TABLE_KIND,
    InvariantSet,
    extract_invariants,
)
from forgeflow.documents.validation.verdict import (
    FAIL,
    LAYER_INVARIANTS,
    PASS,
    LayerVerdict,
)

__all__ = [
    "InvariantComparison",
    "invariants_verdict",
    "paragraph_texts",
    "table_signatures",
]

#: Explicit invariant kinds whose ``value`` must **survive** the edit verbatim.
_SURVIVING_KINDS: tuple[str, ...] = (
    AMOUNT_KIND,
    DATE_KIND,
    NUMBERING_KIND,
    REFERENCE_KIND,
    IMAGE_KIND,
    TABLE_KIND,
)


# --------------------------------------------------------------------------- #
# document readers (shared with L3)                                            #
# --------------------------------------------------------------------------- #
def paragraph_texts(source: bytes) -> list[str]:
    """Body-paragraph texts of ``source`` (single open; the range space L2/L3 use)."""
    return [paragraph.text for paragraph in open_docx(source).paragraphs]


def _cell_merge(tc: Any) -> tuple[int, str]:
    """``(gridSpan, vMerge)`` for a table cell's ``w:tc`` element.

    ``gridSpan`` is the horizontal span (``1`` when absent); ``vMerge`` is the
    literal ``w:val`` (``"restart"`` / ``"continue"``) or ``""`` when the cell is
    not part of a vertical merge.
    """
    from docx.oxml.ns import qn

    tcPr = tc.find(qn("w:tcPr"))
    span = 1
    vmerge = ""
    if tcPr is not None:
        grid = tcPr.find(qn("w:gridSpan"))
        if grid is not None and grid.get(qn("w:val")) is not None:
            try:
                span = int(grid.get(qn("w:val")))
            except (TypeError, ValueError):
                span = 1
        vmerge_el = tcPr.find(qn("w:vMerge"))
        if vmerge_el is not None:
            vmerge = vmerge_el.get(qn("w:val")) or "continue"
    return span, vmerge


def _table_signature(table: Any) -> list[list[dict[str, Any]]]:
    """Per-cell signature of a table: text + ``gridSpan`` + ``vMerge`` (+ nested)."""
    rows: list[list[dict[str, Any]]] = []
    for row in table.rows:
        cells: list[dict[str, Any]] = []
        for cell in row.cells:
            span, vmerge = _cell_merge(cell._tc)
            cells.append(
                {
                    "text": cell.text,
                    "gridSpan": span,
                    "vMerge": vmerge,
                    "nested": [_table_signature(child) for child in cell.tables],
                }
            )
        rows.append(cells)
    return rows


def table_signatures(source: bytes) -> list[list[list[dict[str, Any]]]]:
    """Per-cell signatures of every body table (text + merge structure)."""
    return [_table_signature(table) for table in open_docx(source).tables]


def _heading_count(source: bytes) -> int:
    """Body headings (the rule ``inspect_docx`` groups sections by)."""
    return sum(
        1
        for paragraph in open_docx(source).paragraphs
        if heading_level(paragraph) is not None and paragraph.text.strip()
    )


# --------------------------------------------------------------------------- #
# comparison result                                                            #
# --------------------------------------------------------------------------- #
class InvariantComparison:
    """The measured, per-item outcome of an L2 before/after comparison.

    Every field is a **measured** fact. ``failures`` is the honest, order-stable
    list of concrete violations found (empty ⇒ nothing was violated).
    """

    def __init__(
        self,
        *,
        compared: int,
        survived: int,
        lost: list[str],
        out_of_range_changes: list[str],
        heading_delta: int,
        new_numbers: list[str],
        table_changes: list[str],
    ) -> None:
        self.compared = compared
        self.survived = survived
        self.lost = lost
        self.out_of_range_changes = out_of_range_changes
        self.heading_delta = heading_delta
        self.new_numbers = new_numbers
        self.table_changes = table_changes

    @property
    def failures(self) -> list[str]:
        """Every concrete violation (order: survival → scope → baseline → tables)."""
        out: list[str] = []
        out.extend(f"lost:{item}" for item in self.lost)
        out.extend(f"out-of-range:{item}" for item in self.out_of_range_changes)
        if self.heading_delta:
            out.append(f"headings-delta:{self.heading_delta}")
        out.extend(f"new-number:{item}" for item in self.new_numbers)
        out.extend(f"table:{item}" for item in self.table_changes)
        return out

    def evidence(self) -> str:
        if self.failures:
            return "invariants:FAIL[" + ";".join(self.failures[:8]) + "]"
        # An empty denominator is empty *coverage* (no explicit invariant in range),
        # not "all survived": report it as ``n/a`` so ``0/0`` is never misread as a
        # full pass (红线 4). The numeric ``compared``/``survived`` stay in ``detail``.
        if self.compared == 0:
            survived_note = "survived=n/a (no explicit invariant in range)"
        else:
            survived_note = f"survived={self.survived}/{self.compared}"
        return (
            f"invariants:{survived_note};"
            f"out-of-range=ok;headings=ok;no-new-numbers=ok;tables=unchanged"
        )

    def detail(self) -> dict[str, Any]:
        return {
            "compared": self.compared,
            "survived": self.survived,
            "lost": list(self.lost),
            "out_of_range_changes": list(self.out_of_range_changes),
            "heading_delta": self.heading_delta,
            "new_numbers": list(self.new_numbers),
            "table_changes": list(self.table_changes),
        }


# --------------------------------------------------------------------------- #
# range helpers                                                                #
# --------------------------------------------------------------------------- #
def _resolve_bounds(length: int, start: int, end: int | None) -> tuple[int, int]:
    """Clamp ``[start, end)`` to ``[0, length]`` (interval semantics)."""
    lo = max(0, min(int(start), length))
    hi = length if end is None else max(lo, min(int(end), length))
    return lo, hi


def _survival(before: InvariantSet, after: InvariantSet) -> tuple[int, int, list[str]]:
    """``(compared, survived, lost)`` for the explicit surviving-kind invariants."""
    after_values: dict[str, set[str]] = {}
    for item in after.items:
        after_values.setdefault(item.kind, set()).add(item.value)

    compared = 0
    survived = 0
    lost: list[str] = []
    for item in before.items:
        if not item.explicit or item.kind not in _SURVIVING_KINDS:
            continue
        compared += 1
        if item.value in after_values.get(item.kind, set()):
            survived += 1
        else:
            lost.append(f"{item.kind}={item.value}@{item.location}")
    return compared, survived, lost


def _out_of_range_changes(
    before_texts: list[str], after_texts: list[str], lo: int, hi: int
) -> list[str]:
    """Paragraph indices outside ``[lo, hi)`` whose text changed (head + tail).

    Anchored from both ends: the head is ``[0, lo)``; the tail is compared
    mirrored from the end so an in-range growth of the range does not shift it.
    """
    changes: list[str] = []
    head = min(lo, len(before_texts), len(after_texts))
    for index in range(head):
        if before_texts[index] != after_texts[index]:
            changes.append(f"paragraph[{index}]")

    tail_before = max(0, len(before_texts) - hi)
    tail_after = max(0, len(after_texts) - hi)
    tail = min(tail_before, tail_after)
    for offset in range(1, tail + 1):
        if before_texts[-offset] != after_texts[-offset]:
            changes.append(f"paragraph[tail-{offset}]")
    return changes


def _new_numbers(before: bytes, after: bytes) -> list[str]:
    """Numeric tokens that appear **more** times after than before (不新增数字)."""
    before_counts = Counter(document_numbers(before))
    after_counts = Counter(document_numbers(after))
    extra: list[str] = []
    for token, count in sorted(after_counts.items()):
        surplus = count - before_counts.get(token, 0)
        if surplus > 0:
            extra.append(f"{token}x{surplus}")
    return extra


def _table_changes(before: bytes, after: bytes) -> list[str]:
    """Table ordinals whose per-cell text / merge structure differs."""
    before_tables = table_signatures(before)
    after_tables = table_signatures(after)
    if before_tables == after_tables:
        return []
    if len(before_tables) != len(after_tables):
        return [
            f"count:{len(before_tables)}->{len(after_tables)}"
        ]
    return [
        f"table[{index}]"
        for index, (left, right) in enumerate(zip(before_tables, after_tables))
        if left != right
    ]


# --------------------------------------------------------------------------- #
# verdict                                                                      #
# --------------------------------------------------------------------------- #
def invariants_verdict(
    before: bytes,
    after: bytes,
    *,
    start: int = 0,
    end: int | None = None,
) -> LayerVerdict:
    """L2 — invariant verdict between ``before`` and ``after`` over ``[start, end)``.

    Args:
        before: the original DOCX bytes.
        after: the edited DOCX bytes.
        start: the located range's first paragraph index (inclusive).
        end: the located range's exclusive end paragraph index (``None`` ⇒ EOF).

    Returns:
        A :class:`LayerVerdict`; ``fail`` iff any explicit invariant was lost, any
        out-of-range paragraph changed, the heading count changed, a new number
        appeared, or a table's per-cell text / merge structure changed.
    """
    before_texts = paragraph_texts(before)
    after_texts = paragraph_texts(after)
    lo, hi = _resolve_bounds(len(before_texts), start, end)

    before_set = extract_invariants(before, start=lo, end=hi)
    after_set = extract_invariants(after, start=lo, end=hi)

    compared, survived, lost = _survival(before_set, after_set)
    scope = _out_of_range_changes(before_texts, after_texts, lo, hi)
    heading_delta = _heading_count(after) - _heading_count(before)
    new_numbers = _new_numbers(before, after)
    table_changes = _table_changes(before, after)

    comparison = InvariantComparison(
        compared=compared,
        survived=survived,
        lost=lost,
        out_of_range_changes=scope,
        heading_delta=heading_delta,
        new_numbers=new_numbers,
        table_changes=table_changes,
    )
    status = FAIL if comparison.failures else PASS
    return LayerVerdict(LAYER_INVARIANTS, status, comparison.evidence(), comparison.detail())
